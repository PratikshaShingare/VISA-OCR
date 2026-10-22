"""
Editable preview — Passport and Company Authorization Letters.

Fill the form, press Generate & Preview: the Preview section shows the real
generated page; its Edit button makes the letter's values editable right on
that page (no separate "Letter wording" form). Checks the real flow end to end
(real backend, real LibreOffice PDF, real downloads) — values open seeded with
the generator's own wording, edits persist, Download/Print regenerate with the
edit included, fixed template text is untouched (project rule 5), "Reset all"
restores the auto-generated wording.

Requires: the static file server AND the backend both running (see
run_e2e_suite.py).
"""

import asyncio
import io

from docx import Document
from playwright.async_api import async_playwright

from _common import (
    BASE,
    Checks,
    commit_preview_edit,
    edit_value_on_preview,
    fill_ocr_person,
    fresh_load_and_login,
    pick_company_consulate,
    set_edit_mode,
    start_guided_wizard_application,
    wait_preview_ready,
)

checks = Checks()


def _full_text(doc: Document) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


async def main():
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page(viewport={"width": 1280, "height": 1400})
        console_errors = []
        page.on("console", lambda m: console_errors.append(m.text) if m.type == "error" else None)
        page.on("pageerror", lambda e: console_errors.append(f"pageerror: {e}"))

        await fresh_load_and_login(page, base=BASE)
        await start_guided_wizard_application(page)
        await page.wait_for_timeout(300)

        # ---------- Step 1: applicant ----------
        await page.evaluate("location.hash = '#/new-application/1'")
        await page.wait_for_timeout(300)
        await fill_ocr_person(page, ".step-view[data-step='1']")

        # ---------- Step 2: one traveller (Spouse), so Company has 2
        # selectable applicants to exercise the dynamic field set with ----------
        await page.evaluate("location.hash = '#/new-application/2'")
        await page.wait_for_timeout(300)
        await page.click("[data-action='add-traveller']")
        await page.wait_for_timeout(300)
        await fill_ocr_person(page, ".traveller-panel", "[data-traveller-relation]", "Spouse")
        await page.click("[data-action='close-traveller-panel']")
        await page.wait_for_timeout(200)

        applicant_name = await page.evaluate(
            "(function(){var a=KhannaState.getActiveApplication().applicant; return (a.salutation?a.salutation+'. ':'')+a.fullName;})()"
        )
        # The passport keeps its capitals; letters write the name in ordinary
        # capitalisation ("Sample Testperson"), which is what is looked for below.
        applicant_name = applicant_name.title()

        # ---------- Step 6: Authorization Letters ----------
        await page.evaluate("location.hash = '#/new-application/6'")
        await page.wait_for_timeout(300)

        # ===================== Passport Authorization =====================
        passport_checks = page.locator('[data-auth-checklist="passport"] input[type=checkbox]')
        await passport_checks.nth(0).check()  # applicant only — exercises the single-traveller template
        await page.fill("#auth-passport-recipientCentreName", "Germany Visa Application Centre")
        await page.fill("#auth-passport-recipientAddress", "Andheri, Mumbai")
        await page.fill("#auth-passport-collectorName", "Mr. Suresh Patil")
        await page.locator("#auth-passport-collectorName").blur()
        await page.wait_for_timeout(300)

        checks.check(
            "The old 'Letter wording (editable)' section is gone from the Passport Authorization card",
            await page.locator("[data-passport-field-editors]").count() == 0
            and "Letter wording" not in await page.inner_text("[data-passport-auth-card]"),
        )

        # ---- Generate & Preview, then edit ON the preview ----
        pa_disabled = await page.get_attribute("[data-action='generate-authorization'][data-group='passport']", "disabled")
        checks.check("Passport Authorization Generate enabled", pa_disabled is None)
        scope_pa = "[data-doc-preview='passport']"
        await page.click("[data-action='generate-authorization'][data-group='passport']")
        await wait_preview_ready(page, scope_pa)
        edit_btn = page.locator(f"{scope_pa} [data-dv='edit-toggle']")
        checks.check("The preview has an Edit button", await edit_btn.count() == 1 and await edit_btn.is_enabled())
        checks.check("Nothing is editable (no live fields) while Edit is off", await page.locator(f"{scope_pa} .pe-field").first.is_hidden())
        await edit_btn.click()
        await page.wait_for_timeout(300)
        n_fields = await page.locator(f"{scope_pa} .pe-field").count()
        checks.check("Edit mode makes the whole Passport Authorization Letter editable (at least its 5 named values)", n_fields >= 5, str(n_fields))

        custom_sentence = "I, the undersigned applicant, hereby authorize Khanna Holidays Pvt. Ltd. to collect my original passport."
        seeded_text = await edit_value_on_preview(page, scope_pa, "applicantSentence", custom_sentence)
        checks.check(
            "applicantSentence opens seeded with the real composed wording",
            applicant_name in seeded_text,
            seeded_text,
        )
        await commit_preview_edit(page, scope_pa)
        persisted = await page.evaluate(
            "(function(){var a=KhannaState.getActiveApplication().authorization.passport; "
            "return (a.editedFields||{}).applicantSentence || null;})()"
        )
        checks.check(
            "The Passport Authorization edit is persisted into KhannaState",
            persisted is not None and custom_sentence in persisted,
            str(persisted),
        )
        n_edited = await page.locator(f"{scope_pa} .pe-field.is-edited").count()
        checks.check("The edited value shows on the page straight away", n_edited == 1, str(n_edited))

        # No overlay is offered for the signature/email block.
        checks.check(
            "No editable value exists for the signature/email block",
            await page.locator(f"{scope_pa} .pe-field[data-pe-key='signature']").count() == 0,
        )

        # Download DOCX regenerates first, so the edit is in the file.
        async with page.expect_download() as dl_info:
            await page.click(f"{scope_pa} [data-dv='download-docx']")
        dl = await dl_info.value
        with open(await dl.path(), "rb") as f:
            passport_doc = Document(io.BytesIO(f.read()))
        passport_text = _full_text(passport_doc)

        checks.check("The edited applicant sentence reached the downloaded Passport Authorization .docx", custom_sentence in passport_text)
        checks.check(
            "Fixed subject line is untouched (project rule 5)",
            "Subject: Authorization Letter for Collection of Passport" in passport_text,
        )
        checks.check("Fixed 'Dear Sir/Madam' salutation is untouched", "Dear Sir/Madam" in passport_text)
        checks.check(
            "The auto-filled signature block (never editable) still carries the applicant's own name",
            applicant_name.split(".")[-1].strip() in passport_text or applicant_name in passport_text,
        )

        # The page itself (PDF) now carries the edit after the regeneration.
        await page.wait_for_timeout(500)
        await wait_preview_ready(page, scope_pa)
        n_edited = await page.locator(f"{scope_pa} .pe-field.is-edited").count()
        checks.check("After regeneration the edit is part of the real PDF page", n_edited == 0, str(n_edited))

        # "Reset" restores the auto-generated wording and clears the override.
        await page.locator(f"{scope_pa} .pe-field[data-pe-key='applicantSentence']").first.click()
        await page.wait_for_timeout(700)
        await page.click(f"{scope_pa} [data-pe='reset']")
        await page.wait_for_timeout(700)
        reset_text = (await page.inner_text(f"{scope_pa} .pe-field.is-active .ql-editor")).strip()
        checks.check(
            "'Reset' restores the real composed applicantSentence",
            applicant_name in reset_text and custom_sentence not in reset_text,
            reset_text,
        )
        cleared_passport = await page.evaluate(
            "(function(){var a=KhannaState.getActiveApplication().authorization.passport; "
            "return (a.editedFields||{}).applicantSentence;})()"
        )
        checks.check("Reset clears the persisted Passport Authorization override back to null", cleared_passport is None, str(cleared_passport))
        await commit_preview_edit(page, scope_pa)
        await set_edit_mode(page, scope_pa, False)
        await wait_preview_ready(page, scope_pa)

        # ===================== Company Authorization =====================
        company_checks = page.locator('[data-auth-checklist="company"] input[type=checkbox]')
        company_count = await company_checks.count()
        for i in range(company_count):
            await company_checks.nth(i).check()
        await pick_company_consulate(page, "Germany")

        checks.check(
            "The old 'Letter wording (editable)' section is gone from the Company Authorization card",
            await page.locator("[data-company-field-editors]").count() == 0
            and "Letter wording" not in await page.inner_text("[data-company-auth-card]"),
        )
        ca_disabled = await page.get_attribute("[data-action='generate-authorization'][data-group='company']", "disabled")
        checks.check("Company Authorization Generate enabled", ca_disabled is None)
        scope_ca = "[data-doc-preview='company']"
        await page.click("[data-action='generate-authorization'][data-group='company']")
        await wait_preview_ready(page, scope_ca)
        await set_edit_mode(page, scope_ca, True)
        n_company = await page.locator(f"{scope_ca} .pe-field").count()
        checks.check(
            f"Company Authorization is fully editable: date + recipient + one editable line per selected applicant ({company_count}) plus the other text blocks",
            n_company >= company_count + 2,
            str(n_company),
        )

        custom_line = "Special-handling applicant line edited right on the preview."
        seeded0 = await edit_value_on_preview(page, scope_ca, "applicant_0", custom_line)
        checks.check("applicant_0 opens seeded with real composed wording", "Passport No." in seeded0, seeded0)
        await commit_preview_edit(page, scope_ca)
        persisted_company = await page.evaluate(
            "(function(){var c=KhannaState.getActiveApplication().authorization.company; "
            "return (c.editedFields||{}).applicant_0 || null;})()"
        )
        checks.check(
            "The Company Authorization edit is persisted into KhannaState",
            persisted_company is not None and custom_line in persisted_company,
            str(persisted_company),
        )

        async with page.expect_download() as dl_info:
            await page.click(f"{scope_ca} [data-dv='download-docx']")
        dl2 = await dl_info.value
        with open(await dl2.path(), "rb") as f:
            company_doc = Document(io.BytesIO(f.read()))
        company_text = _full_text(company_doc)
        checks.check("The edited applicant line reached the downloaded Company Authorization .docx", custom_line in company_text)

        if company_count > 1:
            # Deselecting an applicant shifts every later position — the
            # stale "applicant_0" override must be cleared rather than
            # silently reattached to whichever applicant now sits at
            # position 0 (see onAuthPersonChange's own comment on this).
            await company_checks.nth(0).uncheck()
            await page.wait_for_timeout(300)
            cleared_company = await page.evaluate(
                "(function(){var c=KhannaState.getActiveApplication().authorization.company; "
                "return (c.editedFields||{}).applicant_0 !== undefined ? (c.editedFields||{}).applicant_0 : 'ABSENT';})()"
            )
            checks.check(
                "Deselecting an applicant clears the now-stale positional override",
                cleared_company is None,
                str(cleared_company),
            )
            await company_checks.nth(0).check()
            await page.wait_for_timeout(300)

        checks.check(
            "Fixed boilerplate ('We, Khanna Holidays...') is untouched (project rule 5)",
            "We, Khanna Holidays Pvt. Ltd. have been authorized to collect original Passport" in company_text,
        )
        checks.check("Fixed sign-off block ('Ms. Dhvani Chheda') is untouched", "Ms. Dhvani Chheda" in company_text)

        checks.check("Zero console errors throughout", len(console_errors) == 0, "; ".join(console_errors[:8]))

        await browser.close()
    checks.report_and_exit()


if __name__ == "__main__":
    asyncio.run(main())
