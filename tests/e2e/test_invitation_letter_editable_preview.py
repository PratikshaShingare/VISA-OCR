"""
Editable preview — the Invitation Letter.

Fill the form, press Generate & Preview: the Preview section shows the real
generated page; its Edit button makes the letter's values editable right on
that page (no separate "Letter wording" form). Checks the real flow end to end
(real backend, real LibreOffice PDF, real downloads) — values open seeded with
the generator's own wording, edits persist, Download/Print regenerate with the
edit included, fixed template text is untouched (project rule 5), "Reset all"
restores the auto-generated wording.

Requires: the static file server AND the backend both running (see
run_e2e_suite.py).
"""

import asyncio
import io

from docx import Document
from playwright.async_api import async_playwright

from _common import (
    BASE,
    FAKE_SIGNATURE_PNG,
    Checks,
    commit_preview_edit,
    edit_value_on_preview,
    ensure_fixtures,
    fill_ocr_person,
    fresh_load_and_login,
    pick_inviter_country,
    start_guided_wizard_application,
    wait_preview_ready,
)

checks = Checks()


def _full_text(doc: Document) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


def _has_drawing(doc: Document) -> bool:
    return any(
        r._element.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing") is not None
        for p in doc.paragraphs
        for r in p.runs
    )


async def main():
    ensure_fixtures()
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page(viewport={"width": 1280, "height": 1400})
        console_errors = []
        page.on("console", lambda m: console_errors.append(m.text) if m.type == "error" else None)
        page.on("pageerror", lambda e: console_errors.append(f"pageerror: {e}"))

        await fresh_load_and_login(page, base=BASE)
        await start_guided_wizard_application(page)
        await page.wait_for_timeout(300)

        # ---------- Step 1: applicant ----------
        await page.evaluate("location.hash = '#/new-application/1'")
        await page.wait_for_timeout(300)
        await fill_ocr_person(page, ".step-view[data-step='1']")

        # ---------- Step 2: one traveller — the Invitation Letter needs
        # exactly 2 invitees (applicant + this traveller) ----------
        await page.evaluate("location.hash = '#/new-application/2'")
        await page.wait_for_timeout(300)
        await page.click("[data-action='add-traveller']")
        await page.wait_for_timeout(300)
        await fill_ocr_person(page, ".traveller-panel", "[data-traveller-relation]", "Spouse")
        await page.click("[data-action='close-traveller-panel']")
        await page.wait_for_timeout(200)

        applicant_name = await page.evaluate(
            # The passport OCR now fills the salutation too; the invitees sentence
            # lists names without it, so compare against the plain full name.
            "(function(){var a=KhannaState.getActiveApplication().applicant; return a.fullName;})()"
        )

        # ---------- Step 7: Invitation Letter ----------
        await page.evaluate("location.hash = '#/new-application/7'")
        await page.wait_for_timeout(300)

        inv_checks = page.locator("[data-invitation-invitee-checklist] input[type=checkbox]")
        await inv_checks.nth(0).check()
        await inv_checks.nth(1).check()

        fields = {
            "#invitation-field-inviter-fullName": "Ms. Test Inviter",
            "#invitation-field-inviter-passportNumber": "Z9999999",
            "#invitation-field-inviter-addressLine1": "1 Example Street",
            "#invitation-field-inviter-cityPostcode": "Example City EX1",
            "#invitation-field-inviter-fullAddress": "1 Example Street, Example City EX1",
            "#invitation-field-inviter-cityCountry": "Mumbai, India",
            "#invitation-field-inviter-studyingOrWorking": "Working",
            "#invitation-field-inviter-universityOrCompany": "Example Corp",
            "#invitation-field-inviter-phone": "9999999999",
            "#invitation-field-inviter-email": "inviter@example.com",
            "#invitation-field-top-travelStartDate": "2026-11-10",
            "#invitation-field-top-travelEndDate": "2026-11-20",
            "#invitation-field-top-purpose": "attending a family event",
            "#invitation-field-top-accommodationDetails": "with the inviter",
            "#invitation-field-top-returnDate": "2026-11-20",
            "#invitation-field-top-fundingArrangement": "the inviter",
        }
        await pick_inviter_country(page, "Ireland")
        for selector, value in fields.items():
            await page.fill(selector, value)
        await page.locator("#invitation-field-top-fundingArrangement").blur()
        await page.wait_for_timeout(500)

        checks.check(
            "The old 'Letter wording (editable)' section is gone from the Invitation Letter card",
            await page.locator("[data-invitation-field-editors]").count() == 0
            and "Letter wording" not in await page.inner_text("[data-invitation-card]"),
        )

        await page.set_input_files('[data-signature-file-input="invitation"]', FAKE_SIGNATURE_PNG)
        await page.wait_for_timeout(300)

        gen_disabled = await page.get_attribute("[data-action='generate-invitation']", "disabled")
        checks.check("Invitation Letter Generate enabled", gen_disabled is None)
        scope = "[data-invitation-preview]"
        await page.click("[data-action='generate-invitation']")
        await wait_preview_ready(page, scope)

        edit_btn = page.locator(f"{scope} [data-dv='edit-toggle']")
        checks.check("The preview has an Edit button", await edit_btn.count() == 1 and await edit_btn.is_enabled())
        await edit_btn.click()
        await page.wait_for_timeout(300)
        n_fields = await page.locator(f"{scope} .pe-field").count()
        checks.check("Edit mode makes the whole Invitation Letter editable (at least its 15 named values)", n_fields >= 15, str(n_fields))

        seeded_invitees = await edit_value_on_preview(page, scope, "inviteesDetails", " ", select_all=False)
        checks.check(
            "inviteesDetails opens seeded with the real composed wording",
            applicant_name.title() in seeded_invitees,  # letters use ordinary capitalisation
            seeded_invitees,
        )
        await page.keyboard.press("Backspace")  # undo the probe space typed above
        await commit_preview_edit(page, scope)

        custom_purpose = "Attending a very important once-in-a-lifetime family celebration."
        seeded_purpose = await edit_value_on_preview(page, scope, "purposeStatement", custom_purpose)
        checks.check("purposeStatement opens seeded with the real purpose", "attending a family event" in seeded_purpose, seeded_purpose)
        await commit_preview_edit(page, scope)

        persisted = await page.evaluate(
            "(function(){var i=KhannaState.getActiveApplication().invitation; "
            "return (i.editedFields||{}).purposeStatement || null;})()"
        )
        checks.check(
            "The edit is persisted into KhannaState",
            persisted is not None and custom_purpose in persisted,
            str(persisted),
        )

        async with page.expect_download() as dl_info:
            await page.click(f"{scope} [data-dv='download-docx']")
        dl = await dl_info.value
        with open(await dl.path(), "rb") as f:
            content = f.read()
        doc = Document(io.BytesIO(content))
        full_text = _full_text(doc)

        checks.check("The edited purpose sentence reached the downloaded .docx", custom_purpose in full_text)
        checks.check(
            "The auto-generated fields (email) are still correct alongside the one edited field",
            "inviter@example.com" in full_text,
        )
        checks.check(
            "Fixed subject line is untouched (project rule 5)",
            "Invitation Letter for Visitor" in full_text and "my Parents" in full_text,
        )
        checks.check("Fixed 'Respected Sir/Madam,' salutation is untouched", "Respected Sir/Madam," in full_text)
        checks.check("A real uploaded signature image was embedded (never fabricated, never skipped)", _has_drawing(doc))

        # "Reset all" restores the auto-generated wording and clears the override.
        await wait_preview_ready(page, scope)
        reset_all = page.locator(f"{scope} [data-pe='reset-all']")
        checks.check("'Reset all' counts the hand edits", "1" in (await reset_all.inner_text()), await reset_all.inner_text())
        await reset_all.click()
        await wait_preview_ready(page, scope)
        cleared = await page.evaluate(
            "(function(){var i=KhannaState.getActiveApplication().invitation; "
            "return (i.editedFields||{}).purposeStatement;})()"
        )
        checks.check("Reset all clears the persisted override back to null", cleared is None, str(cleared))
        async with page.expect_download() as dl_info:
            await page.click(f"{scope} [data-dv='download-docx']")
        dl = await dl_info.value
        with open(await dl.path(), "rb") as f:
            doc2 = Document(io.BytesIO(f.read()))
        text2 = _full_text(doc2)
        checks.check(
            "After Reset all the document has the auto-generated purpose again",
            "attending a family event" in text2 and custom_purpose not in text2,
        )

        checks.check("Zero console errors throughout", len(console_errors) == 0, "; ".join(console_errors[:8]))

        await browser.close()
    checks.report_and_exit()


if __name__ == "__main__":
    asyncio.run(main())
