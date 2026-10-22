"""
Invitation Letter — Frontend rework, Phase 4 extension: the Word-like
body editor proven on the Europe Cover Letter, Hotel Blocking, and the
Authorization Letters, extended to the Invitation Letter's own FIXED
15-key field set (date, every inviter address/detail line, the invitees'
passport/travel-dates sentence, the purpose/accommodation/funding
sentences, and the sign-off name/phone/email — see
INVITATION_EDITABLE_FIELDS in js/invitation.js /
INVITATION_EDITABLE_FIELD_KEYS in backend/python/
invitation_letter_engine.py), while every fixed line ("From,", "To,",
"Subject:", "Respected Sir/Madam,", "Thank you!", "Yours Sincerely,") and
the blank signature slot stay completely untouched (project rule 5).

Exercises the REAL flow end to end: the editors are seeded via a real
call to /api/documents/invitation-letter/compose-fields, a real edit is
typed into a real Quill contenteditable surface, the edit is persisted
into KhannaState, and the resulting "Generate & Preview" download is a
real .docx produced by the real backend with that edited wording actually
substituted in — checked by parsing the downloaded file with python-docx.
Also covers a real uploaded signature image still landing correctly
(the signature slot is never offered as an editable field).

Requires: the static file server AND the backend both running (see
run_e2e_suite.py) — KHANNA_E2E_BASE/KHANNA_E2E_BACKEND_BASE override the
default localhost ports.
"""

import asyncio
import io

from docx import Document
from playwright.async_api import async_playwright

from _common import BASE, FAKE_SIGNATURE_PNG, Checks, ensure_fixtures, fill_ocr_person, fresh_load_and_login, generate_and_download_docx, start_guided_wizard_application

checks = Checks()


def _full_text(doc: Document) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


def _has_drawing(doc: Document) -> bool:
    return any(
        r._element.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing") is not None
        for p in doc.paragraphs
        for r in p.runs
    )


async def main():
    ensure_fixtures()
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page(viewport={"width": 1280, "height": 1400})
        console_errors = []
        page.on("console", lambda m: console_errors.append(m.text) if m.type == "error" else None)
        page.on("pageerror", lambda e: console_errors.append(f"pageerror: {e}"))

        await fresh_load_and_login(page, base=BASE)
        await start_guided_wizard_application(page)
        await page.wait_for_timeout(300)

        # ---------- Step 1: applicant ----------
        await page.evaluate("location.hash = '#/new-application/1'")
        await page.wait_for_timeout(300)
        await fill_ocr_person(page, ".step-view[data-step='1']")

        # ---------- Step 2: one traveller — the Invitation Letter needs
        # exactly 2 invitees (applicant + this traveller) ----------
        await page.evaluate("location.hash = '#/new-application/2'")
        await page.wait_for_timeout(300)
        await page.click("[data-action='add-traveller']")
        await page.wait_for_timeout(300)
        await fill_ocr_person(page, ".traveller-panel", "[data-traveller-relation]", "Spouse")
        await page.click("[data-action='close-traveller-panel']")
        await page.wait_for_timeout(200)

        applicant_name = await page.evaluate(
            "(function(){var a=KhannaState.getActiveApplication().applicant; return (a.salutation?a.salutation+'. ':'')+a.fullName;})()"
        )

        # ---------- Step 7: Invitation Letter ----------
        await page.evaluate("location.hash = '#/new-application/7'")
        await page.wait_for_timeout(300)

        inv_checks = page.locator("[data-invitation-invitee-checklist] input[type=checkbox]")
        await inv_checks.nth(0).check()
        await inv_checks.nth(1).check()

        fields = {
            "#invitation-field-inviter-fullName": "Ms. Test Inviter",
            "#invitation-field-inviter-passportNumber": "Z9999999",
            "#invitation-field-inviter-addressLine1": "1 Example Street",
            "#invitation-field-inviter-cityPostcode": "Example City EX1",
            "#invitation-field-inviter-fullAddress": "1 Example Street, Example City EX1",
            "#invitation-field-inviter-country": "TestCountry",
            "#invitation-field-inviter-cityCountry": "Example City, TestCountry",
            "#invitation-field-inviter-studyingOrWorking": "Working",
            "#invitation-field-inviter-universityOrCompany": "Example Corp",
            "#invitation-field-inviter-phone": "9999999999",
            "#invitation-field-inviter-email": "inviter@example.com",
            "#invitation-field-top-travelStartDate": "2026-11-10",
            "#invitation-field-top-travelEndDate": "2026-11-20",
            "#invitation-field-top-purpose": "attending a family event",
            "#invitation-field-top-accommodationDetails": "with the inviter",
            "#invitation-field-top-returnDate": "2026-11-20",
            "#invitation-field-top-fundingArrangement": "the inviter",
        }
        for selector, value in fields.items():
            await page.fill(selector, value)
        await page.locator("#invitation-field-top-fundingArrangement").blur()
        await page.wait_for_timeout(500)

        editors = page.locator("[data-invitation-field-editors] [data-invitation-field-editor]")
        checks.check(
            "All 15 Invitation Letter editable-field editors are mounted",
            await editors.count() == 15,
            str(await editors.count()),
        )

        invitees_editor = page.locator("[data-invitation-field-editor='inviteesDetails'] .ql-editor")
        seeded_text = ""
        for _ in range(40):
            if await invitees_editor.count():
                seeded_text = (await invitees_editor.inner_text()).strip()
                if seeded_text and seeded_text != "—":
                    break
            await page.wait_for_timeout(300)
        checks.check(
            "inviteesDetails field self-seeds with the real composed wording (no extra click)",
            applicant_name in seeded_text,
            seeded_text,
        )

        email_editor = page.locator("[data-invitation-field-editor='email'] .ql-editor")
        email_text = (await email_editor.inner_text()).strip()
        checks.check("email field seeds the real inviter email", "inviter@example.com" in email_text, email_text)

        custom_purpose = "Attending a very important once-in-a-lifetime family celebration."
        purpose_editor = page.locator("[data-invitation-field-editor='purposeStatement'] .ql-editor")
        await purpose_editor.click()
        await page.keyboard.press("Control+A")
        await page.keyboard.type(custom_purpose)
        await purpose_editor.blur()
        await page.wait_for_timeout(300)

        persisted = await page.evaluate(
            "(function(){var i=KhannaState.getActiveApplication().invitation; "
            "return (i.editedFields||{}).purposeStatement || null;})()"
        )
        checks.check(
            "The edit is persisted into KhannaState on blur",
            persisted is not None and custom_purpose in persisted,
            str(persisted),
        )

        await page.set_input_files('[data-signature-file-input="invitation"]', FAKE_SIGNATURE_PNG)
        await page.wait_for_timeout(300)

        gen_disabled = await page.get_attribute("[data-action='generate-invitation']", "disabled")
        checks.check("Invitation Letter Generate enabled", gen_disabled is None)
        dl = await generate_and_download_docx(page, "[data-action='generate-invitation']", "[data-invitation-preview]")
        docx_path = await dl.path()
        with open(docx_path, "rb") as f:
            content = f.read()
        doc = Document(io.BytesIO(content))
        full_text = _full_text(doc)

        checks.check("The edited purpose sentence reached the downloaded .docx", custom_purpose in full_text)
        checks.check(
            "The auto-generated fields (email) are still correct alongside the one edited field",
            "inviter@example.com" in full_text,
        )
        checks.check(
            "Fixed subject line is untouched (project rule 5)",
            "Invitation Letter for Visitor" in full_text and "my Parents" in full_text,
        )
        checks.check("Fixed 'Respected Sir/Madam,' salutation is untouched", "Respected Sir/Madam," in full_text)
        checks.check("A real uploaded signature image was embedded (never fabricated, never skipped)", _has_drawing(doc))

        # "Reset" restores the auto-generated wording and clears the override.
        reset_btn = page.locator("[data-invitation-field-editor='purposeStatement'] [data-body-editor-action='reset']")
        await reset_btn.click()
        await page.wait_for_timeout(400)
        reset_text = (await page.locator("[data-invitation-field-editor='purposeStatement'] .ql-editor").inner_text()).strip()
        checks.check(
            "'Reset to auto-generated text' restores the real composed purposeStatement",
            "attending a family event" in reset_text and custom_purpose not in reset_text,
            reset_text,
        )
        cleared = await page.evaluate(
            "(function(){var i=KhannaState.getActiveApplication().invitation; "
            "return (i.editedFields||{}).purposeStatement;})()"
        )
        checks.check("Reset clears the persisted override back to null", cleared is None, str(cleared))

        checks.check("Zero console errors throughout", len(console_errors) == 0, "; ".join(console_errors[:8]))

        await browser.close()
    checks.report_and_exit()


if __name__ == "__main__":
    asyncio.run(main())
