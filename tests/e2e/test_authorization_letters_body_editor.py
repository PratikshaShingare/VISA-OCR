"""
Authorization Letters — Frontend rework, Phase 4 extension: the Word-like
body editor proven on the Europe Cover Letter and Hotel Blocking, extended
to both the Passport Authorization Letter (a FIXED 5-key field set — date,
recipient centre, recipient address, applicant sentence, collector
sentence — see PASSPORT_EDITABLE_FIELDS in js/authorization.js /
compose_passport_editable_fields in backend/python/
authorization_letter_engine.py) and the Company Authorization Letter (a
DYNAMIC field set — date, recipient, plus one "applicant_<index>" editor
per currently-selected applicant — see companyFieldDefs()/
mountCompanyFieldEditors() in js/authorization.js /
compose_company_editable_fields in the same engine file), while every
fixed heading, boilerplate sentence and signature/email block stays
completely untouched (project rule 5).

Exercises the REAL flow end to end: each compact field editor is seeded
via a real call to the matching /compose-fields endpoint (never
placeholder/fake text), a real edit is typed into a real Quill
contenteditable surface, the edit is persisted into KhannaState, and the
resulting "Generate & Preview" download is a real .docx produced by the
real backend with that edited wording actually substituted in — checked
by parsing the downloaded file with python-docx. Also covers the Passport
Authorization Letter's signature/email block staying entirely untouched
(never offered as an editable field — see authorization_letter_engine.
_fill_passport_single's own comment for why), and the Company
Authorization Letter's per-applicant editors correctly REMOUNTING (and
clearing stale positional overrides) when the selected-applicant list
changes.

Requires: the static file server AND the backend both running (see
run_e2e_suite.py) — KHANNA_E2E_BASE/KHANNA_E2E_BACKEND_BASE override the
default localhost ports.
"""

import asyncio
import io

from docx import Document
from playwright.async_api import async_playwright

from _common import BASE, Checks, fill_ocr_person, fresh_load_and_login, generate_and_download_docx, start_guided_wizard_application

checks = Checks()


def _full_text(doc: Document) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


async def main():
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page(viewport={"width": 1280, "height": 1400})
        console_errors = []
        page.on("console", lambda m: console_errors.append(m.text) if m.type == "error" else None)
        page.on("pageerror", lambda e: console_errors.append(f"pageerror: {e}"))

        await fresh_load_and_login(page, base=BASE)
        await start_guided_wizard_application(page)
        await page.wait_for_timeout(300)

        # ---------- Step 1: applicant ----------
        await page.evaluate("location.hash = '#/new-application/1'")
        await page.wait_for_timeout(300)
        await fill_ocr_person(page, ".step-view[data-step='1']")

        # ---------- Step 2: one traveller (Spouse), so Company has 2
        # selectable applicants to exercise the dynamic field set with ----------
        await page.evaluate("location.hash = '#/new-application/2'")
        await page.wait_for_timeout(300)
        await page.click("[data-action='add-traveller']")
        await page.wait_for_timeout(300)
        await fill_ocr_person(page, ".traveller-panel", "[data-traveller-relation]", "Spouse")
        await page.click("[data-action='close-traveller-panel']")
        await page.wait_for_timeout(200)

        applicant_name = await page.evaluate(
            "(function(){var a=KhannaState.getActiveApplication().applicant; return (a.salutation?a.salutation+'. ':'')+a.fullName;})()"
        )

        # ---------- Step 6: Authorization Letters ----------
        await page.evaluate("location.hash = '#/new-application/6'")
        await page.wait_for_timeout(300)

        # ===================== Passport Authorization =====================
        passport_checks = page.locator('[data-auth-checklist="passport"] input[type=checkbox]')
        await passport_checks.nth(0).check()  # applicant only — exercises the single-traveller template
        await page.fill("#auth-passport-recipientCentreName", "Germany Visa Application Centre")
        await page.fill("#auth-passport-recipientAddress", "Andheri, Mumbai")
        await page.fill("#auth-passport-collectorName", "Mr. Suresh Patil")
        await page.locator("#auth-passport-collectorName").blur()
        await page.wait_for_timeout(300)

        passport_editors = page.locator("[data-passport-field-editors] [data-passport-field-editor]")
        checks.check(
            "All 5 Passport Authorization editable-field editors are mounted",
            await passport_editors.count() == 5,
            str(await passport_editors.count()),
        )

        applicant_sentence_editor = page.locator("[data-passport-field-editor='applicantSentence'] .ql-editor")
        seeded_text = ""
        for _ in range(40):
            if await applicant_sentence_editor.count():
                seeded_text = (await applicant_sentence_editor.inner_text()).strip()
                if seeded_text and seeded_text != "—":
                    break
            await page.wait_for_timeout(300)
        checks.check(
            "applicantSentence field self-seeds with the real composed wording (no extra click)",
            applicant_name in seeded_text,
            seeded_text,
        )

        centre_editor = page.locator("[data-passport-field-editor='recipientCentre'] .ql-editor")
        centre_text = (await centre_editor.inner_text()).strip()
        checks.check(
            "recipientCentre field seeds the real centre name",
            "Germany Visa Application Centre" in centre_text,
            centre_text,
        )

        # No editor is offered for the signature/email block — see
        # authorization_letter_engine's own comment on why that stays
        # auto-only (a real <w:br/> + mailto: hyperlink it would destroy).
        signature_editor_count = await page.locator("[data-passport-field-editor='signature']").count()
        checks.check("No editable-field editor exists for the signature/email block", signature_editor_count == 0)

        custom_sentence = "I, the undersigned applicant, hereby authorize Khanna Holidays Pvt. Ltd. to collect my original passport."
        await applicant_sentence_editor.click()
        await page.keyboard.press("Control+A")
        await page.keyboard.type(custom_sentence)
        await applicant_sentence_editor.blur()
        await page.wait_for_timeout(300)

        persisted = await page.evaluate(
            "(function(){var a=KhannaState.getActiveApplication().authorization.passport; "
            "return (a.editedFields||{}).applicantSentence || null;})()"
        )
        checks.check(
            "The Passport Authorization edit is persisted into KhannaState on blur",
            persisted is not None and custom_sentence in persisted,
            str(persisted),
        )

        pa_disabled = await page.get_attribute("[data-action='generate-authorization'][data-group='passport']", "disabled")
        checks.check("Passport Authorization Generate enabled", pa_disabled is None)
        dl = await generate_and_download_docx(
            page, "[data-action='generate-authorization'][data-group='passport']", "[data-doc-preview='passport']"
        )
        docx_path = await dl.path()
        with open(docx_path, "rb") as f:
            passport_doc = Document(io.BytesIO(f.read()))
        passport_text = _full_text(passport_doc)

        checks.check("The edited applicant sentence reached the downloaded Passport Authorization .docx", custom_sentence in passport_text)
        checks.check(
            "Fixed subject line is untouched (project rule 5)",
            "Subject: Authorization Letter for Collection of Passport" in passport_text,
        )
        checks.check("Fixed 'Dear Sir/Madam' salutation is untouched", "Dear Sir/Madam" in passport_text)
        checks.check(
            "The auto-filled signature block (never editable) still carries the applicant's own name",
            applicant_name.split(".")[-1].strip() in passport_text or applicant_name in passport_text,
        )

        # "Reset" restores the auto-generated wording and clears the override.
        reset_btn = page.locator("[data-passport-field-editor='applicantSentence'] [data-body-editor-action='reset']")
        await reset_btn.click()
        await page.wait_for_timeout(400)
        reset_text = (await page.locator("[data-passport-field-editor='applicantSentence'] .ql-editor").inner_text()).strip()
        checks.check(
            "'Reset to auto-generated text' restores the real composed applicantSentence",
            applicant_name in reset_text and custom_sentence not in reset_text,
            reset_text,
        )
        cleared_passport = await page.evaluate(
            "(function(){var a=KhannaState.getActiveApplication().authorization.passport; "
            "return (a.editedFields||{}).applicantSentence;})()"
        )
        checks.check("Reset clears the persisted Passport Authorization override back to null", cleared_passport is None, str(cleared_passport))

        # ===================== Company Authorization =====================
        company_checks = page.locator('[data-auth-checklist="company"] input[type=checkbox]')
        company_count = await company_checks.count()
        for i in range(company_count):
            await company_checks.nth(i).check()
        await page.fill("#auth-company-recipientText", "The Consulate General of Germany, Mumbai")
        await page.locator("#auth-company-recipientText").blur()
        await page.wait_for_timeout(300)

        company_editors = page.locator("[data-company-field-editors] [data-company-field-editor]")
        checks.check(
            f"Company Authorization has date + recipient + one editor per selected applicant ({company_count})",
            await company_editors.count() == company_count + 2,
            str(await company_editors.count()),
        )

        applicant0_editor = page.locator("[data-company-field-editor='applicant_0'] .ql-editor")
        seeded0 = ""
        for _ in range(40):
            if await applicant0_editor.count():
                seeded0 = (await applicant0_editor.inner_text()).strip()
                if seeded0 and seeded0 != "—":
                    break
            await page.wait_for_timeout(300)
        checks.check("applicant_0 field self-seeds with real composed wording", "Passport No." in seeded0, seeded0)

        custom_line = "Special-handling applicant line edited via the Word-like editor."
        await applicant0_editor.click()
        await page.keyboard.press("Control+A")
        await page.keyboard.type(custom_line)
        await applicant0_editor.blur()
        await page.wait_for_timeout(300)

        persisted_company = await page.evaluate(
            "(function(){var c=KhannaState.getActiveApplication().authorization.company; "
            "return (c.editedFields||{}).applicant_0 || null;})()"
        )
        checks.check(
            "The Company Authorization edit is persisted into KhannaState on blur",
            persisted_company is not None and custom_line in persisted_company,
            str(persisted_company),
        )

        if company_count > 1:
            # Deselecting an applicant shifts every later position — the
            # stale "applicant_0" override must be cleared rather than
            # silently reattached to whichever applicant now sits at
            # position 0 (see onAuthPersonChange's own comment on this).
            await company_checks.nth(0).uncheck()
            await page.wait_for_timeout(300)
            cleared_company = await page.evaluate(
                "(function(){var c=KhannaState.getActiveApplication().authorization.company; "
                "return (c.editedFields||{}).applicant_0 !== undefined ? (c.editedFields||{}).applicant_0 : 'ABSENT';})()"
            )
            checks.check(
                "Deselecting an applicant clears the now-stale positional override",
                cleared_company is None,
                str(cleared_company),
            )
            new_editor_count = await page.locator("[data-company-field-editors] [data-company-field-editor]").count()
            checks.check(
                "The editable-field editor SET remounts to match the new selection count",
                new_editor_count == company_count + 1,
                str(new_editor_count),
            )
            # Re-select so the letter still has every applicant for the
            # download below (an honest "at least 1" letter, not a
            # regression side effect of the check above).
            await company_checks.nth(0).check()
            await page.wait_for_timeout(300)

        ca_disabled = await page.get_attribute("[data-action='generate-authorization'][data-group='company']", "disabled")
        checks.check("Company Authorization Generate enabled", ca_disabled is None)
        dl2 = await generate_and_download_docx(
            page, "[data-action='generate-authorization'][data-group='company']", "[data-doc-preview='company']"
        )
        docx_path2 = await dl2.path()
        with open(docx_path2, "rb") as f:
            company_doc = Document(io.BytesIO(f.read()))
        company_text = _full_text(company_doc)

        checks.check(
            "Fixed boilerplate ('We, Khanna Holidays...') is untouched (project rule 5)",
            "We, Khanna Holidays Pvt. Ltd. have been authorized to collect original Passport" in company_text,
        )
        checks.check("Fixed sign-off block ('Ms. Dhvani Chheda') is untouched", "Ms. Dhvani Chheda" in company_text)

        checks.check("Zero console errors throughout", len(console_errors) == 0, "; ".join(console_errors[:8]))

        await browser.close()
    checks.report_and_exit()


if __name__ == "__main__":
    asyncio.run(main())
