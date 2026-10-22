"""
Editable preview — the Checklist Letter.

Fill the form, press Generate & Preview: the Preview section shows the real
generated page; its Edit button makes the letter's values editable right on
that page (no separate "Letter wording" form). Checks the real flow end to end
(real backend, real LibreOffice PDF, real downloads) — values open seeded with
the generator's own wording, edits persist, Download/Print regenerate with the
edit included, fixed template text is untouched (project rule 5), "Reset all"
restores the auto-generated wording.

Requires: the static file server AND the backend both running (see
run_e2e_suite.py).
"""

import asyncio
import io

from docx import Document
from playwright.async_api import async_playwright

from _common import BASE, Checks, commit_preview_edit, edit_value_on_preview, fresh_load_and_login, wait_preview_ready

checks = Checks()


def _full_text(doc: Document) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


async def main():
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page(viewport={"width": 1280, "height": 1400})
        console_errors = []
        page.on("console", lambda m: console_errors.append(m.text) if m.type == "error" else None)
        page.on("pageerror", lambda e: console_errors.append(f"pageerror: {e}"))

        await fresh_load_and_login(page, base=BASE)

        await page.evaluate("location.hash = '#/checklist-letter'")
        await page.wait_for_timeout(500)

        # The editable-fields section shares the Generate button's own
        # readiness gate (country filled) — the ONLY field this document
        # type genuinely requires, since processing time/visa fees are
        # legitimately optional (an honest "_____" placeholder otherwise —
        # project rule 9/10). Exactly like every other document type's own
        # compact editors (see hotels.js's own mountHotelFieldEditors header
        # comment), a field seeds ONCE, the moment its gate flips true, and
        # does not keep re-syncing to a later keystroke into the fields
        # above it — so this fills every field THAT gate does not already
        # cover first, then fills country LAST, matching the realistic
        # order a staff member filling in supporting details before the
        # one truly required field would produce.
        r = "[data-checklist-letter-root]"
        await page.select_option(f"{r} [data-field='visaType']", "Tourist")
        await page.select_option(f"{r} [data-field='employmentStatus']", "Employed")
        await page.fill(f"{r} [data-field='visaFees']", "5,000")
        await page.locator(f"{r} [data-field='visaFees']").blur()
        await page.fill(f"{r} [data-field='country']", "Germany")
        await page.locator(f"{r} [data-field='country']").blur()
        await page.wait_for_timeout(800)  # let the async suggested-documents lookup settle first

        # ---------- The checkbox-driven document SELECTION mechanism is
        # untouched by this phase's own work — confirm it still behaves
        # exactly as before: at least one document got auto-suggested and
        # is checked. ----------
        doc_checkboxes = page.locator(f"{r} [data-rdl-document]")
        checked_count = 0
        for i in range(await doc_checkboxes.count()):
            if await doc_checkboxes.nth(i).is_checked():
                checked_count += 1
        checks.check("At least one document is auto-suggested/checked (selection UI unchanged)", checked_count > 0, str(checked_count))

        checks.check(
            "The old 'Letter wording (editable)' section is gone from the Checklist Letter page",
            await page.locator(f"{r} [data-rdl-field-editors]").count() == 0 and "Letter wording" not in await page.inner_text(r),
        )

        # ---------- Generate & Preview, then edit on the page ----------
        gen_disabled = await page.get_attribute(f"{r} [data-action='generate-rdl']", "disabled")
        checks.check("Checklist Letter Generate enabled", gen_disabled is None)
        scope = f"{r} [data-rdl-preview]"
        await page.click(f"{r} [data-action='generate-rdl']")
        await wait_preview_ready(page, scope)
        edit_btn = page.locator(f"{scope} [data-dv='edit-toggle']")
        checks.check("The preview has an Edit button", await edit_btn.count() == 1 and await edit_btn.is_enabled())
        await edit_btn.click()
        await page.wait_for_timeout(300)
        n_fields = await page.locator(f"{scope} .pe-field").count()
        keys = await page.locator(f"{scope} .pe-field").evaluate_all("els=>els.map(e=>e.getAttribute('data-pe-key'))")
        checks.check("Edit mode makes the whole Checklist Letter editable (the 6 values plus every other text block)", n_fields >= 6 and all(k in keys for k in ("checklistItems", "greetingSentence", "visaFees")), str(n_fields))

        # The whole checklist range is one editable block.
        custom_checklist = "A completely staff-rewritten checklist line."
        checklist_seed = await edit_value_on_preview(page, scope, "checklistItems", custom_checklist)
        checks.check("checklistItems opens seeded with the real auto-generated document labels", "Passport Copy" in checklist_seed, checklist_seed)
        await commit_preview_edit(page, scope)

        custom_greeting = "A warm, staff-edited greeting for this enquiry."
        greeting_seed = await edit_value_on_preview(page, scope, "greetingSentence", custom_greeting)
        checks.check("greetingSentence opens seeded with the real composed wording", "Thank you for your inquiry" in greeting_seed and "Germany" in greeting_seed, greeting_seed)
        await commit_preview_edit(page, scope)

        fees_seed = await edit_value_on_preview(page, scope, "visaFees", " ", select_all=False)
        checks.check("visaFees opens seeded with the real typed-in figure", "5,000" in fees_seed, fees_seed)
        await page.keyboard.press("Backspace")
        await commit_preview_edit(page, scope)

        async with page.expect_download() as dl_info:
            await page.click(f"{scope} [data-dv='download-docx']")
        dl = await dl_info.value
        with open(await dl.path(), "rb") as f:
            content = f.read()
        doc = Document(io.BytesIO(content))
        full_text = _full_text(doc)

        checks.check("The edited checklist line reached the downloaded .docx", custom_checklist in full_text)
        checks.check("The edited greeting sentence reached the downloaded .docx", custom_greeting in full_text)
        checks.check("The un-edited auto-generated documents are gone (replaced, not appended)", "Passport Copy" not in full_text)
        checks.check(
            "The un-edited fields (country/visa type/fees) are still correct alongside the two edited fields",
            "Country: Germany" in full_text and "5,000" in full_text,
        )
        checks.check("Fixed closing wording untouched (project rule 5)", "Warm Regards," in full_text)

        # ---------- Reset all restores the auto-generated wording ----------
        await wait_preview_ready(page, scope)
        reset_all = page.locator(f"{scope} [data-pe='reset-all']")
        label = await reset_all.inner_text()
        checks.check("'Reset all' counts exactly the two hand edits (a touched-but-unchanged value is not an edit)", "(2)" in label, label)
        await reset_all.click()
        await wait_preview_ready(page, scope)
        async with page.expect_download() as dl_info:
            await page.click(f"{scope} [data-dv='download-docx']")
        dl = await dl_info.value
        with open(await dl.path(), "rb") as f:
            text2 = _full_text(Document(io.BytesIO(f.read())))
        checks.check(
            "After Reset all the document has the real composed wording again",
            "Thank you for your inquiry" in text2 and custom_greeting not in text2 and "Passport Copy" in text2,
        )

        checks.check("Zero console errors throughout", len(console_errors) == 0, "; ".join(console_errors[:8]))

        await browser.close()
    checks.report_and_exit()


if __name__ == "__main__":
    asyncio.run(main())
