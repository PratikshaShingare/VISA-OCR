"""
Khanna Travels & Holidays — Cover Letter document generators
================================================================

Generates the three real, region-specific Khanna Travels & Holidays
Tourist Visa Cover Letters, each from its own real, unmodified reference
template (project rule 5 — Document Template Rule):

    reference-cover-letter-formats/Europe/Europe_covering_letter_template_clean.docx
    reference-cover-letter-formats/Japan/Japan_covering_letter_template_clean.docx
    reference-cover-letter-formats/Singapore/Singapore_covering_letter_template_clean.docx

Unlike the Passport Authorization templates (Phase 8), these three ARE
genuine blank templates — every dynamic paragraph in all three is a
single run holding bracketed `[Token]` placeholders, confirmed by direct
inspection before writing any fill logic here. Token substitution is
still done paragraph-by-paragraph (`docx_utils.apply_token_map_to_paragraph`)
rather than as one document-wide sweep, because the SAME bracket token
(e.g. `[Passport Number]`, `[Job Title]`) legitimately needs two
different real values in two different places — an applicant's vs. a
companion's — and a blind find-and-replace-everywhere would silently
write the same value into both. A few Europe sentences additionally have
the identical token appearing *twice within one paragraph* with two
different values (e.g. "[Number of Nights]" for a primary destination
and an optional second-country leg) — those specific paragraphs are
composed directly as plain Python strings instead of token-substituted,
for the same reason.

Two deliberate, documented judgement calls on real template quirks
(neither left in the generated output — project rule 9, no faked/broken
content reaches a client):
  - Japan's companion sentence carries a literal, human-authored
    instruction, `[Add additional family member details here if
    applicable].` — stripped outright rather than ever shown to a client,
    consistent with this letter's own real person cap (see below).
  - Singapore's file has a second, broken, mislabeled signature block
    directly duplicating the first ("Passanger 1 Name" — misspelled and
    not even bracketed — "P1's Email ID") — deleted outright as leftover
    template cruft, the same judgement call already made for the Company
    Authorization Letter's flawed second applicant line in Phase 8.

Real-person-count limits, each directly justified by what that
template's own FIXED sentences actually say — never an invented sentence
structure beyond what the reference material supports (project rule 9):
  - Europe: exactly applicant + 1 companion. Every dynamic sentence is
    written for exactly one named companion ("[Relation/Family]"
    singular, "We both are...", "My [Relation]...").
  - Japan: exactly applicant + 1 companion, for the same reason — its
    own template literally asks a human to hand-add more, which this
    engine does not attempt to do automatically.
  - Singapore: applicant alone, or applicant + any number of companions —
    its passenger table is a genuinely repeatable row (Sr no./Name/
    Passport/Relation/Occupation), cloned the same way the Company
    Authorization Letter's applicant line is cloned (Phase 8), and its
    intro sentence is composed (not just token-filled) so "along with my
    Family" is dropped cleanly when travelling solo.
"""

from __future__ import annotations

import io
import os
from typing import Optional

from docx import Document
from docx.table import Table

import docx_utils
import letter_shared

_THIS_DIR = os.path.dirname(os.path.abspath(__file__))
_REFERENCE_ROOT = os.path.normpath(os.path.join(_THIS_DIR, "..", "..", "reference-cover-letter-formats"))

EUROPE_TEMPLATE_PATH = os.path.join(_REFERENCE_ROOT, "Europe", "Europe_covering_letter_template_clean.docx")
JAPAN_TEMPLATE_PATH = os.path.join(_REFERENCE_ROOT, "Japan", "Japan_covering_letter_template_clean.docx")
SINGAPORE_TEMPLATE_PATH = os.path.join(_REFERENCE_ROOT, "Singapore", "Singapore_covering_letter_template_clean.docx")


class CoverLetterError(docx_utils.DocxBuildError):
    """Raised whenever a cover letter genuinely cannot be generated —
    never swallowed to produce a fake/empty document (project rule 9)."""


def _load_template(path: str, label: str) -> Document:
    if not os.path.isfile(path):
        raise CoverLetterError(
            f"{label} reference template not found at {path}. This file is part "
            "of the project's read-only reference material (reference-cover-letter-"
            "formats/) and must already exist on this machine — it is never created "
            "by this app."
        )
    return Document(path)  # only ever read, never written back to


def _relation_word_or_fallback(relation: Optional[str], sex: Optional[str]) -> str:
    return letter_shared.relation_word(relation, sex) or "family member"


def _relation_phrase_or_fallback(relation: Optional[str], sex: Optional[str]) -> str:
    return letter_shared.relation_phrase(relation, sex) or "my family member"


def _phone_or_dash(phone: Optional[str]) -> str:
    digits = letter_shared.local_phone_digits(phone)
    return f"+91 {digits}" if digits else "—"


# ---------------------------------------------------------------------------
# Europe
# ---------------------------------------------------------------------------


def generate_europe_cover_letter_docx(
    applicant: dict,
    companion: dict,
    fields: dict,
    date_override: Optional[str] = None,
) -> bytes:
    """`fields` keys used: recipientText, destinationCountry,
    travelStartDate, travelEndDate, cityCountryOfResidence,
    numberOfNights, applicantEmploymentStatus, applicantJobTitle,
    applicantEmployerName, companionJobTitle, companionEmployerName,
    companionEmploymentStartYear, nextCountry, nextTravelStartDate,
    nextTravelEndDate, nextNumberOfNights, fundingArrangement."""
    if not applicant or not companion:
        raise ValueError(
            "The Europe Cover Letter is written for the applicant plus exactly one "
            "companion (matching the reference template's own fixed wording) — "
            "please select both."
        )

    document = _load_template(EUROPE_TEMPLATE_PATH, "Europe Cover Letter")
    paragraphs = document.paragraphs

    docx_utils.apply_token_map_to_paragraph(paragraphs[0], {"[Date]": letter_shared.format_letter_date(date_override)})
    docx_utils.set_paragraph_text(paragraphs[3], (fields.get("recipientText") or "").strip() or "—")
    docx_utils.apply_token_map_to_paragraph(
        paragraphs[5], {"[Destination Country]": fields.get("destinationCountry") or "—"}
    )

    relation_with_name = f"{_relation_word_or_fallback(companion.get('relationToApplicant'), companion.get('sex'))}, {companion.get('fullName') or '—'}"
    relation_bare = _relation_word_or_fallback(companion.get("relationToApplicant"), companion.get("sex"))

    para8 = (
        f"I, {letter_shared.display_name(applicant)} (holding Indian Passport No.: "
        f"{applicant.get('passportNumber') or '—'} issued at {applicant.get('placeOfIssue') or '—'} on "
        f"{applicant.get('passportIssueDate') or '—'}) would like to visit your admired country along with "
        f"my {relation_with_name} from {fields.get('travelStartDate') or '—'} to "
        f"{fields.get('travelEndDate') or '—'} for the purpose of Tourism."
    )
    docx_utils.set_paragraph_text(paragraphs[8], para8)

    para9 = (
        f"I and my {relation_with_name} reside in {fields.get('cityCountryOfResidence') or '—'} along with our "
        "parents and few close relatives. We both are financially independent and Employed – "
        f"{fields.get('applicantEmploymentStatus') or '—'}. I am a {fields.get('applicantJobTitle') or '—'} at "
        f"{fields.get('applicantEmployerName') or '—'}. While my {relation_bare} is working for "
        f"{fields.get('companionEmployerName') or '—'} since {fields.get('companionEmploymentStartYear') or '—'}, "
        f"currently as {fields.get('companionJobTitle') or '—'}. Kindly find our relevant documents for your "
        "reference."
    )
    docx_utils.set_paragraph_text(paragraphs[9], para9)

    destination = fields.get("destinationCountry") or "—"
    next_country = (fields.get("nextCountry") or "").strip()
    destination_or_countries = f"{destination} and {next_country}" if next_country else destination
    next_leg = (
        f" and then to {next_country} for {fields.get('nextNumberOfNights') or '—'} from "
        f"{fields.get('nextTravelStartDate') or '—'} to {fields.get('nextTravelEndDate') or '—'}"
        if next_country
        else ""
    )
    para10 = (
        f"My {relation_with_name} and I have planned to visit {destination_or_countries} for the purpose of "
        "Tourism. We look forward to visiting the country, exploring the beautiful sights, witnessing its rich "
        f"culture and spending leisure time together. We would be visiting {destination} for "
        f"{fields.get('numberOfNights') or '—'} from {fields.get('travelStartDate') or '—'} to "
        f"{fields.get('travelEndDate') or '—'}{next_leg}. Kindly find our relevant documents for your perusal. "
        "We will return back to India after the end of our visit as we need to rejoin and look after our work "
        "and personal commitments."
    )
    docx_utils.set_paragraph_text(paragraphs[10], para10)

    docx_utils.apply_token_map_to_paragraph(
        paragraphs[11], {"[Funding Arrangement]": fields.get("fundingArrangement") or "—"}
    )

    docx_utils.apply_token_map_to_paragraph(paragraphs[17], {"[Applicant Full Name]": letter_shared.display_name(applicant)})
    docx_utils.apply_token_map_to_paragraph(paragraphs[18], {"[Phone Number]": _phone_or_dash(applicant.get("phone"))})
    docx_utils.apply_token_map_to_paragraph(paragraphs[19], {"[Email Address]": (applicant.get("email") or "").strip() or "—"})

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def generate_europe_cover_letter_pdf(applicant: dict, companion: dict, fields: dict, date_override: Optional[str] = None) -> bytes:
    docx_bytes = generate_europe_cover_letter_docx(applicant, companion, fields, date_override)
    try:
        return docx_utils.convert_docx_bytes_to_pdf(docx_bytes, base_name="europe_cover_letter")
    except docx_utils.DocxBuildError as e:
        raise CoverLetterError(str(e))


# ---------------------------------------------------------------------------
# Japan
# ---------------------------------------------------------------------------


def _fill_japan_hotel_table(document: Document, hotels: list[dict]) -> None:
    table = document.tables[0]
    header = [c.text.strip() for c in table.rows[0].cells]
    if header != ["NAME", "DATE", "CONTACT NO."]:
        raise CoverLetterError(
            "Japan Cover Letter template structure has changed (hotel table header "
            f"was {header!r}) — cannot safely fill it."
        )
    if len(table.rows) < 3:
        raise CoverLetterError(
            "Japan Cover Letter template structure has changed (expected 2 example "
            "hotel rows) — cannot safely fill it."
        )
    template_row_el = table.rows[1]._tr
    extra_row_el = table.rows[2]._tr
    extra_row_el.getparent().remove(extra_row_el)

    if not hotels:
        # No hotel info given at all — remove the now-templateless row too
        # rather than leave literal "[Hotel Name 1]" placeholder text in a
        # real client-facing letter.
        template_row_el.getparent().remove(template_row_el)
        return

    clone_els = docx_utils.clone_and_insert_after(template_row_el, len(hotels) - 1)
    row_els = [template_row_el] + clone_els
    for row_el, hotel in zip(row_els, hotels):
        cells = _cells_of_row_element(table, row_el)
        docx_utils.set_cell_text(cells[0], hotel.get("name") or "—")
        stay = f"{hotel.get('checkIn') or '—'} - {hotel.get('checkOut') or '—'}"
        docx_utils.set_cell_text(cells[1], stay)
        docx_utils.set_cell_text(cells[2], hotel.get("contactNo") or "—")


def _cells_of_row_element(table: Table, row_el):
    """Returns the docx Cell objects belonging to a specific `<w:tr>`
    element, even one just cloned and not yet reflected in `table.rows`
    (python-docx's Table wraps the live XML, so a freshly inserted row is
    already valid to address this way — the same technique the Hotel
    Voucher engine's multi-hotel cloning already relies on)."""
    from docx.table import _Row

    row = _Row(row_el, table)
    return row.cells


def generate_japan_cover_letter_docx(
    applicant: dict,
    companion: dict,
    fields: dict,
    hotels: list[dict],
    date_override: Optional[str] = None,
) -> bytes:
    """`fields` keys used: destinationCountry, travelStartDate,
    travelEndDate, applicantEmployerOccupation, companionOccupation,
    fundingArrangement."""
    if not applicant or not companion:
        raise ValueError(
            "The Japan Cover Letter is written for the applicant plus exactly one "
            "companion (matching the reference template's own fixed wording) — "
            "please select both."
        )

    document = _load_template(JAPAN_TEMPLATE_PATH, "Japan Cover Letter")
    paragraphs = document.paragraphs

    docx_utils.apply_token_map_to_paragraph(paragraphs[1], {"[Date]": letter_shared.format_letter_date(date_override)})
    docx_utils.apply_token_map_to_paragraph(
        paragraphs[8], {"[Destination Country]": fields.get("destinationCountry") or "—"}
    )
    docx_utils.apply_token_map_to_paragraph(
        paragraphs[10],
        {
            "[Passenger 1 Name]": letter_shared.display_name(applicant),
            "[Passport Number]": applicant.get("passportNumber") or "—",
            "[Place of Issue]": applicant.get("placeOfIssue") or "—",
            "[Passport Issue Date]": applicant.get("passportIssueDate") or "—",
            "[Travel Start Date]": fields.get("travelStartDate") or "—",
            "[Travel End Date]": fields.get("travelEndDate") or "—",
        },
    )
    docx_utils.apply_token_map_to_paragraph(
        paragraphs[11], {"[Employer / Occupation]": fields.get("applicantEmployerOccupation") or "—"}
    )
    docx_utils.apply_token_map_to_paragraph(
        paragraphs[12],
        {
            "[Relation]": _relation_word_or_fallback(companion.get("relationToApplicant"), companion.get("sex")).capitalize(),
            "[Passenger 2 Name]": letter_shared.display_name(companion),
            "[Passport Number]": companion.get("passportNumber") or "—",
            "[Place of Issue]": companion.get("placeOfIssue") or "—",
            "[Passport Issue Date]": companion.get("passportIssueDate") or "—",
            "[Occupation]": fields.get("companionOccupation") or "—",
            # The template's own human-authored instruction — never shown
            # to a client (see module docstring).
            " [Add additional family member details here if applicable].": "",
        },
    )
    docx_utils.apply_token_map_to_paragraph(
        paragraphs[15], {"[Sponsor Name / Funding Arrangement]": fields.get("fundingArrangement") or "—"}
    )
    docx_utils.apply_token_map_to_paragraph(paragraphs[21], {"[Passenger 1 Name]": letter_shared.display_name(applicant)})
    docx_utils.apply_token_map_to_paragraph(paragraphs[22], {"[Phone Number]": _phone_or_dash(applicant.get("phone"))})
    docx_utils.apply_token_map_to_paragraph(
        paragraphs[23], {"[Email Address]": (applicant.get("email") or "").strip() or "—"}
    )

    _fill_japan_hotel_table(document, hotels)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def generate_japan_cover_letter_pdf(
    applicant: dict, companion: dict, fields: dict, hotels: list[dict], date_override: Optional[str] = None
) -> bytes:
    docx_bytes = generate_japan_cover_letter_docx(applicant, companion, fields, hotels, date_override)
    try:
        return docx_utils.convert_docx_bytes_to_pdf(docx_bytes, base_name="japan_cover_letter")
    except docx_utils.DocxBuildError as e:
        raise CoverLetterError(str(e))


# ---------------------------------------------------------------------------
# Singapore
# ---------------------------------------------------------------------------


def _fill_singapore_passenger_table(document: Document, applicant: dict, companions: list[dict], occupations: dict) -> None:
    table = document.tables[0]
    header = [c.text.strip() for c in table.rows[0].cells]
    if header != ["Sr no.", "Passengers Name", "Passport No", "Relation", "Occupation"]:
        raise CoverLetterError(
            f"Singapore Cover Letter template structure has changed (passenger table header was {header!r}) "
            "— cannot safely fill it."
        )
    if len(table.rows) < 3:
        raise CoverLetterError(
            "Singapore Cover Letter template structure has changed (expected applicant + one example companion "
            "row) — cannot safely fill it."
        )

    applicant_row_el = table.rows[1]._tr
    companion_template_row_el = table.rows[2]._tr

    applicant_cells = _cells_of_row_element(table, applicant_row_el)
    docx_utils.set_cell_text(applicant_cells[0], "1")
    docx_utils.set_cell_text(applicant_cells[1], letter_shared.display_name(applicant))
    docx_utils.set_cell_text(applicant_cells[2], applicant.get("passportNumber") or "—")
    # applicant_cells[3] ("Self") is fixed reference content — left untouched.
    docx_utils.set_cell_text(applicant_cells[4], occupations.get(applicant.get("id"), "") or "—")

    if not companions:
        # No companion rows needed — remove the now-templateless row rather
        # than leave literal "[Passenger 2 Name]" placeholder text in a
        # real client-facing letter (same judgement call as Japan's empty
        # hotel table above).
        companion_template_row_el.getparent().remove(companion_template_row_el)
        return

    # Clone BEFORE the template row is touched further — it must still be
    # attached to the table for clone_and_insert_after's addnext() to have
    # somewhere to anchor the new siblings (the same working technique
    # already proven in _fill_japan_hotel_table above; do not remove the
    # template row first, that leaves addnext() with no parent to insert
    # into).
    clone_els = docx_utils.clone_and_insert_after(companion_template_row_el, len(companions) - 1)
    row_els = [companion_template_row_el] + clone_els

    for idx, (row_el, companion) in enumerate(zip(row_els, companions), start=2):
        cells = _cells_of_row_element(table, row_el)
        docx_utils.set_cell_text(cells[0], str(idx))
        docx_utils.set_cell_text(cells[1], letter_shared.display_name(companion))
        docx_utils.set_cell_text(cells[2], companion.get("passportNumber") or "—")
        docx_utils.set_cell_text(
            cells[3], _relation_word_or_fallback(companion.get("relationToApplicant"), companion.get("sex")).capitalize()
        )
        docx_utils.set_cell_text(cells[4], occupations.get(companion.get("id"), "") or "—")


def generate_singapore_cover_letter_docx(
    applicant: dict,
    companions: list[dict],
    fields: dict,
    date_override: Optional[str] = None,
) -> bytes:
    """`fields` keys used: recipientText, travelStartDate, travelEndDate,
    fundingArrangement, hotelName, hotelAddress, occupations (a dict
    keyed by person id)."""
    if not applicant:
        raise ValueError("At least the applicant is required to generate this letter.")

    document = _load_template(SINGAPORE_TEMPLATE_PATH, "Singapore Cover Letter")
    paragraphs = document.paragraphs

    docx_utils.apply_token_map_to_paragraph(paragraphs[0], {"[Date]": letter_shared.format_letter_date(date_override)})
    docx_utils.set_paragraph_text(paragraphs[2], (fields.get("recipientText") or "").strip() or "—")

    family_clause = " along with my Family" if companions else ""
    para7 = (
        f"I, {letter_shared.display_name(applicant)} planning to visit Singapore for Tourism purposes"
        f"{family_clause} & would like to travel on {fields.get('travelStartDate') or '—'} and will stay until "
        f"{fields.get('travelEndDate') or '—'}."
    )
    docx_utils.set_paragraph_text(paragraphs[7], para7)

    docx_utils.apply_token_map_to_paragraph(
        paragraphs[9], {"[Funding Arrangement]": fields.get("fundingArrangement") or "—"}
    )
    docx_utils.apply_token_map_to_paragraph(
        paragraphs[10],
        {"[Hotel Name]": fields.get("hotelName") or "—", "[Hotel Address]": fields.get("hotelAddress") or "—"},
    )
    docx_utils.apply_token_map_to_paragraph(paragraphs[14], {"[Passenger 1 Name]": letter_shared.display_name(applicant)})
    docx_utils.apply_token_map_to_paragraph(
        paragraphs[15], {"[Email Address]": (applicant.get("email") or "").strip() or "—"}
    )
    docx_utils.apply_token_map_to_paragraph(paragraphs[16], {"[Contact No.]": _phone_or_dash(applicant.get("phone"))})

    # The template's own second, broken/mislabeled duplicate signature
    # block — deleted, not preserved (see module docstring). Captured by
    # reference before any of the above edits shift nothing (paragraph
    # indices 17-19 are untouched by any fill above), then removed last.
    for idx in (17, 18, 19):
        docx_utils.delete_paragraph(paragraphs[idx])

    _fill_singapore_passenger_table(document, applicant, companions, fields.get("occupations") or {})

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def generate_singapore_cover_letter_pdf(
    applicant: dict, companions: list[dict], fields: dict, date_override: Optional[str] = None
) -> bytes:
    docx_bytes = generate_singapore_cover_letter_docx(applicant, companions, fields, date_override)
    try:
        return docx_utils.convert_docx_bytes_to_pdf(docx_bytes, base_name="singapore_cover_letter")
    except docx_utils.DocxBuildError as e:
        raise CoverLetterError(str(e))
