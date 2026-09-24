"""
Khanna Travels & Holidays — Admin: Master Excel export
=======================================================

Phase 11. Unlike every prior document engine (Hotel Voucher, Authorization
Letters, Cover Letters, Invitation/Initors Letters), this one has no
reference `.docx` template to preserve — there is no existing "master Excel
export" format on file to match (project rule 5 only applies when a
reference template actually exists). Its layout is instead a straightforward,
professional data export designed from the real Application data model
(data_model.py / js/core/state.js): one row per Application, one row per
Traveller, and a real computed status breakdown — never fabricated figures.

This app has no persistent backend datastore for Application records (they
live in the browser's localStorage — see state.js's own header comment), so
the frontend POSTs its current in-memory list of applications and this
module turns that JSON straight into a workbook. Nothing is written to disk
here and nothing is logged — the same "don't expose passport data
unnecessarily" discipline as every other document engine (project rule 11),
just applied to a spreadsheet instead of a .docx.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


class AdminExportError(Exception):
    """Raised when the input isn't shaped like a list of Application
    records — never on merely empty/missing individual fields, which are
    real, legitimate states for a Draft application (rendered as an em
    dash, matching the em-dash-for-missing convention already used by
    hotel_voucher_engine.py)."""


HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
SUBTITLE_FONT = Font(size=9, italic=True, color="6B7280")

# Canonical status list — kept in sync with data_model.py / state.js
# (STATUSES). Used so the Status Summary sheet always shows every real
# status, including ones with a genuine zero count, rather than only the
# statuses that happen to appear in this particular export (an admin
# reading "0 Rejected" is real information; a missing row would look like
# an oversight).
STATUSES: List[str] = [
    "Draft",
    "New",
    "Documents Pending",
    "Documents Received",
    "Under Review",
    "Ready for Submission",
    "Submitted",
    "Processing",
    "Approved",
    "Rejected",
    "Cancelled",
]


def _dash(value: Any) -> str:
    """Missing/blank data renders as an em dash, matching the convention
    already established in hotel_voucher_engine.py — never blank (which
    could be misread as a spreadsheet error) and never a fabricated
    default."""
    if value is None:
        return "—"
    text = str(value).strip()
    return text if text else "—"


def _fmt_dt(value: Any) -> str:
    """Application timestamps are stored as ISO-8601 strings (state.js's
    own `nowIso()`). Rendered as a plain readable date/time; an unparsable
    value is shown exactly as given rather than hidden (project rule 9 —
    never silently swallow a value that doesn't fit)."""
    if not value:
        return "—"
    text = str(value)
    try:
        # Python's fromisoformat doesn't accept a trailing "Z" before 3.11's
        # relaxed parser; normalize it ourselves so this works across
        # whichever Python 3 the staff machine happens to have.
        normalized = text[:-1] + "+00:00" if text.endswith("Z") else text
        dt = datetime.fromisoformat(normalized)
        return dt.strftime("%d %b %Y %H:%M")
    except ValueError:
        return text


def _applicant_name(app: Dict[str, Any]) -> str:
    applicant = app.get("applicant") or {}
    name = (applicant.get("fullName") or "").strip()
    return name or "Untitled applicant"


def _autofit(ws: Worksheet, header_row_idx: int, min_width: int = 10, max_width: int = 42) -> None:
    """Real column auto-sizing computed from the sheet's own actual cell
    content (max string length per column, header included) — not a set of
    hardcoded widths guessed up front, since the right width genuinely
    depends on what data ends up in each export."""
    widths: Dict[int, int] = {}
    for row in ws.iter_rows(min_row=header_row_idx):
        for cell in row:
            if cell.value is None:
                continue
            length = len(str(cell.value))
            col = cell.column
            if length > widths.get(col, 0):
                widths[col] = length
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, width + 2))


def _write_title_block(ws: Worksheet, title: str, generated_at: datetime) -> int:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Khanna Travels & Holidays — generated " + generated_at.strftime("%d %b %Y %H:%M") + " UTC"
    ws["A2"].font = SUBTITLE_FONT
    return 4  # header row starts here, leaving one blank row after the title block


def _write_header(ws: Worksheet, row_idx: int, headers: List[str]) -> None:
    for col_idx, label in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = ws.cell(row=row_idx + 1, column=1).coordinate


def _applications_sheet(wb: Workbook, applications: List[Dict[str, Any]], generated_at: datetime) -> None:
    ws = wb.active
    ws.title = "Applications"
    header_row = _write_title_block(ws, "Applications Export", generated_at)

    headers = [
        "Application ID",
        "Status",
        "Applicant Name",
        "Passport Number",
        "Destination",
        "Visa Category",
        "Travellers",
        "Documents Verified %",
        "Travel Start",
        "Travel End",
        "Created By",
        "Created At",
        "Updated At",
    ]
    _write_header(ws, header_row, headers)

    row_idx = header_row + 1
    for app in applications:
        applicant = app.get("applicant") or {}
        travel = app.get("travel") or {}
        checklist = app.get("checklist") or {}
        passport = (applicant.get("passport") or {}).get("current") or {}
        travellers = app.get("travellers") or []

        values = [
            _dash(app.get("id")),
            _dash(app.get("status")),
            _applicant_name(app),
            _dash(passport.get("number")),
            _dash(travel.get("destination")),
            _dash(travel.get("visaCategory")),
            len(travellers),
            checklist.get("completionPct", 0),
            _dash(travel.get("startDate")),
            _dash(travel.get("endDate")),
            _dash(app.get("createdBy")),
            _fmt_dt(app.get("createdAt")),
            _fmt_dt(app.get("updatedAt")),
        ]
        for col_idx, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
        row_idx += 1

    if not applications:
        ws.cell(row=row_idx, column=1, value="No applications to export yet.")

    _autofit(ws, header_row)


def _travellers_sheet(wb: Workbook, applications: List[Dict[str, Any]], generated_at: datetime) -> None:
    ws = wb.create_sheet("Travellers")
    header_row = _write_title_block(ws, "Travellers Export", generated_at)

    headers = [
        "Application ID",
        "Primary Applicant",
        "Traveller Name",
        "Relation",
        "Sex",
        "Passport Number",
        "Passport Verification",
    ]
    _write_header(ws, header_row, headers)

    row_idx = header_row + 1
    for app in applications:
        applicant_name = _applicant_name(app)
        for traveller in app.get("travellers") or []:
            passport = (traveller.get("passport") or {}).get("current") or {}
            values = [
                _dash(app.get("id")),
                applicant_name,
                _dash(traveller.get("fullName")),
                _dash(traveller.get("relation")),
                _dash(traveller.get("sex")),
                _dash(passport.get("number")),
                _dash((traveller.get("passport") or {}).get("verificationStatus")),
            ]
            for col_idx, value in enumerate(values, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)
            row_idx += 1

    if row_idx == header_row + 1:
        ws.cell(row=row_idx, column=1, value="No travellers recorded across these applications.")

    _autofit(ws, header_row)


def _status_summary_sheet(wb: Workbook, applications: List[Dict[str, Any]], generated_at: datetime) -> None:
    ws = wb.create_sheet("Status Summary")
    header_row = _write_title_block(ws, "Status Summary", generated_at)
    _write_header(ws, header_row, ["Status", "Applications"])

    counts: Dict[str, int] = {s: 0 for s in STATUSES}
    for app in applications:
        status = app.get("status")
        if status not in counts:
            # An application carrying a status this export doesn't
            # recognize is still counted, not silently dropped — it just
            # gets its own row rather than being folded into a known one.
            counts[status] = 0
        counts[status] = counts.get(status, 0) + 1

    row_idx = header_row + 1
    for status in STATUSES:
        ws.cell(row=row_idx, column=1, value=status)
        ws.cell(row=row_idx, column=2, value=counts.get(status, 0))
        row_idx += 1
    for status, count in counts.items():
        if status not in STATUSES:
            ws.cell(row=row_idx, column=1, value=status + " (unrecognized status)")
            ws.cell(row=row_idx, column=2, value=count)
            row_idx += 1

    ws.cell(row=row_idx + 1, column=1, value="Total").font = Font(bold=True)
    ws.cell(row=row_idx + 1, column=2, value=len(applications)).font = Font(bold=True)

    _autofit(ws, header_row)


def build_applications_workbook(applications: List[Dict[str, Any]]) -> bytes:
    """Builds the real Master Excel export workbook and returns its raw
    .xlsx bytes. `applications` must be a list (an empty list is valid —
    a freshly-set-up install with zero applications yet is a real,
    honest state, not an error)."""
    if not isinstance(applications, list):
        raise AdminExportError("applications must be a list of Application records.")

    generated_at = datetime.now(timezone.utc)
    wb = Workbook()
    _applications_sheet(wb, applications, generated_at)
    _travellers_sheet(wb, applications, generated_at)
    _status_summary_sheet(wb, applications, generated_at)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
