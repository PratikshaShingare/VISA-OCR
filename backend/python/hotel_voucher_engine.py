"""
Khanna Travels & Holidays — Hotel Voucher document generator
==============================================================

Generates the company's actual "Hotel Voucher" document for one or more
hotel bookings on an Application, using the real, unmodified reference
template as the exact source of truth for layout, fonts, colours and the
company letterhead/footer (project rule 5 — "Document Template Rule").

Source of truth (read-only, NEVER modified by this module):
    reference-hotel-blocking-formats/Single Hotel/Hotel Booking Format.docx

For more than one hotel, the template's own "HOTEL VOUCHER" heading + table
block is cloned once per additional hotel. This exactly mirrors how the
real reference-hotel-blocking-formats/Multiple Hotels/ example itself
repeats that identical block per hotel (verified by inspecting that file
directly) — so a 3-hotel voucher generated here has the same structure as
the agency's own real 3-hotel example, not an approximation of it.

Only the applicant/booking data changes per block: hotel name, address,
city, phone, lead guest, room/guest counts, check-in/out dates, nights,
confirmation number and the guest-name list. Everything else — the red
"HOTEL VOUCHER" heading, the "Booking details" table structure, the hotel
policy text, the header/footer letterhead images — comes through untouched
because we only ever edit specific cells' text runs, never rebuild the
document from scratch.
"""

from __future__ import annotations

import copy
import io
import os
from datetime import datetime
from typing import Optional

from docx import Document
from docx.table import Table

import docx_utils

_THIS_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_PATH = os.path.normpath(
    os.path.join(
        _THIS_DIR,
        "..",
        "..",
        "reference-hotel-blocking-formats",
        "Single Hotel",
        "Hotel Booking Format.docx",
    )
)


class HotelVoucherError(docx_utils.DocxBuildError):
    """Raised whenever a voucher genuinely cannot be generated — never
    swallowed to produce a fake/empty document (project rule 9). Subclasses
    the shared DocxBuildError so callers that only catch the shared type
    (e.g. a future generic document endpoint) still see it, while existing
    callers that catch HotelVoucherError specifically are unaffected."""


def _load_template() -> Document:
    if not os.path.isfile(TEMPLATE_PATH):
        raise HotelVoucherError(
            "Hotel voucher reference template not found at "
            f"{TEMPLATE_PATH}. This file is part of the project's read-only "
            "reference material (reference-hotel-blocking-formats/) and must "
            "already exist on this machine — it is never created by this app."
        )
    # Document() only reads this file; nothing is ever written back to it.
    return Document(TEMPLATE_PATH)


def _find_hotel_block_table(document: Document):
    """Locates the template's single self-contained "HOTEL VOUCHER" block.
    In the reference file this is ONE top-level table whose single cell
    nests both the red "HOTEL VOUCHER" heading paragraph and the wrapper
    table holding everything else — not a heading paragraph sitting as a
    separate sibling before the table (verified directly against the
    template's own XML: document.paragraphs has no "HOTEL VOUCHER" text at
    all, because it lives nested inside the table cell). Cloning this one
    table element whole is therefore enough to duplicate an entire block.
    Fails loudly (not silently) if the reference template's structure ever
    changes in a way this engine can't safely handle — an honest error
    beats a wrong document."""
    if not document.tables:
        raise HotelVoucherError(
            "Hotel voucher template structure has changed (no table found) — "
            "cannot safely generate a voucher. Check reference-hotel-blocking-"
            "formats/Single Hotel/Hotel Booking Format.docx."
        )
    table = document.tables[0]
    try:
        heading_text = table.rows[0].cells[0].paragraphs[0].text.strip()
    except IndexError:
        heading_text = ""
    if heading_text != "HOTEL VOUCHER":
        raise HotelVoucherError(
            "Hotel voucher template structure has changed (expected the first "
            "table's first cell to start with a 'HOTEL VOUCHER' heading, found "
            f"{heading_text!r}) — cannot safely generate a voucher. Check "
            "reference-hotel-blocking-formats/Single Hotel/Hotel Booking Format.docx."
        )
    return table._tbl


_set_paragraph_text = docx_utils.set_paragraph_text
_set_cell_text = docx_utils.set_cell_text
_replace_token_in_cell = docx_utils.replace_token_in_cell


def _set_guest_names(cell, names: list[str]) -> None:
    """Fills the guest-name cell with one paragraph per name, adding or
    removing paragraphs (cloned from the cell's own first paragraph, so
    formatting matches) to fit however many names were supplied."""
    if not names:
        names = ["—"]

    while len(cell.paragraphs) < len(names):
        template_p = cell.paragraphs[0]._p
        new_p = copy.deepcopy(template_p)
        cell.paragraphs[-1]._p.addnext(new_p)
    while len(cell.paragraphs) > len(names) and len(cell.paragraphs) > 1:
        last = cell.paragraphs[-1]
        last._p.getparent().remove(last._p)

    for i, name in enumerate(names):
        _set_paragraph_text(cell.paragraphs[i], name or "—")


def _fmt_date(value: Optional[str]) -> str:
    if not value:
        return "—"
    try:
        return datetime.strptime(value, "%Y-%m-%d").strftime("%d %b %Y")
    except ValueError:
        return str(value)  # show whatever was given rather than hide it


def _format_guest_counts(hotel: dict) -> str:
    """Combines the (explicitly split, on request) adults/children counts
    into one "2 Adult(s), 2 Child(s)"-style string for both cells that used
    to just print the single, undifferentiated noOfGuests value. Children
    is only appended when a real, non-zero count was actually entered — a
    booking with no children shouldn't print "0 Child(s)"."""
    adults = str(hotel.get("noOfAdults") or "").strip()
    children = str(hotel.get("noOfChildren") or "").strip()
    parts = []
    if adults:
        parts.append(f"{adults} Adult(s)")
    if children and children != "0":
        parts.append(f"{children} Child(s)")
    return ", ".join(parts) if parts else "—"


def _nights(check_in: Optional[str], check_out: Optional[str]) -> Optional[int]:
    try:
        d1 = datetime.strptime(check_in, "%Y-%m-%d").date()
        d2 = datetime.strptime(check_out, "%Y-%m-%d").date()
        n = (d2 - d1).days
        return n if n >= 0 else None
    except (ValueError, TypeError):
        return None


def _fill_hotel_block(table: Table, hotel: dict) -> None:
    """Fills one HOTEL VOUCHER block's fields at their known structural cell
    positions (verified against the real template — see module docstring),
    rather than text-searching for placeholders: several placeholders (e.g.
    '[Date]', '[Hotel Address]') repeat identically more than once in the
    template, so only position reliably distinguishes them."""
    outer_cell = table.rows[0].cells[0]
    wrapper = outer_cell.tables[0]  # 2 rows: [0]=confirmation sentence, [1]=nested detail tables
    thanks_cell = wrapper.rows[0].cells[0]
    detail_cell = wrapper.rows[1].cells[0]
    booking = detail_cell.tables[0]  # 7x4 "Booking details" table
    guest = detail_cell.tables[1]  # 2x3 guest/room table
    # detail_cell.tables[2] is the static Hotel Policy table — left untouched.

    _replace_token_in_cell(thanks_cell, "[Confirmation No.]", hotel.get("confirmationNumber") or "—")

    _set_cell_text(booking.rows[1].cells[3], hotel.get("hotelName") or "—")
    _set_cell_text(booking.rows[2].cells[1], hotel.get("leadGuestName") or "—")
    _set_cell_text(booking.rows[2].cells[3], _format_guest_counts(hotel))
    _set_cell_text(booking.rows[3].cells[1], str(hotel.get("noOfRooms") or "—"))
    _set_cell_text(booking.rows[3].cells[3], hotel.get("phone") or "—")
    _set_cell_text(booking.rows[4].cells[1], _fmt_date(hotel.get("checkIn")))
    _set_cell_text(booking.rows[4].cells[3], _fmt_date(hotel.get("checkOut")))
    nights = _nights(hotel.get("checkIn"), hotel.get("checkOut"))
    _set_cell_text(booking.rows[5].cells[1], (f"{nights} Night(s)" if nights is not None else "—"))
    _set_cell_text(booking.rows[5].cells[3], hotel.get("city") or "—")
    _set_cell_text(booking.rows[6].cells[1], hotel.get("address") or "—")  # merged across cols 1-3

    guest_names = hotel.get("guestNames") or []
    guest_names = [n for n in guest_names if n and n.strip()]
    if not guest_names:
        guest_names = [hotel.get("leadGuestName") or "—"]
    _set_guest_names(guest.rows[1].cells[0], guest_names)
    _set_cell_text(guest.rows[1].cells[1], hotel.get("roomType") or "—")
    _set_cell_text(guest.rows[1].cells[2], _format_guest_counts(hotel))


def generate_hotel_voucher_docx(hotels: list[dict]) -> bytes:
    """Builds the Hotel Voucher .docx for one or more hotel bookings.
    Returns raw .docx bytes. Raises HotelVoucherError / ValueError on any
    problem — never returns a silently-wrong or empty document."""
    if not hotels:
        raise ValueError("At least one hotel is required to generate a voucher.")

    document = _load_template()
    table_el = _find_hotel_block_table(document)
    # The template's own trailing blank paragraph (body-level, after the
    # table) — reused as a spacer between blocks when there's more than one
    # hotel, so multiple vouchers in the same document aren't jammed
    # directly against each other. Real template content, not invented.
    spacer_el = document.paragraphs[-1]._p if document.paragraphs else None

    blocks_el = [table_el]
    last_el = table_el
    for _ in hotels[1:]:
        if spacer_el is not None:
            new_spacer_el = copy.deepcopy(spacer_el)
            last_el.addnext(new_spacer_el)
            last_el = new_spacer_el
        new_table_el = copy.deepcopy(table_el)
        last_el.addnext(new_table_el)
        blocks_el.append(new_table_el)
        last_el = new_table_el

    for t_el, hotel in zip(blocks_el, hotels):
        table = Table(t_el, document)
        _fill_hotel_block(table, hotel)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def generate_hotel_voucher_pdf(hotels: list[dict]) -> bytes:
    """Converts the generated .docx to PDF via a real LibreOffice headless
    conversion, so the PDF is the actual rendered document rather than a
    hand-recreated approximation. Requires LibreOffice ('soffice') to be
    installed on the machine running this backend — if it isn't, this
    raises a clear, honest error instead of returning nothing or a fake
    file (project rule 9)."""
    docx_bytes = generate_hotel_voucher_docx(hotels)
    try:
        return docx_utils.convert_docx_bytes_to_pdf(docx_bytes, base_name="hotel_voucher")
    except docx_utils.DocxBuildError as e:
        # Re-raised as HotelVoucherError so existing callers (app.py) that
        # catch that specific type are unaffected by this refactor.
        raise HotelVoucherError(str(e))
