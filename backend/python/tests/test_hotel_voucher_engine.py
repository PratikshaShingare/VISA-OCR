"""
Tests for hotel_voucher_engine.py, generating against the REAL
reference-hotel-blocking-formats/Single Hotel/Hotel Booking Format.docx
template (project rule 5 — the reference template is the source of truth;
these tests must fail loudly if that real file's structure ever changes in
a way the engine can't safely handle, not pass against a fake stand-in).
"""

import io

import pytest
from docx import Document

import hotel_voucher_engine as hve
from conftest import skip_if_missing, skip_if_missing_binary


def _all_paragraphs_including_nested_tables(container):
    """python-docx's `document.tables` / `cell.tables` are NOT recursive —
    a table nested inside a table cell (exactly the Hotel Voucher
    template's own real structure: an outer wrapper table containing a
    'Booking details' table and a guest/room table inside one cell) is
    invisible to a plain `doc.tables` walk. This helper walks paragraphs
    and nested tables recursively so tests can assert on the full,
    real rendered text."""
    paragraphs = list(container.paragraphs)
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(_all_paragraphs_including_nested_tables(cell))
    return paragraphs


def _full_text(doc: Document) -> str:
    return "\n".join(p.text for p in _all_paragraphs_including_nested_tables(doc))


def test_generate_hotel_voucher_docx_requires_at_least_one_hotel():
    with pytest.raises(ValueError):
        hve.generate_hotel_voucher_docx([])


def test_single_hotel_voucher_contains_the_real_booking_data(sample_hotel):
    skip_if_missing(hve.TEMPLATE_PATH)
    docx_bytes = hve.generate_hotel_voucher_docx([sample_hotel])
    assert docx_bytes[:2] == b"PK", "must be a real .docx (zip) file"

    doc = Document(io.BytesIO(docx_bytes))
    full_text = _full_text(doc)
    assert sample_hotel["hotelName"] in full_text
    assert sample_hotel["confirmationNumber"] in full_text
    assert sample_hotel["leadGuestName"] in full_text
    assert "01 Dec 2026" in full_text, "check-in date must be formatted, not left as raw ISO"
    assert "4 Night(s)" in full_text, "nights must be computed from check-in/check-out"


def test_multiple_hotels_clone_a_full_block_per_hotel(sample_hotels):
    skip_if_missing(hve.TEMPLATE_PATH)
    docx_bytes = hve.generate_hotel_voucher_docx(sample_hotels)
    doc = Document(io.BytesIO(docx_bytes))

    heading_count = sum(
        1
        for t in doc.tables
        if t.rows and t.rows[0].cells and t.rows[0].cells[0].paragraphs
        and t.rows[0].cells[0].paragraphs[0].text.strip() == "HOTEL VOUCHER"
    )
    assert heading_count == len(sample_hotels), (
        f"expected one cloned 'HOTEL VOUCHER' block per hotel ({len(sample_hotels)}), found {heading_count}"
    )

    full_text = _full_text(doc)
    for hotel in sample_hotels:
        assert hotel["hotelName"] in full_text
        assert hotel["confirmationNumber"] in full_text


def test_guest_names_list_is_filled_with_one_paragraph_per_guest():
    skip_if_missing(hve.TEMPLATE_PATH)
    hotel = {
        "hotelName": "Guest List Hotel",
        "confirmationNumber": "CONF-GL",
        "leadGuestName": "Lead Guest",
        "checkIn": "2026-06-01",
        "checkOut": "2026-06-03",
        "guestNames": ["Guest One", "Guest Two", "Guest Three"],
    }
    docx_bytes = hve.generate_hotel_voucher_docx([hotel])
    doc = Document(io.BytesIO(docx_bytes))
    full_text = _full_text(doc)
    for name in hotel["guestNames"]:
        assert name in full_text


def test_missing_fields_render_as_an_em_dash_not_blank_or_none():
    skip_if_missing(hve.TEMPLATE_PATH)
    minimal_hotel = {"hotelName": "Bare Hotel"}
    docx_bytes = hve.generate_hotel_voucher_docx([minimal_hotel])
    doc = Document(io.BytesIO(docx_bytes))
    full_text = _full_text(doc)
    assert "—" in full_text, "unfilled fields should show an honest placeholder, never 'None' or 'undefined'"
    assert "None" not in full_text


def test_generate_hotel_voucher_pdf_produces_a_real_pdf(sample_hotel):
    skip_if_missing(hve.TEMPLATE_PATH)
    skip_if_missing_binary("soffice")
    pdf_bytes = hve.generate_hotel_voucher_pdf([sample_hotel])
    assert pdf_bytes[:5] == b"%PDF-"
