"""
Tests for cover_letter_engine.py — the Europe / Japan / Singapore Cover
Letters, generated against the REAL reference-cover-letter-formats/ files
(project rule 5). Each region has its own real template with a different
fixed shape (Europe/Japan are always applicant + exactly one companion;
Singapore takes any number of companions, including zero), so each is
exercised on its own real template.
"""

import io

import pytest
from docx import Document

import cover_letter_engine as cle
from conftest import skip_if_missing, skip_if_missing_binary


def _all_paragraphs_including_nested_tables(container):
    paragraphs = list(container.paragraphs)
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(_all_paragraphs_including_nested_tables(cell))
    return paragraphs


def _full_text(doc: Document) -> str:
    return "\n".join(p.text for p in _all_paragraphs_including_nested_tables(doc))


EUROPE_FIELDS = {
    "recipientText": "Consulate General of France, Mumbai",
    "destinationCountry": "France",
    "travelStartDate": "2026-11-10",
    "travelEndDate": "2026-11-20",
    "cityCountryOfResidence": "Mumbai, India",
    "numberOfNights": "10 nights",
    "applicantEmploymentStatus": "Salaried",
    "applicantJobTitle": "Software Engineer",
    "applicantEmployerName": "Acme Corp",
    "companionJobTitle": "Designer",
    "companionEmployerName": "Beta LLC",
    "companionEmploymentStartYear": "2015",
    "fundingArrangement": "Self-funded by the applicant",
}


def test_europe_cover_letter_requires_applicant_and_companion():
    with pytest.raises(ValueError):
        cle.generate_europe_cover_letter_docx({}, {}, EUROPE_FIELDS)


def test_europe_cover_letter_fills_real_template(sample_applicant, sample_companion):
    skip_if_missing(cle.EUROPE_TEMPLATE_PATH)
    docx_bytes = cle.generate_europe_cover_letter_docx(sample_applicant, sample_companion, EUROPE_FIELDS)
    assert docx_bytes[:2] == b"PK"
    text = _full_text(Document(io.BytesIO(docx_bytes)))
    assert sample_applicant["fullName"] in text
    assert sample_applicant["passportNumber"] in text
    assert EUROPE_FIELDS["destinationCountry"] in text
    assert EUROPE_FIELDS["applicantEmployerName"] in text
    assert sample_applicant["email"] in text


def test_europe_cover_letter_pdf(sample_applicant, sample_companion):
    skip_if_missing(cle.EUROPE_TEMPLATE_PATH)
    skip_if_missing_binary("soffice")
    pdf_bytes = cle.generate_europe_cover_letter_pdf(sample_applicant, sample_companion, EUROPE_FIELDS)
    assert pdf_bytes[:5] == b"%PDF-"


# ---------------------------------------------------------------------------
# Phase 4 proof of concept — Europe Cover Letter "Word-like" body editor
# ---------------------------------------------------------------------------


def test_compose_europe_body_paragraphs_matches_generated_docx(sample_applicant, sample_companion):
    """The shared composer (also used to seed the Quill editor) must
    produce the EXACT same wording that ends up in the real generated
    document when nothing has been edited — never a second,
    hand-duplicated copy of this business logic (project rule 15)."""
    skip_if_missing(cle.EUROPE_TEMPLATE_PATH)
    blocks = cle.compose_europe_body_paragraphs(sample_applicant, sample_companion, EUROPE_FIELDS)
    assert len(blocks) == 3
    assert all(b["align"] == "justify" for b in blocks)

    docx_bytes = cle.generate_europe_cover_letter_docx(sample_applicant, sample_companion, EUROPE_FIELDS)
    paragraphs = Document(io.BytesIO(docx_bytes)).paragraphs
    assert paragraphs[8].text == blocks[0]["text"]
    assert paragraphs[9].text == blocks[1]["text"]
    assert paragraphs[10].text == blocks[2]["text"]


def test_compose_europe_body_html_is_escaped_and_carries_justify_class(sample_applicant, sample_companion):
    skip_if_missing(cle.EUROPE_TEMPLATE_PATH)
    html = cle.compose_europe_body_html(sample_applicant, sample_companion, EUROPE_FIELDS)
    assert html.count("<p") == 3
    assert html.count('class="ql-align-justify"') == 3
    assert sample_applicant["fullName"] in html


def test_compose_europe_body_html_escapes_special_characters():
    applicant = {"fullName": "A & B <Co>", "passportNumber": "X1", "placeOfIssue": "Y", "passportIssueDate": "Z"}
    companion = {"fullName": "C & D", "relationToApplicant": "Spouse", "sex": "F"}
    html = cle.compose_europe_body_html(applicant, companion, EUROPE_FIELDS)
    assert "&amp;" in html
    assert "<Co>" not in html


def test_europe_cover_letter_with_edited_body_html_replaces_only_body_paragraphs(sample_applicant, sample_companion):
    """Every paragraph OUTSIDE the editable 8-10 range — including the
    signature block, which comes AFTER a range whose paragraph count the
    edit shrinks from 3 to 2 — must still be filled correctly. This is
    the concrete regression test for the 'capture stable paragraph
    references before mutating the live document.paragraphs view' rule
    documented in docx_utils.replace_paragraph_range_with_html."""
    skip_if_missing(cle.EUROPE_TEMPLATE_PATH)
    edited_html = (
        '<p class="ql-align-justify">Edited paragraph one with <strong>bold</strong> text.</p>'
        '<p class="ql-align-center">Edited paragraph two, centered and <em>italic</em>.</p>'
    )
    docx_bytes = cle.generate_europe_cover_letter_docx(
        sample_applicant, sample_companion, EUROPE_FIELDS, edited_body_html=edited_html
    )
    paragraphs = Document(io.BytesIO(docx_bytes)).paragraphs

    assert paragraphs[8].text == "Edited paragraph one with bold text."
    assert paragraphs[8].runs[1].text == "bold" and paragraphs[8].runs[1].bold is True
    assert paragraphs[9].text == "Edited paragraph two, centered and italic."
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    assert paragraphs[9].alignment == WD_ALIGN_PARAGRAPH.CENTER

    # The body shrank from 3 paragraphs to 2, so every paragraph the real
    # generator fills AFTER the body range now sits one index earlier —
    # this must still be the real funding/signature text, not garbage
    # left over from a stale index into the old paragraph list.
    full_text = "\n".join(p.text for p in paragraphs)
    assert EUROPE_FIELDS["fundingArrangement"] in full_text
    assert sample_applicant["fullName"] in full_text
    assert sample_applicant["email"] in full_text
    assert "9819347139" in full_text or (sample_applicant.get("phone") or "") in full_text


def test_europe_cover_letter_with_empty_edited_body_html_leaves_one_blank_paragraph(sample_applicant, sample_companion):
    """An explicitly-cleared editor (editedBodyHtml="") is a real edit,
    not 'no edit was made' — the letter must never end up silently
    missing a paragraph where the template expects one."""
    skip_if_missing(cle.EUROPE_TEMPLATE_PATH)
    docx_bytes = cle.generate_europe_cover_letter_docx(
        sample_applicant, sample_companion, EUROPE_FIELDS, edited_body_html=""
    )
    paragraphs = Document(io.BytesIO(docx_bytes)).paragraphs
    full_text = "\n".join(p.text for p in paragraphs)
    # The signature block must still be present and correct even though
    # the body range collapsed to a single blank paragraph.
    assert sample_applicant["fullName"] in full_text
    assert sample_applicant["email"] in full_text


JAPAN_FIELDS = {
    "destinationCountry": "Japan",
    "travelStartDate": "2026-10-01",
    "travelEndDate": "2026-10-10",
    "applicantEmployerOccupation": "Software Engineer at Acme Corp",
    "companionOccupation": "Designer",
    "fundingArrangement": "Self-funded",
}


def test_japan_cover_letter_requires_applicant_and_companion():
    with pytest.raises(ValueError):
        cle.generate_japan_cover_letter_docx({}, {}, JAPAN_FIELDS, [])


def test_japan_cover_letter_fills_real_template_with_hotels(sample_applicant, sample_companion):
    skip_if_missing(cle.JAPAN_TEMPLATE_PATH)
    hotels = [
        {"name": "Tokyo Test Hotel", "checkIn": "2026-10-01", "checkOut": "2026-10-05", "contactNo": "+81 3 0000 0000"},
        {"name": "Osaka Test Hotel", "checkIn": "2026-10-05", "checkOut": "2026-10-10", "contactNo": "+81 6 0000 0000"},
    ]
    docx_bytes = cle.generate_japan_cover_letter_docx(sample_applicant, sample_companion, JAPAN_FIELDS, hotels)
    text = _full_text(Document(io.BytesIO(docx_bytes)))
    assert sample_applicant["fullName"] in text
    assert sample_companion["fullName"] in text
    for hotel in hotels:
        assert hotel["name"] in text


def test_japan_cover_letter_with_no_hotels_removes_the_templateless_row(sample_applicant, sample_companion):
    skip_if_missing(cle.JAPAN_TEMPLATE_PATH)
    # Must not raise, and must not leave literal "[Hotel Name 1]" placeholder text.
    docx_bytes = cle.generate_japan_cover_letter_docx(sample_applicant, sample_companion, JAPAN_FIELDS, [])
    text = _full_text(Document(io.BytesIO(docx_bytes)))
    assert "[Hotel Name" not in text


def test_japan_cover_letter_pdf(sample_applicant, sample_companion):
    skip_if_missing(cle.JAPAN_TEMPLATE_PATH)
    skip_if_missing_binary("soffice")
    pdf_bytes = cle.generate_japan_cover_letter_pdf(sample_applicant, sample_companion, JAPAN_FIELDS, [])
    assert pdf_bytes[:5] == b"%PDF-"


SINGAPORE_FIELDS = {
    "recipientText": "Consulate General of Singapore, Mumbai",
    "travelStartDate": "2026-09-01",
    "travelEndDate": "2026-09-08",
    "fundingArrangement": "Self-funded",
    "hotelName": "Singapore Test Hotel",
    "hotelAddress": "1 Test Ave, Singapore",
    "occupations": {},
}


def test_singapore_cover_letter_requires_at_least_the_applicant():
    with pytest.raises(ValueError):
        cle.generate_singapore_cover_letter_docx({}, [], SINGAPORE_FIELDS)


def test_singapore_cover_letter_with_no_companions(sample_applicant):
    skip_if_missing(cle.SINGAPORE_TEMPLATE_PATH)
    docx_bytes = cle.generate_singapore_cover_letter_docx(sample_applicant, [], SINGAPORE_FIELDS)
    text = _full_text(Document(io.BytesIO(docx_bytes)))
    assert sample_applicant["fullName"] in text
    assert "along with my Family" not in text, "solo traveller must not get the family clause"


def test_singapore_cover_letter_with_companions(sample_applicant, sample_companion):
    skip_if_missing(cle.SINGAPORE_TEMPLATE_PATH)
    docx_bytes = cle.generate_singapore_cover_letter_docx(sample_applicant, [sample_companion], SINGAPORE_FIELDS)
    text = _full_text(Document(io.BytesIO(docx_bytes)))
    assert sample_applicant["fullName"] in text
    assert "along with my Family" in text
    assert sample_companion["fullName"] in text, "companion must appear in the passenger table"


def test_singapore_cover_letter_pdf(sample_applicant):
    skip_if_missing(cle.SINGAPORE_TEMPLATE_PATH)
    skip_if_missing_binary("soffice")
    pdf_bytes = cle.generate_singapore_cover_letter_pdf(sample_applicant, [], SINGAPORE_FIELDS)
    assert pdf_bytes[:5] == b"%PDF-"
