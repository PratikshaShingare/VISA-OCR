"""
Whole-document editable blocks (document_blocks.py): every body paragraph of
a generated letter can be edited, not just the engine's named fields.
"""

import io

from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

import app as app_module
import authorization_letter_engine as auth
import cover_letter_engine as cover
import document_blocks
import docx_utils
import invitation_letter_engine as inv
from conftest import skip_if_missing_binary

client = TestClient(app_module.app)

PEOPLE = [{"fullName": "A One", "passportNumber": "P1"}]


def _texts(docx_bytes):
    """Every paragraph's text in document order, however deeply nested."""
    doc = Document(io.BytesIO(docx_bytes))
    return ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in doc.element.body.iter(qn("w:p"))]


def _passport_docx(edited=None, shader=None):
    return auth.generate_passport_authorization_docx(
        PEOPLE, "VFS Delhi", "Some address", "Staff Name", edited_fields=edited, _shader=shader
    )


def _blocks_of(docx_bytes):
    return document_blocks.enumerate_blocks(Document(io.BytesIO(docx_bytes)))


def test_block_keys_are_stable_and_unique_even_for_repeated_text():
    doc = Document()
    for t in ("Yours faithfully,", "Body", "Yours faithfully,", ""):
        doc.add_paragraph(t)
    blocks = document_blocks.enumerate_blocks(doc)
    assert len(blocks) == 3, "empty paragraphs are never blocks"
    assert len({b.key for b in blocks}) == 3
    assert blocks[0].key.endswith("-0") and blocks[2].key.endswith("-1")
    assert blocks[0].key[:-2] == blocks[2].key[:-2]
    again = document_blocks.enumerate_blocks(doc)
    assert [b.key for b in again] == [b.key for b in blocks]


def test_keys_ignore_case_and_whitespace_but_not_words():
    a, b = Document(), Document()
    a.add_paragraph("Dear  Sir/Madam,")
    b.add_paragraph("dear sir/madam,")
    assert document_blocks.enumerate_blocks(a)[0].key == document_blocks.enumerate_blocks(b)[0].key
    c = Document()
    c.add_paragraph("Dear Sir,")
    assert document_blocks.enumerate_blocks(c)[0].key != document_blocks.enumerate_blocks(a)[0].key


def test_images_fields_and_section_breaks_are_never_blocks():
    doc = Document()
    p = doc.add_paragraph("page ")
    run = p.add_run()
    fld = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"})
    run._r.append(fld)
    doc.add_paragraph("real text")
    keys_text = [b.text for b in document_blocks.enumerate_blocks(doc)]
    assert keys_text == ["real text"]


def test_a_block_edit_replaces_exactly_that_paragraph():
    base = _passport_docx()
    block = next(b for b in _blocks_of(base) if b.text.strip().startswith("Thanking You"))
    edited = _passport_docx({block.key: "<p>Warm regards,</p>"})
    before, after = _texts(base), _texts(edited)
    assert "Thanking You," in before and "Thanking You," not in after
    assert after.count("Warm regards,") == 1
    assert len(after) == len(before)
    assert [t for t in before if not t.startswith("Thanking")] == [t for t in after if t != "Warm regards,"]


def test_an_edit_can_add_paragraphs_a_list_and_a_table():
    base = _passport_docx()
    block = next(b for b in _blocks_of(base) if b.text.strip().startswith("Thanking You"))
    html = (
        "<p>Please note:</p>"
        '<ol><li data-list="bullet"><span class="ql-ui"></span>First point</li></ol>'
        '<table><tbody><tr><td data-row="a">X</td><td data-row="a">Y</td></tr></tbody></table>'
    )
    doc = Document(io.BytesIO(_passport_docx({block.key: html})))
    texts = [p.text for p in doc.paragraphs]
    assert "Please note:" in texts and "First point" in texts
    assert len(doc.tables) == 1 and [c.text for c in doc.tables[0].rows[0].cells] == ["X", "Y"]


def test_an_edit_with_no_matching_block_is_reported_not_dropped_silently():
    shader = docx_utils.ProbeShader()
    _passport_docx({"blk-deadbeef-0": "<p>x</p>"}, shader=shader)
    assert shader.unmatched == ["blk-deadbeef-0"]


def test_every_block_is_shaded_exactly_once_and_seeded():
    shader = docx_utils.ProbeShader()
    data = _passport_docx(shader=shader)
    block_keys = [k for k in shader.keys if document_blocks.is_block_key(k)]
    assert block_keys, "the letter has blocks beyond its named fields"
    assert set(block_keys) == set(shader.seeds)
    named = [k for k in shader.keys if not document_blocks.is_block_key(k)]
    assert set(named) == {"date", "recipientCentre", "recipientAddress", "applicantSentence", "collectorSentence"}
    # the same block keys exist in the real (unshaded) document
    real_keys = {b.key for b in _blocks_of(_passport_docx())}
    assert set(block_keys) <= real_keys
    # probe text == real text
    assert _texts(data) == _texts(_passport_docx())


def test_seed_html_states_formatting_alignment_and_lists():
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("Bold red")
    r.bold = True
    from docx.shared import Pt, RGBColor

    r.font.color.rgb = RGBColor(0xFF, 0, 0)
    r.font.size = Pt(13)
    p.alignment = 1
    seed = document_blocks.paragraph_seed_html(doc, p._p)
    assert seed.startswith('<p class="ql-align-center">')
    assert "<strong>" in seed and "color: #FF0000" in seed and "font-size: 13pt" in seed


def test_seed_then_save_is_lossless_for_text_and_run_formatting():
    base = _passport_docx()
    shader = docx_utils.ProbeShader()
    _passport_docx(shader=shader)
    key = next(k for k in shader.keys if document_blocks.is_block_key(k) and "contact" in shader.seeds[k].lower())
    out = Document(io.BytesIO(_passport_docx({key: shader.seeds[key]})))
    assert [p.text for p in out.paragraphs] == [p.text for p in Document(io.BytesIO(base)).paragraphs]


def test_edits_work_on_every_letter_engine():
    a = {"fullName": "A B", "passportNumber": "P", "salutation": "Mr", "sex": "Male", "phone": "9999999999",
         "email": "a@b.c", "isApplicant": True, "occupation": "x"}
    b = dict(a, fullName="C D", relationToApplicant="Spouse", sex="Female", salutation="Mrs", isApplicant=False)
    f = {"recipientText": "The Consulate", "destinationCountry": "Japan", "travelStartDate": "2026-12-01",
         "travelEndDate": "2026-12-10", "applicantEmployerOccupation": "Mgr", "companionOccupation": "x",
         "fundingArrangement": "Self", "hotelName": "H", "hotelAddress": "Addr", "occupations": {}}
    cases = {
        "japan": lambda e: cover.generate_japan_cover_letter_docx(a, b, f, [], None, edited_blocks=e),
        "singapore": lambda e: cover.generate_singapore_cover_letter_docx(a, [b], f, None, edited_blocks=e),
    }
    for name, build in cases.items():
        base = build(None)
        blocks = _blocks_of(base)
        assert len(blocks) > 5, name
        target = blocks[-1]
        edited = build({target.key: "<p>EDITED LINE</p>"})
        assert "EDITED LINE" in _texts(edited), name
        assert target.text.strip() not in _texts(edited), name


def test_layout_reports_blocks_with_seed_html_over_http():
    skip_if_missing_binary("soffice")
    res = client.post(
        "/api/documents/company-authorization",
        json={"people": [{"fullName": "A", "passportNumber": "P1"}], "recipientText": "Visa Officer", "format": "layout"},
    )
    groups = res.json()["groups"][0]
    blocks = {k: v for k, v in groups.items() if k.startswith("blk-")}
    assert blocks and all(v["html"].startswith(("<p", "<ol")) for v in blocks.values())


def test_edit_changes_the_real_pdf_and_the_layout_follows():
    skip_if_missing_binary("soffice")
    shader = docx_utils.ProbeShader()
    _passport_docx(shader=shader)
    key = next(k for k in shader.keys if document_blocks.is_block_key(k) and "contact" in shader.seeds[k].lower())
    plain = auth.compute_passport_field_layout(PEOPLE, "VFS Delhi", "Some address", "Staff Name")
    long_html = "<p>" + "A much longer closing remark. " * 30 + "</p>"
    edited = auth.compute_passport_field_layout(
        PEOPLE, "VFS Delhi", "Some address", "Staff Name", edited_fields={key: long_html}
    )
    h0 = plain["groups"][0][key]["rect"]
    h1 = edited["groups"][0][key]["rect"]
    assert (h1[3] - h1[1]) > (h0[3] - h0[1]) * 3
    assert edited["unmatched"] == []
    # the original seed is kept so "reset" returns to the generated wording
    assert edited["groups"][0][key]["html"] == plain["groups"][0][key]["html"]


# ------------------------------------------------------------------ hotel voucher

import hotel_voucher_engine as hv

HOTEL = {
    "hotelName": "Hotel X", "phone": "1", "address": "addr", "city": "Paris", "confirmationNumber": "C1",
    "leadGuestName": "Mr A", "noOfAdults": "2", "noOfChildren": "0", "noOfRooms": "1", "roomType": "Dbl",
    "checkIn": "2026-12-01", "checkOut": "2026-12-05", "guestNames": ["Mr A", "Mrs B"],
}


def test_hotel_voucher_blocks_are_scoped_to_each_hotel():
    skip_if_missing_binary("soffice")
    layout = hv.compute_field_layout([HOTEL, dict(HOTEL, hotelName="Hotel Y")])
    g0, g1 = layout["hotels"]
    blocks0 = {k for k in g0 if k.startswith("blk-")}
    assert blocks0 and blocks0 == {k for k in g1 if k.startswith("blk-")}
    assert layout["unmatched"] == [[], []]
    assert all("html" in g0[k] for k in blocks0)
    # named cells are not also offered as blocks
    assert "hotelName" in g0 and not any(k for k in blocks0 if "Hotel X" in g0[k]["html"])


def test_hotel_voucher_block_edit_changes_only_that_hotel():
    skip_if_missing_binary("soffice")
    layout = hv.compute_field_layout([HOTEL, HOTEL])
    key = next(k for k, v in layout["hotels"][0].items() if k.startswith("blk-") and "policy" in v["html"].lower())
    edited = dict(HOTEL, editedFields={key: "<p>CUSTOM POLICY LINE</p>"})
    docx = hv.generate_hotel_voucher_docx([edited, HOTEL])
    texts = _texts(docx)
    assert texts.count("CUSTOM POLICY LINE") == 1
    assert sum(1 for t in texts if "policy" in t.lower()) >= 1
