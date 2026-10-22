"""
Editable-preview layout probes — where each editable value of a letter sits
on the page of its REAL generated PDF (docx_utils.ProbeShader +
read_probe_layout, and each engine's compute_*_field_layout).

Covers: the shared helpers on a small synthetic document, then every letter
engine against its real reference template (one rectangle per editable
value, inside the page, none overlapping another, the template's own
alignment reported, and the probe copy's TEXT identical to the real
document's — the probe must change nothing but shading), plus the
`format: "layout"` HTTP contract.
"""

import io

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

import app as app_module
import authorization_letter_engine as auth
import cover_letter_engine as cover
import docx_utils
import invitation_letter_engine as inv
import required_document_letter_engine as rdl
import visa_requirements_data
from conftest import skip_if_missing_binary

client = TestClient(app_module.app)


def _all_text(docx_bytes: bytes) -> list[str]:
    doc = Document(io.BytesIO(docx_bytes))
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                out.extend(p.text for p in cell.paragraphs)
    return out


def _assert_sane_layout(layout: dict, expected_keys: set):
    assert len(layout["groups"]) == 1
    fields = layout["groups"][0]
    # Named fields are exactly the engine's own; every OTHER key is a
    # whole-document block (document_blocks.py), checked in test_document_blocks.
    assert {k for k in fields if not k.startswith("blk-")} == expected_keys
    rects = []
    for key, info in fields.items():
        page = layout["pages"][info["page"]]
        l, t, r, b = info["rect"]
        assert 0 <= l < r <= page["width"] + 0.5, key
        assert 0 <= t < b <= page["height"] + 0.5, key
        assert info["fontSize"] and info["fontSize"] > 4, key
        assert info["align"] in ("left", "center", "right", "justify"), key
        rects.append((key, info["page"], info["rect"]))
    for i, (ka, pa, ra) in enumerate(rects):
        for kb, pb, rb in rects[i + 1:]:
            if pa != pb:
                continue
            overlap_h = min(ra[3], rb[3]) - max(ra[1], rb[1])
            assert overlap_h <= 1.0, f"{ka} and {kb} overlap on the page"


# ---------------------------------------------------------------- helpers


def test_probe_rgb_is_unique_and_decodable():
    seen = {docx_utils.probe_rgb(i) for i in range(500)}
    assert len(seen) == 500


def test_shade_element_fill_puts_shd_before_rpr_and_jc():
    doc = Document()
    p = doc.add_paragraph("x")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = None
    ppr = p._p.get_or_add_pPr()
    docx_utils.shade_element_fill(ppr, (30, 40, 177))
    tags = [c.tag for c in ppr]
    assert qn("w:shd") in tags
    assert tags.index(qn("w:shd")) < tags.index(qn("w:jc"))
    # shading again replaces rather than stacks
    docx_utils.shade_element_fill(ppr, (31, 40, 177))
    assert len([c for c in ppr if c.tag == qn("w:shd")]) == 1


def test_probe_shader_reports_the_templates_own_alignment():
    doc = Document()
    p = doc.add_paragraph("justified")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    q = doc.add_paragraph("plain")
    shader = docx_utils.ProbeShader()
    shader.shade("a", p)
    shader.shade("b", q)
    assert shader.aligns == {"a": "justify", "b": "left"}
    assert list(shader.decode().values()) == [(0, "a"), (0, "b")]


def test_read_probe_layout_on_a_synthetic_document():
    skip_if_missing_binary("soffice")
    doc = Document()
    doc.add_paragraph("Fixed heading")
    target = doc.add_paragraph("An editable value")
    doc.add_paragraph("Fixed footer line")
    shader = docx_utils.ProbeShader()
    shader.shade("value", target)
    buf = io.BytesIO()
    doc.save(buf)
    pdf = docx_utils.convert_docx_bytes_to_pdf(buf.getvalue(), base_name="synthetic_probe")
    layout = shader.layout_from_pdf(pdf)
    info = layout["groups"][0]["value"]
    assert info["page"] == 0
    l, t, r, b = info["rect"]
    assert r > l and b > t
    assert info["fontSize"] and info["fontSize"] > 4


# ------------------------------------------------------- real templates

PEOPLE = [
    {"salutation": "Mr", "fullName": "Test One", "passportNumber": "X1234567", "phone": "9999999999",
     "email": "a@b.co", "sex": "Male", "relationToApplicant": "", "isApplicant": True},
]


def test_passport_authorization_layout():
    skip_if_missing_binary("soffice")
    layout = auth.compute_passport_field_layout(PEOPLE, "VFS Delhi", "Some address", "Staff Name")
    _assert_sane_layout(layout, {"date", "recipientCentre", "recipientAddress", "applicantSentence", "collectorSentence"})
    assert layout["groups"][0]["date"]["align"] == "right"
    assert layout["groups"][0]["applicantSentence"]["align"] == "justify"  # read from the template, not guessed


def test_passport_authorization_layout_follows_an_edited_value():
    skip_if_missing_binary("soffice")
    base = auth.compute_passport_field_layout(PEOPLE, "VFS Delhi", "Some address", "Staff Name")
    long_edit = "<p>" + ("A much longer authorization sentence. " * 12) + "</p>"
    edited = auth.compute_passport_field_layout(
        PEOPLE, "VFS Delhi", "Some address", "Staff Name", edited_fields={"applicantSentence": long_edit}
    )
    h0 = base["groups"][0]["applicantSentence"]["rect"]
    h1 = edited["groups"][0]["applicantSentence"]["rect"]
    assert (h1[3] - h1[1]) > (h0[3] - h0[1]) * 2, "the field must grow with its edited text"
    assert edited["groups"][0]["collectorSentence"]["rect"][1] > base["groups"][0]["collectorSentence"]["rect"][1]


def test_company_authorization_layout_has_one_line_per_applicant():
    skip_if_missing_binary("soffice")
    people = [{"fullName": "A One", "passportNumber": "P1"}, {"fullName": "B Two", "passportNumber": "P2"}]
    layout = auth.compute_company_field_layout(people, "The Visa Officer")
    _assert_sane_layout(layout, {"date", "recipient", "applicant_0", "applicant_1"})


def test_probe_copy_text_matches_the_real_document():
    people = [{"fullName": "A One", "passportNumber": "P1"}]
    real = auth.generate_company_authorization_docx(people, "The Visa Officer")
    probe = auth.generate_company_authorization_docx(people, "The Visa Officer", _shader=docx_utils.ProbeShader())
    assert _all_text(real) == _all_text(probe)


INVITER = {
    "fullName": "Inv Test", "addressLine1": "1 St", "addressLine2": "", "cityPostcode": "X1", "country": "UK",
    "cityCountry": "London, UK", "passportNumber": "P", "fullAddress": "1 St, X1", "studyingOrWorking": "working",
    "universityOrCompany": "Co", "phone": "123", "email": "a@b.c",
}
INVITEES = [
    {"fullName": "A B", "passportNumber": "P1", "placeOfIssue": "Delhi", "passportIssueDate": "2020-01-01"},
    {"fullName": "C D", "passportNumber": "P2", "placeOfIssue": "Delhi", "passportIssueDate": "2020-01-01"},
]
INV_FIELDS = {
    "travelStartDate": "2026-12-01", "travelEndDate": "2026-12-10", "purpose": "Tourism",
    "accommodationDetails": "with the inviter", "returnDate": "2026-12-10", "fundingArrangement": "the inviter",
}


def test_invitation_layout_has_every_editable_value():
    skip_if_missing_binary("soffice")
    layout = inv.compute_invitation_field_layout(INVITER, INVITEES, INV_FIELDS)
    _assert_sane_layout(layout, set(inv.INVITATION_EDITABLE_FIELD_KEYS))


def test_invitation_probe_copy_text_matches_the_real_document():
    real = inv.generate_invitation_letter_docx(INVITER, INVITEES, INV_FIELDS)
    probe = inv.generate_invitation_letter_docx(INVITER, INVITEES, INV_FIELDS, _shader=docx_utils.ProbeShader())
    assert _all_text(real) == _all_text(probe)


def test_required_document_letter_layout():
    skip_if_missing_binary("soffice")
    ids = list(visa_requirements_data.DOCUMENT_ORDER)[:4]
    layout = rdl.compute_required_document_letter_field_layout(
        "France", "Tourist", "Employed", ids, processing_time="3 weeks", visa_fees="EUR 90"
    )
    _assert_sane_layout(layout, set(rdl.REQUIRED_DOCUMENT_LETTER_EDITABLE_FIELD_KEYS))
    checklist = layout["groups"][0]["checklistItems"]["rect"]
    assert (checklist[3] - checklist[1]) > 40, "the whole checklist is one block spanning every line"


def test_europe_cover_letter_layout_is_one_body_block():
    skip_if_missing_binary("soffice")
    a = {"fullName": "A B", "passportNumber": "P", "salutation": "Mr", "sex": "Male", "phone": "9999999999",
         "email": "a@b.c", "isApplicant": True, "occupation": "x"}
    b = dict(a, fullName="C D", relationToApplicant="Spouse", sex="Female", salutation="Mrs", isApplicant=False)
    f = {"recipientText": "The Consulate", "destinationCountry": "France", "travelStartDate": "2026-12-01",
         "travelEndDate": "2026-12-10", "cityCountryOfResidence": "Delhi, India", "numberOfNights": "9",
         "applicantEmploymentStatus": "Employed", "applicantJobTitle": "Mgr", "applicantEmployerName": "Co",
         "companionJobTitle": "x", "companionEmployerName": "y", "companionEmploymentStartYear": "2015",
         "fundingArrangement": "Self"}
    layout = cover.compute_europe_field_layout(a, b, f)
    _assert_sane_layout(layout, {"body"})


# ------------------------------------------------------------- HTTP


def test_layout_format_returns_json_for_company_authorization():
    skip_if_missing_binary("soffice")
    res = client.post(
        "/api/documents/company-authorization",
        json={"people": [{"fullName": "A", "passportNumber": "P1"}], "recipientText": "Visa Officer", "format": "layout"},
    )
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/json")
    body = res.json()
    assert set(body) == {"pages", "groups", "unmatched"}
    assert "applicant_0" in body["groups"][0]


def test_layout_format_works_for_singapore_cover_letter_too():
    skip_if_missing_binary("soffice")
    payload = {
        "region": "Singapore",
        "people": [{"id": "applicant", "fullName": "A B", "passportNumber": "P", "isApplicant": True}],
        "recipientText": "x", "destinationCountry": "Singapore", "travelStartDate": "2026-12-01",
        "travelEndDate": "2026-12-05", "fundingArrangement": "Self", "hotelName": "H", "hotelAddress": "Addr",
        "format": "layout",
    }
    res = client.post("/api/documents/cover-letter", json=payload)
    assert res.status_code == 200
    groups = res.json()["groups"][0]
    assert groups and all(k.startswith("blk-") for k in groups)


def test_layout_format_passes_the_edited_fields_through():
    skip_if_missing_binary("soffice")
    base = {"people": [{"fullName": "A", "passportNumber": "P1"}], "recipientText": "Visa Officer", "format": "layout"}
    plain = client.post("/api/documents/company-authorization", json=base).json()
    edited = client.post(
        "/api/documents/company-authorization",
        json=dict(base, editedFields={"recipient": "<p>" + "Line. " * 80 + "</p>"}),
    ).json()
    r0 = plain["groups"][0]["recipient"]["rect"]
    r1 = edited["groups"][0]["recipient"]["rect"]
    assert (r1[3] - r1[1]) > (r0[3] - r0[1]) * 2
