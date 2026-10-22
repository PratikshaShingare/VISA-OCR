"""
Tests for authorization_letter_engine.py — Passport Authorization Letter
(single + two travellers) and Company Authorization Letter, generated
against the REAL reference-templates/ files (project rule 5). These
templates are real former-client letters with the client's own data mostly
NOT bracketed (see the module's own docstring) — every dynamic field is
filled by verified structural position, so these tests exist specifically
to catch a template-structure regression the position-based fills would
otherwise silently mis-fill.
"""

import io

import pytest
from docx import Document

import authorization_letter_engine as ale
from conftest import skip_if_missing, skip_if_missing_binary


def _full_text(doc: Document) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


# ---------------------------------------------------------------------------
# Passport Authorization Letter
# ---------------------------------------------------------------------------


def test_passport_auth_requires_at_least_one_traveller():
    with pytest.raises(ValueError):
        ale.generate_passport_authorization_docx([], "Some Centre", "Some Address", "Some Collector")


def test_passport_auth_rejects_more_than_two_travellers(sample_applicant, sample_companion):
    third = dict(sample_applicant, fullName="Third Person", passportNumber="T0000003")
    with pytest.raises(ValueError):
        ale.generate_passport_authorization_docx(
            [sample_applicant, sample_companion, third], "Some Centre", "Some Address", "Some Collector"
        )


def test_passport_auth_requires_collector_and_centre_name(sample_applicant):
    with pytest.raises(ValueError):
        ale.generate_passport_authorization_docx([sample_applicant], "Some Centre", "Some Address", "")
    with pytest.raises(ValueError):
        ale.generate_passport_authorization_docx([sample_applicant], "", "Some Address", "Some Collector")


def test_passport_auth_single_traveller_fills_real_template(sample_applicant):
    skip_if_missing(ale.PASSPORT_SINGLE_TEMPLATE_PATH)
    docx_bytes = ale.generate_passport_authorization_docx(
        [sample_applicant],
        recipient_centre_name="Test Visa Application Centre",
        recipient_address="123 Test Road, Test City",
        collector_name="Ms. Test Collector",
    )
    assert docx_bytes[:2] == b"PK"
    doc = Document(io.BytesIO(docx_bytes))
    text = _full_text(doc)
    assert sample_applicant["fullName"] in text
    assert sample_applicant["passportNumber"] in text
    assert "Ms. Test Collector" in text
    assert "Test Visa Application Centre" in text
    assert sample_applicant["email"] in text


def test_passport_auth_two_travellers_fills_both_names(sample_applicant, sample_companion):
    skip_if_missing(ale.PASSPORT_MULTIPLE_TEMPLATE_PATH)
    docx_bytes = ale.generate_passport_authorization_docx(
        [sample_applicant, sample_companion],
        recipient_centre_name="Test Visa Application Centre",
        recipient_address="123 Test Road, Test City",
        collector_name="Ms. Test Collector",
    )
    doc = Document(io.BytesIO(docx_bytes))
    text = _full_text(doc)
    assert sample_applicant["fullName"] in text
    assert sample_applicant["passportNumber"] in text
    assert sample_companion["fullName"] in text
    assert sample_companion["passportNumber"] in text


def test_passport_auth_never_leaks_the_real_former_clients_data(sample_applicant):
    """Project rule 11 — the real former client's name/passport/phone/email
    baked into these reference letters must never survive into generated
    output. A regression here would be a genuine data-leak bug, not just a
    cosmetic one."""
    skip_if_missing(ale.PASSPORT_SINGLE_TEMPLATE_PATH)
    docx_bytes = ale.generate_passport_authorization_docx(
        [sample_applicant], "Test Centre", "Test Address", "Test Collector"
    )
    doc = Document(io.BytesIO(docx_bytes))
    text = _full_text(doc)
    assert "Dilip Bijlani" not in text, "the real former client's name must never appear in generated output"
    assert "9819347139" not in text, "the real former client's phone number must never appear in generated output"


def test_passport_auth_pdf_produces_a_real_pdf(sample_applicant):
    skip_if_missing(ale.PASSPORT_SINGLE_TEMPLATE_PATH)
    skip_if_missing_binary("soffice")
    pdf_bytes = ale.generate_passport_authorization_pdf(
        [sample_applicant], "Test Centre", "Test Address", "Test Collector"
    )
    assert pdf_bytes[:5] == b"%PDF-"


# ---------------------------------------------------------------------------
# Company Authorization Letter
# ---------------------------------------------------------------------------


def test_company_auth_requires_at_least_one_applicant():
    with pytest.raises(ValueError):
        ale.generate_company_authorization_docx([], "Visa Officer, Some Consulate")


def test_company_auth_single_applicant(sample_applicant):
    skip_if_missing(ale.COMPANY_TEMPLATE_PATH)
    docx_bytes = ale.generate_company_authorization_docx([sample_applicant], "Visa Officer, Some Consulate")
    doc = Document(io.BytesIO(docx_bytes))
    text = _full_text(doc)
    assert sample_applicant["fullName"] in text
    assert sample_applicant["passportNumber"] in text
    assert "Visa Officer, Some Consulate" in text


def test_company_auth_clones_one_line_per_applicant(sample_applicant, sample_companion):
    skip_if_missing(ale.COMPANY_TEMPLATE_PATH)
    third = {"fullName": "Third Traveller", "passportNumber": "T0000003"}
    people = [sample_applicant, sample_companion, third]
    docx_bytes = ale.generate_company_authorization_docx(people, "Visa Officer, Some Consulate")
    doc = Document(io.BytesIO(docx_bytes))
    text = _full_text(doc)
    for person in people:
        assert person["fullName"].strip() in text
        assert person["passportNumber"] in text


def test_company_auth_pdf_produces_a_real_pdf(sample_applicant):
    skip_if_missing(ale.COMPANY_TEMPLATE_PATH)
    skip_if_missing_binary("soffice")
    pdf_bytes = ale.generate_company_authorization_pdf([sample_applicant], "Visa Officer, Some Consulate")
    assert pdf_bytes[:5] == b"%PDF-"


# ---------------------------------------------------------------------------
# Frontend rework, Phase 4 — Word-like body editor extended to both
# Authorization Letter types
# ---------------------------------------------------------------------------


def test_compose_company_editable_fields_matches_the_real_generated_text(sample_applicant, sample_companion):
    skip_if_missing(ale.COMPANY_TEMPLATE_PATH)
    people = [sample_applicant, sample_companion]
    fields = ale.compose_company_editable_fields(people, "Visa Officer, Some Consulate")
    assert sample_applicant["fullName"] in fields["applicant_0"]
    assert sample_applicant["passportNumber"] in fields["applicant_0"]
    assert sample_companion["fullName"] in fields["applicant_1"]
    assert "Visa Officer, Some Consulate" in fields["recipient"]
    assert "Date-" in fields["date"]


def test_company_auth_edited_applicant_line_replaces_only_that_line(sample_applicant, sample_companion):
    skip_if_missing(ale.COMPANY_TEMPLATE_PATH)
    people = [sample_applicant, sample_companion]
    docx_bytes = ale.generate_company_authorization_docx(
        people, "Visa Officer, Some Consulate",
        edited_fields={"applicant_0": "<p><strong>Edited</strong> Applicant Line</p>"},
    )
    doc = Document(io.BytesIO(docx_bytes))
    text = _full_text(doc)
    assert "Edited Applicant Line" in text
    assert sample_applicant["fullName"] not in text
    # Every other field (2nd applicant, recipient, date, fixed boilerplate,
    # signature) is untouched — still auto-generated.
    assert sample_companion["fullName"] in text
    assert sample_companion["passportNumber"] in text
    assert "Visa Officer, Some Consulate" in text
    assert "Dear Sir/Madam" in text
    assert "Ms. Dhvani Chheda" in text


def test_company_auth_edited_recipient_and_date_fields(sample_applicant):
    skip_if_missing(ale.COMPANY_TEMPLATE_PATH)
    docx_bytes = ale.generate_company_authorization_docx(
        [sample_applicant], "Visa Officer, Some Consulate",
        edited_fields={"recipient": "<p>Edited Recipient Text</p>", "date": "<p>Date- Edited Date Line</p>"},
    )
    doc = Document(io.BytesIO(docx_bytes))
    text = _full_text(doc)
    assert "Edited Recipient Text" in text
    assert "Visa Officer, Some Consulate" not in text
    assert "Date- Edited Date Line" in text
    # Applicant line (comes AFTER the edited date/recipient paragraphs)
    # must still be correct — the real regression this guards against.
    assert sample_applicant["fullName"] in text
    assert sample_applicant["passportNumber"] in text
    assert "Ms. Dhvani Chheda" in text


def test_compose_passport_editable_fields_matches_the_real_generated_text(sample_applicant):
    skip_if_missing(ale.PASSPORT_SINGLE_TEMPLATE_PATH)
    fields = ale.compose_passport_editable_fields(
        [sample_applicant], "Test Visa Application Centre", "123 Test Road, Test City", "Ms. Test Collector"
    )
    assert sample_applicant["fullName"] in fields["applicantSentence"]
    assert sample_applicant["passportNumber"] in fields["applicantSentence"]
    assert "Ms. Test Collector" in fields["collectorSentence"]
    assert "Test Visa Application Centre" in fields["recipientCentre"]
    assert "123 Test Road, Test City" in fields["recipientAddress"]
    assert "Date-" in fields["date"]
    # The signature/email block is deliberately not offered as a field.
    assert "signature" not in fields
    assert "email" not in fields


def test_passport_auth_edited_applicant_sentence_leaves_signature_correct(sample_applicant):
    """The applicant sentence (paragraph 15) sits BEFORE the signature/
    email block (paragraphs 25/26) — editing it must not corrupt the
    signature's real line break or the email's real mailto: hyperlink,
    the exact hazard this feature's own stable-paragraph-reference design
    exists to prevent."""
    skip_if_missing(ale.PASSPORT_SINGLE_TEMPLATE_PATH)
    docx_bytes = ale.generate_passport_authorization_docx(
        [sample_applicant], "Test Centre", "Test Address", "Test Collector",
        edited_fields={"applicantSentence": "<p>Custom authorization wording for this traveller.</p>"},
    )
    doc = Document(io.BytesIO(docx_bytes))
    text = _full_text(doc)
    assert "Custom authorization wording for this traveller." in text
    assert sample_applicant["fullName"] in text  # still in the signature block
    assert sample_applicant["email"] in text  # still in the (real, working) hyperlink line

    email_para = None
    for p in doc.paragraphs:
        if sample_applicant["email"] in p.text:
            email_para = p
    assert email_para is not None
    hyperlink_el = email_para._p.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink")
    assert hyperlink_el is not None, "the email line's real hyperlink must survive an edit to an earlier paragraph"


def test_passport_auth_edited_collector_sentence_and_multi_traveller_variant(sample_applicant, sample_companion):
    skip_if_missing(ale.PASSPORT_MULTIPLE_TEMPLATE_PATH)
    docx_bytes = ale.generate_passport_authorization_docx(
        [sample_applicant, sample_companion], "Test Centre", "Test Address", "Test Collector",
        edited_fields={"collectorSentence": "<p>Please release the passports to our custom collector.</p>"},
    )
    doc = Document(io.BytesIO(docx_bytes))
    text = _full_text(doc)
    assert "Please release the passports to our custom collector." in text
    assert "Test Collector" not in text
    assert sample_applicant["fullName"] in text
    assert sample_companion["fullName"] in text


def test_passport_auth_edited_date_and_recipient_fields_before_applicant_sentence(sample_applicant):
    """date (paragraph 4) and recipient centre/address (8/9) all sit
    BEFORE the applicant sentence (15) and signature (25/26) — editing
    them must not shift those later paragraphs' own content."""
    skip_if_missing(ale.PASSPORT_SINGLE_TEMPLATE_PATH)
    docx_bytes = ale.generate_passport_authorization_docx(
        [sample_applicant], "Test Centre", "Test Address", "Test Collector",
        edited_fields={
            "date": "<p>Date- A Custom Date Line</p>",
            "recipientCentre": "<p>Custom Centre Name</p>",
            "recipientAddress": "<p>Custom Address Line One</p><p>Custom Address Line Two</p>",
        },
    )
    doc = Document(io.BytesIO(docx_bytes))
    text = _full_text(doc)
    assert "Date- A Custom Date Line" in text
    assert "Custom Centre Name" in text
    assert "Custom Address Line One" in text
    assert "Custom Address Line Two" in text
    assert "Test Centre" not in text
    # Applicant sentence + signature (both AFTER the edited paragraphs)
    # must still be exactly correct.
    assert sample_applicant["fullName"] in text
    assert sample_applicant["passportNumber"] in text
    assert sample_applicant["email"] in text
