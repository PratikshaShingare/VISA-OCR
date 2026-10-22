"""
Tests for admin_export_engine.py — the Master Excel export (Applications /
Travellers sheets), built with openpyxl rather than from a
reference-template file (there is no .xlsx reference template for this one
— it's a fresh report, not a reproduction of an existing document format).
"""

import io

import openpyxl
import pytest

import admin_export_engine
import data_model


def _sample_application(status="Under Review", traveller_count=1):
    app = data_model.create_empty_application(created_by="qa@khannatravels.com")
    app["status"] = status  # hidden data field; must never appear in the export
    app["applicant"]["fullName"] = "Test Applicant One"
    app["applicant"]["passport"]["current"]["number"] = "T1111111"
    app["travel"]["destination"] = "Europe (Schengen)"
    app["travel"]["visaCategory"] = "Tourist"
    for i in range(traveller_count):
        app["travellers"].append(
            {
                "id": f"trav_{i}",
                "fullName": f"Test Traveller {i}",
                "passport": {"current": {"number": f"T200000{i}"}},
                "relation": "Spouse",
            }
        )
    return app


def test_build_applications_workbook_with_no_applications_still_produces_a_valid_workbook():
    content = admin_export_engine.build_applications_workbook([])
    assert content[:2] == b"PK", "an .xlsx is a real zip file"
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert len(wb.sheetnames) >= 1


def test_build_applications_workbook_has_the_two_expected_sheets_and_no_status_sheet():
    apps = [_sample_application()]
    content = admin_export_engine.build_applications_workbook(apps)
    wb = openpyxl.load_workbook(io.BytesIO(content))
    names = set(wb.sheetnames)
    assert any("Application" in n for n in names)
    assert any("Traveller" in n for n in names)
    assert not any("Status" in n for n in names), "application status was removed from the product"
    assert len(names) == 2


def test_applications_sheet_contains_real_application_data():
    apps = [_sample_application(status="Approved")]
    content = admin_export_engine.build_applications_workbook(apps)
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = next(ws for ws in wb.worksheets if "Application" in ws.title)
    all_values = [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None]
    assert any("Test Applicant One" in v for v in all_values)
    assert not any("Approved" in v for v in all_values), "status must not be exported"
    assert not any(str(cell.value) == "Status" for row in ws.iter_rows() for cell in row), "no Status column"
    assert any("T1111111" in v for v in all_values)


def test_travellers_sheet_lists_every_travellers_record():
    apps = [_sample_application(traveller_count=3)]
    content = admin_export_engine.build_applications_workbook(apps)
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = next(ws for ws in wb.worksheets if "Traveller" in ws.title)
    all_values = [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None]
    for i in range(3):
        assert any(f"Test Traveller {i}" in v for v in all_values)


def test_multiple_applications_all_appear_in_the_applications_sheet():
    apps = [_sample_application() for _ in range(5)]
    for i, app in enumerate(apps):
        app["id"] = f"APP_test_{i}"
        app["applicant"]["fullName"] = f"Applicant Number {i}"
    content = admin_export_engine.build_applications_workbook(apps)
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = next(ws for ws in wb.worksheets if "Application" in ws.title)
    all_values = [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None]
    for i in range(5):
        assert any(f"Applicant Number {i}" in v for v in all_values)
