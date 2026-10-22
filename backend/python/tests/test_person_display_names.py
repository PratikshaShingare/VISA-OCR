"""
Passport names are stored exactly as printed (capitals); only what is written
into a generated document is shown in ordinary capitalisation. Synthetic names.
"""

import pytest

import app as backend
import letter_shared


@pytest.mark.parametrize(
    "raw, expected",
    [
        ("MIHIR HARILAL KOTAK", "Mihir Harilal Kotak"),
        ("Mr. AMARDEEP DIVEKAR", "Mr. Amardeep Divekar"),
        ("O'BRIEN-SMITH", "O'Brien-Smith"),
        ("K. R. SHARMA", "K. R. Sharma"),
        ("Already Mixed Case", "Already Mixed Case"),
        ("McDonald", "McDonald"),  # mixed-case words are never touched
        ("", ""),
        ("  SPACED   NAME ", "Spaced Name"),
    ],
)
def test_person_display_name(raw, expected):
    assert letter_shared.person_display_name(raw) == expected


def test_none_is_left_alone():
    assert letter_shared.person_display_name(None) in (None, "")


def test_hotel_item_names_are_title_cased_including_each_guest():
    h = backend.HotelItem(leadGuestName="MR. TEST PERSON SAMPLE", guestNames=["MRS. TEST SPOUSE SAMPLE", "Mstr. Small Kid"])
    assert h.leadGuestName == "Mr. Test Person Sample"
    assert h.guestNames == ["Mrs. Test Spouse Sample", "Mstr. Small Kid"]


@pytest.mark.parametrize(
    "model",
    [backend.AuthPerson, backend.CompanyAuthPerson, backend.CoverLetterPerson, backend.InvitationLetterInviter],
)
def test_people_models_title_case_full_name_but_not_passport_number(model):
    p = model(fullName="TEST PERSON SAMPLE", passportNumber="X1234567")
    assert p.fullName == "Test Person Sample"
    assert p.passportNumber == "X1234567"


def test_company_authorization_endpoint_writes_readable_name_but_keeps_passport_number():
    import io

    from docx import Document
    from fastapi.testclient import TestClient

    client = TestClient(backend.app)
    resp = client.post(
        "/api/documents/company-authorization",
        json={
            "people": [{"fullName": "TEST PERSON SAMPLE", "passportNumber": "X1234567"}],
            "recipientText": "The Consulate General of Testland",
            "format": "docx",
        },
    )
    assert resp.status_code == 200
    doc = Document(io.BytesIO(resp.content))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "Test Person Sample" in text
    assert "TEST PERSON SAMPLE" not in text
    assert "X1234567" in text
