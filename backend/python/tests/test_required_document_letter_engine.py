"""
Tests for required_document_letter_engine.py, generating against the REAL
reference-templates/Required Document Letter/Required Document Letter.docx
template (project rule 5 — the reference template is the source of truth).

Rewritten for the checkbox-based, data-driven document-catalog rebuild
(visa_requirements_data.DOCUMENT_CATALOG) — the old fixed 4-conditional-
bullet API (invited=/has_us_visa_copy=) no longer exists; every document is
now selected via an explicit `document_ids` list resolved against that
catalog.
"""

import io

import pytest
from docx import Document

import required_document_letter_engine as rdl
import visa_requirements_data as vrd
from conftest import skip_if_missing, skip_if_missing_binary


def _full_text(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    return "\n".join(p.text for p in doc.paragraphs)


def test_requires_country():
    skip_if_missing(rdl.TEMPLATE_PATH)
    with pytest.raises(ValueError):
        rdl.generate_required_document_letter_docx(country="  ", visa_type="Tourist")


def test_rejects_unsupported_visa_type():
    skip_if_missing(rdl.TEMPLATE_PATH)
    with pytest.raises(ValueError):
        rdl.generate_required_document_letter_docx(country="Nepal", visa_type="Not A Real Visa Type")


def test_rejects_unsupported_employment_status():
    skip_if_missing(rdl.TEMPLATE_PATH)
    with pytest.raises(ValueError):
        rdl.generate_required_document_letter_docx(
            country="Nepal", visa_type="Tourist", employment_status="Not A Real Status"
        )


def test_every_expanded_visa_type_and_employment_status_is_accepted():
    """The 31-section rebuild spec explicitly expanded both lists — confirm
    every single one is actually wired through to the engine, not just
    declared in visa_requirements_data.py and silently unused."""
    skip_if_missing(rdl.TEMPLATE_PATH)
    for visa_type in vrd.VISA_TYPES:
        rdl.generate_required_document_letter_docx(country="Nepal", visa_type=visa_type)
    for status in vrd.EMPLOYMENT_STATUSES:
        rdl.generate_required_document_letter_docx(country="Nepal", visa_type="Tourist", employment_status=status)


def test_unknown_document_id_is_rejected_not_silently_dropped():
    """ValueError (not RequiredDocumentLetterError) — this is client-supplied
    bad input, mapped to a clean HTTP 400 by app.py, same as an invalid
    country/visa type/employment status. RequiredDocumentLetterError is
    reserved for the reference template itself being malformed."""
    skip_if_missing(rdl.TEMPLATE_PATH)
    with pytest.raises(ValueError):
        rdl.generate_required_document_letter_docx(
            country="Nepal", visa_type="Tourist", document_ids=["not_a_real_document_id"]
        )


def test_no_documents_selected_produces_an_honest_empty_checklist():
    skip_if_missing(rdl.TEMPLATE_PATH)
    text = _full_text(rdl.generate_required_document_letter_docx(country="Japan", visa_type="Tourist", document_ids=[]))
    assert "Japan" in text and "Tourist Visa" in text
    assert "Documents Required (Digital Copies):" in text
    # No bullet lines at all — not a single leftover fixed-template phrase.
    for phrase in ("Passport Copy", "Aadhar Card", "If Employed", "If Self-Employed", "If Invited", "Valid USA Visa Copy"):
        assert phrase not in text


def test_only_selected_documents_appear_in_canonical_order():
    """Selecting a scattered subset of ids must produce exactly those
    documents, in DOCUMENT_ORDER (not selection order, not dict order) —
    the core "checkbox-based, only selected documents appear" requirement."""
    skip_if_missing(rdl.TEMPLATE_PATH)
    selected = ["us_visa_copy", "passport_copy", "medical_certificate"]  # deliberately out of canonical order
    text = _full_text(rdl.generate_required_document_letter_docx(country="France", visa_type="Business", document_ids=selected))
    assert "- Passport Copy (First and Last Page)" in text
    assert "- Medical Certificate / Fitness Report from Registered Physician" in text
    assert "- Valid USA Visa Copy (if any)" in text
    # Never-selected catalog documents must not leak in.
    assert "Aadhar Card" not in text
    assert "Salary Slips" not in text
    passport_idx = text.index("Passport Copy")
    medical_idx = text.index("Medical Certificate")
    usvisa_idx = text.index("Valid USA Visa Copy")
    assert passport_idx < medical_idx < usvisa_idx, "documents must render in canonical DOCUMENT_ORDER, not selection order"


def test_bank_statement_and_itr_placeholders_substitute_into_their_own_bullets():
    skip_if_missing(rdl.TEMPLATE_PATH)
    text = _full_text(rdl.generate_required_document_letter_docx(
        country="United Kingdom", visa_type="Tourist",
        document_ids=["bank_statement", "itr"],
        bank_statement_duration="6 Months", itr_years="2 Years",
    ))
    assert "- Personal Bank Statement of Last 6 Months (with sufficient balance)" in text
    assert "- Personal Income Tax Returns of Last 2 Years (Acknowledgment Page Only)" in text


def test_bank_statement_and_itr_fall_back_to_honest_placeholder_when_blank():
    skip_if_missing(rdl.TEMPLATE_PATH)
    text = _full_text(rdl.generate_required_document_letter_docx(
        country="United Kingdom", visa_type="Tourist", document_ids=["bank_statement", "itr"],
    ))
    assert "[DURATION]" in text
    assert "Last __ (Acknowledgment" in text


def test_processing_time_and_visa_fees_fill_the_real_blanks_without_fabricating_a_unit_clash():
    skip_if_missing(rdl.TEMPLATE_PATH)
    text = _full_text(rdl.generate_required_document_letter_docx(
        country="United Kingdom", visa_type="Tourist",
        processing_time="2-3 Weeks", visa_fees="8,500",
    ))
    assert "Processing Time: 2-3 Weeks (Subject to Consulate Approval)" in text
    assert "Weeks Months" not in text, "the template's own fixed 'Months' word must not clash with a supplied unit"
    assert "INR 8,500 /- pp" in text


def test_unfilled_processing_time_and_fees_show_honest_placeholders():
    skip_if_missing(rdl.TEMPLATE_PATH)
    text = _full_text(rdl.generate_required_document_letter_docx(country="Nepal", visa_type="Tourist"))
    assert "Processing Time: __ Months" in text
    assert "INR _____" in text


def test_label_bold_formatting_is_preserved_value_is_not():
    """The template's 'Country:' label is bold and the value after it is
    not — a document-wide token sweep would incorrectly bold the whole
    line (see this module's own docstring). Confirms the per-run fill
    approach keeps them visually distinct, matching the real template."""
    skip_if_missing(rdl.TEMPLATE_PATH)
    docx_bytes = rdl.generate_required_document_letter_docx(country="Kenya", visa_type="Tourist")
    doc = Document(io.BytesIO(docx_bytes))
    country_para = next(p for p in doc.paragraphs if p.text.startswith("Country:"))
    assert country_para.runs[0].bold is True
    assert country_para.runs[0].text == "Country"
    assert not any(r.bold for r in country_para.runs[1:])


def test_dynamic_bullet_lines_preserve_the_template_bullet_formatting():
    """Every cloned bullet line must inherit the real template's own plain
    (non-bold) run formatting — a document-wide rebuild that lost per-run
    fidelity would silently bold or resize the checklist."""
    skip_if_missing(rdl.TEMPLATE_PATH)
    docx_bytes = rdl.generate_required_document_letter_docx(
        country="Kenya", visa_type="Tourist", document_ids=["passport_copy", "aadhar_card", "photographs"],
    )
    doc = Document(io.BytesIO(docx_bytes))
    bullet_paragraphs = [p for p in doc.paragraphs if p.text.startswith("- ")]
    assert len(bullet_paragraphs) == 3
    for p in bullet_paragraphs:
        assert not any(r.bold for r in p.runs), f"bullet line should not be bold: {p.text!r}"


def test_generate_required_document_letter_pdf_produces_a_real_pdf():
    skip_if_missing(rdl.TEMPLATE_PATH)
    skip_if_missing_binary("soffice")
    pdf_bytes = rdl.generate_required_document_letter_pdf(country="Nepal", visa_type="Tourist")
    assert pdf_bytes[:5] == b"%PDF-"


# ---------------------------------------------------------------------------
# visa_requirements_data.py's own rule engine (DOCUMENT_CATALOG /
# resolve_suggested_document_ids / resolve_document_label) — the
# "structured, data-driven configuration... not hundreds of if/else" the
# user explicitly asked for.
# ---------------------------------------------------------------------------


def test_document_catalog_and_order_are_consistent():
    assert set(vrd.DOCUMENT_ORDER) == set(vrd.DOCUMENT_CATALOG)
    assert len(vrd.DOCUMENT_ORDER) == len(set(vrd.DOCUMENT_ORDER)), "no duplicate ids in DOCUMENT_ORDER"


def test_always_documents_are_suggested_regardless_of_visa_type_or_status():
    always_ids = {doc_id for doc_id, entry in vrd.DOCUMENT_CATALOG.items() if entry.get("always")}
    suggested = set(vrd.resolve_suggested_document_ids("Transit", "Minor", []))
    assert always_ids.issubset(suggested)


def test_employment_status_conditional_documents_are_suggested_only_for_their_own_status():
    employed = set(vrd.resolve_suggested_document_ids("Tourist", "Employed", []))
    self_employed = set(vrd.resolve_suggested_document_ids("Tourist", "Self-Employed", []))
    assert "salary_slips" in employed
    assert "salary_slips" not in self_employed
    assert "company_registration_docs" in self_employed
    assert "company_registration_docs" not in employed


def test_visa_type_conditional_documents_are_suggested_only_for_their_own_visa_types():
    student = set(vrd.resolve_suggested_document_ids("Student", "Other", []))
    tourist = set(vrd.resolve_suggested_document_ids("Tourist", "Other", []))
    assert "admission_letter" in student
    assert "admission_letter" not in tourist


def test_circumstance_flags_add_their_own_documents():
    plain = set(vrd.resolve_suggested_document_ids("Tourist", "Other", []))
    invited = set(vrd.resolve_suggested_document_ids("Tourist", "Other", ["invited"]))
    assert "invitation_docs" not in plain
    assert "invitation_docs" in invited


def test_resolve_document_label_substitutes_dynamic_values():
    assert vrd.resolve_document_label("bank_statement", "12 Months", "") == (
        "Personal Bank Statement of Last 12 Months (with sufficient balance)"
    )
    assert vrd.resolve_document_label("itr", "", "5 Years") == (
        "Personal Income Tax Returns of Last 5 Years (Acknowledgment Page Only)"
    )
    # Static (non-templated) documents ignore the duration/years args entirely.
    assert vrd.resolve_document_label("passport_copy", "12 Months", "5 Years") == "Passport Copy (First and Last Page)"


def test_resolve_document_label_rejects_unknown_id():
    with pytest.raises(KeyError):
        vrd.resolve_document_label("not_a_real_id")
