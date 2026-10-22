"""
Unit tests for docx_utils.py — the shared, formatting-preserving .docx
helpers every document-generation engine builds on.

These tests build small SYNTHETIC .docx documents on the fly with
python-docx (not the project's real reference templates — see
test_hotel_voucher_engine.py / test_authorization_letter_engine.py /
test_cover_letter_engine.py / test_invitation_letter_engine.py for tests
against the real templates). That is deliberate here: these are unit tests
of the low-level mechanics themselves (run-preserving text replacement,
`<w:br/>`-aware two-line paragraphs, hyperlink rewriting, element cloning),
so a minimal, purpose-built document makes the exact structure under test
unambiguous.
"""

import io

import pytest
from docx import Document
from docx.oxml.ns import qn

import docx_utils
from conftest import skip_if_missing_binary


def _new_doc_with_paragraph(text: str = "Hello world"):
    doc = Document()
    doc.add_paragraph(text)
    return doc


def test_set_paragraph_text_preserves_first_run_formatting():
    doc = _new_doc_with_paragraph("Original text")
    para = doc.paragraphs[0]
    para.runs[0].bold = True
    docx_utils.set_paragraph_text(para, "Replaced text")
    assert para.text == "Replaced text"
    assert para.runs[0].bold is True


def test_set_paragraph_text_blanks_extra_runs_without_removing_them():
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Hello ")
    para.add_run("World")
    assert len(para.runs) == 2
    docx_utils.set_paragraph_text(para, "New")
    assert para.text == "New"
    assert len(para.runs) == 2, "run count must stay the same — never delete a run"
    assert para.runs[1].text == ""


def test_set_cell_text():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run("old")
    docx_utils.set_cell_text(table.cell(0, 0), "new value")
    assert table.cell(0, 0).text == "new value"


def test_replace_token_in_paragraph_found_and_not_found():
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Dear [Name], welcome.")
    found = docx_utils.replace_token_in_paragraph(para, "[Name]", "Rohan")
    assert found is True
    assert para.text == "Dear Rohan, welcome."

    found_again = docx_utils.replace_token_in_paragraph(para, "[Name]", "Someone Else")
    assert found_again is False, "token no longer present after the first replacement"


def test_replace_token_in_paragraph_merges_a_token_split_across_runs():
    doc = Document()
    para = doc.add_paragraph()
    # Simulate Word having split a single bracket token across three runs —
    # a real-world authoring artifact this helper exists specifically to
    # survive (see its own docstring).
    para.add_run("Dear [Na")
    para.add_run("me")
    para.add_run("], welcome.")
    found = docx_utils.replace_token_in_paragraph(para, "[Name]", "Priya")
    assert found is True
    assert para.text == "Dear Priya, welcome."


def test_apply_token_map_to_paragraph_replaces_every_pair():
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("[A] and [B] and [A] again")
    docx_utils.apply_token_map_to_paragraph(para, {"[A]": "x", "[B]": "y"})
    assert para.text == "x and y and x again"


def test_replace_token_in_cell_scans_every_paragraph():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.paragraphs[0].add_run("first paragraph")
    cell.add_paragraph("second paragraph with [Token] inside")
    found = docx_utils.replace_token_in_cell(cell, "[Token]", "REPLACED")
    assert found is True
    assert "REPLACED" in cell.text


def _add_run_with_break(paragraph, before_text: str, after_text: str):
    """Builds a paragraph with three runs — text, a real <w:br/>, more text
    — matching the exact shape docx_utils.set_two_line_paragraph documents
    itself as handling (verified against the real Passport Authorization
    Letter templates)."""
    r1 = paragraph.add_run(before_text)
    r_break = paragraph.add_run()
    br = r_break._element.makeelement(qn("w:br"), {})
    r_break._element.append(br)
    paragraph.add_run(after_text)
    return r1, r_break


def test_has_break_run_true_and_false():
    doc = Document()
    p_with_break = doc.add_paragraph()
    _add_run_with_break(p_with_break, "line1", "line2")
    assert docx_utils.has_break_run(p_with_break) is True

    p_without = doc.add_paragraph("plain text, no break")
    assert docx_utils.has_break_run(p_without) is False


def test_set_two_line_paragraph_preserves_the_break_and_sets_both_lines():
    doc = Document()
    para = doc.add_paragraph()
    _add_run_with_break(para, "Old Name", "Old Phone")
    docx_utils.set_two_line_paragraph(para, "Rohan Mehta", "Phone No.: +91 9876500001")

    assert docx_utils.has_break_run(para), "the <w:br/> element itself must survive"
    full_text_via_xml = "".join(t.text or "" for t in para._p.findall(f".//{qn('w:t')}"))
    assert "Rohan Mehta" in full_text_via_xml
    assert "Phone No.: +91 9876500001" in full_text_via_xml


def test_set_two_line_paragraph_falls_back_gracefully_without_a_break_run():
    doc = Document()
    para = doc.add_paragraph("single run, no break at all")
    docx_utils.set_two_line_paragraph(para, "Line One", "Line Two")
    assert "Line One" in para.text
    assert "Line Two" in para.text


def _add_hyperlink(paragraph, document, text: str, target: str):
    part = document.part
    r_id = part.relate_to(target, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = paragraph._p.makeelement(qn("w:hyperlink"), {qn("r:id"): r_id})
    run = paragraph.add_run(text)
    run_el = run._element
    paragraph._p.remove(run_el)
    hyperlink.append(run_el)
    paragraph._p.append(hyperlink)
    return r_id


def test_find_hyperlink_text_run_and_update_text_and_target():
    doc = Document()
    para = doc.add_paragraph("Email id: ")
    _add_hyperlink(para, doc, "old@example.com", "mailto:old@example.com")

    found = docx_utils.find_hyperlink_text_run(para)
    assert found is not None
    _hyperlink_el, r_id, t_el = found
    assert t_el.text == "old@example.com"

    updated = docx_utils.set_hyperlink_text_and_target(para, doc, "new@example.com", "mailto:new@example.com")
    assert updated is True
    _hyperlink_el2, r_id2, t_el2 = docx_utils.find_hyperlink_text_run(para)
    assert t_el2.text == "new@example.com"


def test_set_hyperlink_text_and_target_returns_false_when_no_hyperlink():
    doc = Document()
    para = doc.add_paragraph("plain paragraph, no hyperlink")
    assert docx_utils.set_hyperlink_text_and_target(para, doc, "irrelevant") is False


def test_remove_hyperlinks_in_paragraph():
    doc = Document()
    para = doc.add_paragraph("Email Id: ")
    _add_hyperlink(para, doc, "leftover@example.com", "mailto:leftover@example.com")
    assert docx_utils.find_hyperlink_text_run(para) is not None

    removed = docx_utils.remove_hyperlinks_in_paragraph(para)
    assert removed == 1
    assert docx_utils.find_hyperlink_text_run(para) is None


def test_insert_signature_image_with_a_real_image(tmp_path):
    from PIL import Image

    img_path = tmp_path / "sig.png"
    Image.new("RGB", (100, 40), color="white").save(img_path)

    doc = Document()
    para = doc.add_paragraph()
    docx_utils.insert_signature_image(para, img_path.read_bytes())
    assert len(para.runs) == 1
    # A picture-bearing run contains a <w:drawing> element, not text.
    assert para.runs[0]._element.find(qn("w:drawing")) is not None


def test_insert_signature_image_raises_on_garbage_bytes():
    doc = Document()
    para = doc.add_paragraph()
    with pytest.raises(docx_utils.DocxBuildError):
        docx_utils.insert_signature_image(para, b"not a real image")


def test_delete_paragraph_removes_it_from_the_document():
    doc = Document()
    doc.add_paragraph("keep me")
    doc.add_paragraph("delete me")
    assert len(doc.paragraphs) == 2
    docx_utils.delete_paragraph(doc.paragraphs[1])
    assert len(doc.paragraphs) == 1
    assert doc.paragraphs[0].text == "keep me"


def test_clone_and_insert_after_duplicates_a_table_row_n_times():
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    template_row = table.rows[1]
    template_row.cells[0].text = "template"

    new_elements = docx_utils.clone_and_insert_after(template_row._tr, 3)
    assert len(new_elements) == 3
    assert len(table.rows) == 5, "1 header + 1 template + 3 clones = 5 rows"
    assert all(r.cells[0].text == "template" for r in table.rows[1:])


def test_convert_docx_bytes_to_pdf_produces_a_real_pdf():
    skip_if_missing_binary("soffice")
    doc = Document()
    doc.add_paragraph("PDF conversion smoke test")
    buf = io.BytesIO()
    doc.save(buf)
    pdf_bytes = docx_utils.convert_docx_bytes_to_pdf(buf.getvalue(), base_name="smoke_test")
    assert pdf_bytes[:5] == b"%PDF-", "output must be a real PDF, not a stub"
    assert len(pdf_bytes) > 200


# ---------------------------------------------------------------------------
# Phase 4 proof of concept — Quill HTML parsing / paragraph-range replace
# ---------------------------------------------------------------------------


def test_parse_editor_html_plain_paragraphs():
    blocks = docx_utils.parse_editor_html("<p>First.</p><p>Second.</p>")
    assert [b["align"] for b in blocks] == [None, None]
    assert [b["runs"][0]["text"] for b in blocks] == ["First.", "Second."]


def test_parse_editor_html_alignment_classes():
    html = (
        '<p class="ql-align-center">Centered.</p>'
        '<p class="ql-align-right">Righted.</p>'
        '<p class="ql-align-justify">Justified.</p>'
        "<p>Default.</p>"
    )
    blocks = docx_utils.parse_editor_html(html)
    assert [b["align"] for b in blocks] == ["center", "right", "both", None]


def test_parse_editor_html_bold_italic_underline():
    html = "<p>Plain <strong>bold</strong> <em>italic</em> <u>under</u> text.</p>"
    blocks = docx_utils.parse_editor_html(html)
    runs = blocks[0]["runs"]
    by_text = {r["text"].strip(): r for r in runs}
    assert by_text["bold"]["bold"] is True
    assert by_text["italic"]["italic"] is True
    assert by_text["under"]["underline"] is True
    assert by_text["Plain"]["bold"] is False


def test_parse_editor_html_lists_carry_type_and_level():
    html = (
        '<ol><li data-list="bullet"><span class="ql-ui" contenteditable="false"></span>One</li>'
        '<li data-list="bullet" class="ql-indent-1"><span class="ql-ui"></span>Two</li>'
        '<li data-list="ordered"><span class="ql-ui"></span>First</li></ol>'
        "<ul><li>Legacy</li></ul>"
    )
    blocks = docx_utils.parse_editor_html(html)
    assert [("".join(r["text"] for r in b["runs"]), b["list"]) for b in blocks] == [
        ("One", {"type": "bullet", "level": 0}),
        ("Two", {"type": "bullet", "level": 1}),
        ("First", {"type": "ordered", "level": 0}),
        ("Legacy", {"type": "bullet", "level": 0}),
    ]


def test_parse_editor_html_block_formats_and_run_formats():
    html = (
        '<p class="ql-indent-2" style="line-height: 1.5; margin-top: 12pt; margin-bottom: 0pt;">'
        "<s>gone</s> H<sub>2</sub>O x<sup>2</sup></p>"
    )
    (block,) = docx_utils.parse_editor_html(html)
    assert block["indent"] == 2 and block["line_height"] == 1.5
    assert block["space_before"] == 12.0 and block["space_after"] == 0.0
    flags = {r["text"].strip(): r for r in block["runs"] if r["text"].strip()}
    assert flags["gone"]["strike"] and flags["2"]["sub"] or flags["2"]["sup"]
    assert any(r["sub"] for r in block["runs"]) and any(r["sup"] for r in block["runs"])


def test_parse_editor_html_tables():
    html = (
        '<table><tbody><tr><td data-row="r1">A1</td><td data-row="r1"><strong>B1</strong></td></tr>'
        '<tr><td data-row="r2">A2</td><td data-row="r2"><br></td></tr></tbody></table><p>after</p>'
    )
    blocks = docx_utils.parse_editor_html(html)
    assert [b["kind"] for b in blocks] == ["table", "p"]
    rows = blocks[0]["rows"]
    assert len(rows) == 2 and len(rows[0]) == 2
    assert rows[0][1]["blocks"][0]["runs"][0]["bold"] is True
    assert rows[1][1]["blocks"][0]["runs"] == [] or not rows[1][1]["blocks"][0]["runs"][0]["text"].strip()


def test_parse_editor_html_drops_blank_blocks():
    blocks = docx_utils.parse_editor_html("<p><br></p><p>Real text.</p><p></p>")
    assert len(blocks) == 1
    assert blocks[0]["runs"][0]["text"] == "Real text."


def _new_doc_with_body_range():
    """A synthetic 5-paragraph document standing in for the Europe Cover
    Letter's real 'paragraph before the body / 3 body paragraphs /
    paragraph after the body' shape, so replace_paragraph_range_with_html
    can be exercised without depending on the real reference template."""
    doc = Document()
    doc.add_paragraph("Before the body.")
    doc.add_paragraph("Body para one.")
    doc.add_paragraph("Body para two.")
    doc.add_paragraph("Body para three.")
    doc.add_paragraph("After the body.")
    return doc


def test_replace_paragraph_range_with_html_replaces_only_the_given_range():
    doc = _new_doc_with_body_range()
    paragraphs = doc.paragraphs
    first, last = paragraphs[1], paragraphs[3]

    docx_utils.replace_paragraph_range_with_html(
        doc, first, last, "<p>New one.</p><p>New two.</p>"
    )

    texts = [p.text for p in doc.paragraphs]
    assert texts == ["Before the body.", "New one.", "New two.", "After the body."]


def test_replace_paragraph_range_with_html_preserves_bold_italic_underline_and_alignment():
    doc = _new_doc_with_body_range()
    paragraphs = doc.paragraphs
    first, last = paragraphs[1], paragraphs[3]

    html = (
        '<p class="ql-align-center">Plain <strong>bold</strong> and <em>italic</em> and <u>under</u>.</p>'
    )
    new_paragraphs = docx_utils.replace_paragraph_range_with_html(doc, first, last, html)

    assert len(new_paragraphs) == 1
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    assert new_paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    runs = new_paragraphs[0].runs
    by_text = {r.text.strip(): r for r in runs if r.text.strip()}
    assert by_text["bold"].bold is True
    assert by_text["italic"].italic is True
    assert by_text["under"].underline is True


def test_replace_paragraph_range_with_html_always_leaves_at_least_one_paragraph():
    doc = _new_doc_with_body_range()
    paragraphs = doc.paragraphs
    first, last = paragraphs[1], paragraphs[3]

    docx_utils.replace_paragraph_range_with_html(doc, first, last, "")

    texts = [p.text for p in doc.paragraphs]
    # 5 originally, 3-paragraph range collapsed to exactly 1 blank paragraph.
    assert len(texts) == 3
    assert texts == ["Before the body.", "", "After the body."]


def test_replace_paragraph_range_with_html_raises_when_last_precedes_first():
    doc = Document()
    doc.add_paragraph("A")
    doc.add_paragraph("B")
    doc.add_paragraph("C")
    paragraphs = doc.paragraphs
    # last_paragraph (A) is NOT reachable walking forward from
    # first_paragraph (C) — an honest error, not a silent no-op or a
    # corrupted document.
    with pytest.raises(docx_utils.DocxBuildError):
        docx_utils.replace_paragraph_range_with_html(doc, paragraphs[2], paragraphs[0], "<p>x</p>")


def test_parse_editor_html_reads_inline_style_colour_highlight_size_font():
    html = (
        '<p><strong style="font-size: 20pt; color: rgb(230, 0, 0); background-color: #ffff00; '
        'font-family: Arial;">Hi</strong> plain</p>'
    )
    runs = docx_utils.parse_editor_html(html)[0]["runs"]
    assert runs[0]["text"] == "Hi"
    assert runs[0]["bold"] is True
    assert runs[0]["color"] == "E60000"
    assert runs[0]["highlight"] == "FFFF00"
    assert runs[0]["size_pt"] == 20.0
    assert runs[0]["font"] == "Arial"
    assert runs[1]["text"] == " plain" and runs[1]["bold"] is False and runs[1]["color"] is None


def test_parse_editor_html_ignores_unusable_style_values():
    runs = docx_utils.parse_editor_html('<p><span style="color: banana; font-size: 900pt">x</span></p>')[0]["runs"]
    assert runs[0]["color"] is None and runs[0]["size_pt"] is None


# ---------------------------------------------------------------------------
# Real Word structures written from the editor's HTML
# ---------------------------------------------------------------------------


def _fresh_doc():
    doc = Document()
    doc.add_paragraph("Before")
    target = doc.add_paragraph("Replace me")
    doc.add_paragraph("After")
    return doc, target


def test_writer_creates_real_bulleted_and_numbered_lists():
    doc, target = _fresh_doc()
    html = (
        '<ol><li data-list="bullet"><span class="ql-ui"></span>b1</li><li data-list="bullet" class="ql-indent-1">'
        '<span class="ql-ui"></span>b2</li><li data-list="ordered"><span class="ql-ui"></span>n1</li>'
        '<li data-list="ordered"><span class="ql-ui"></span>n2</li></ol><p>gap</p>'
        '<ol><li data-list="ordered"><span class="ql-ui"></span>again</li></ol>'
    )
    paras = docx_utils.replace_paragraph_range_with_html(doc, target, target, html)
    num_ids = []
    for p in paras:
        num_pr = p._p.pPr.find(qn("w:numPr")) if p._p.pPr is not None else None
        num_ids.append(None if num_pr is None else (num_pr.find(qn("w:numId")).get(qn("w:val")), num_pr.find(qn("w:ilvl")).get(qn("w:val"))))
    assert num_ids[0][0] == num_ids[1][0] and num_ids[0][1] == "0" and num_ids[1][1] == "1"
    assert num_ids[2][0] == num_ids[3][0] != num_ids[0][0]
    assert num_ids[4] is None
    assert num_ids[5][0] not in (num_ids[2][0], num_ids[0][0]), "a second numbered list restarts at 1"
    # The numbering part really exists and defines those ids.
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    numbering = doc.part.part_related_by(RT.NUMBERING).element
    defined = {n.get(qn("w:numId")) for n in numbering.findall(qn("w:num"))}
    assert {n[0] for n in num_ids if n} <= defined
    # and it survives a save/reload
    buf = io.BytesIO()
    doc.save(buf)
    assert Document(io.BytesIO(buf.getvalue())).paragraphs[1].text == "b1"


def test_writer_applies_indent_line_height_and_spacing():
    from docx.shared import Pt

    doc, target = _fresh_doc()
    (para,) = docx_utils.replace_paragraph_range_with_html(
        doc, target, target,
        '<p class="ql-indent-2" style="line-height: 1.5; margin-top: 12pt; margin-bottom: 0pt;">text</p>',
    )
    fmt = para.paragraph_format
    assert abs(fmt.left_indent.inches - 1.0) < 0.01
    assert fmt.line_spacing == 1.5
    assert fmt.space_before == Pt(12) and fmt.space_after == Pt(0)


def test_writer_applies_strike_sub_sup_and_word_highlight():
    from docx.enum.text import WD_COLOR_INDEX

    doc, target = _fresh_doc()
    (para,) = docx_utils.replace_paragraph_range_with_html(
        doc, target, target,
        '<p><s>a</s><sub>b</sub><sup>c</sup><span style="background-color: #00ffff">d</span></p>',
    )
    by = {r.text: r for r in para.runs}
    assert by["a"].font.strike is True
    assert by["b"].font.subscript is True
    assert by["c"].font.superscript is True
    assert by["d"].font.highlight_color == WD_COLOR_INDEX.TURQUOISE


def test_writer_inserts_a_real_table_between_paragraphs():
    doc, target = _fresh_doc()
    html = (
        '<p>intro</p><table><tbody><tr><td data-row="r1">A1</td><td data-row="r1">B1</td></tr>'
        '<tr><td data-row="r2">A2</td><td data-row="r2">B2</td></tr></tbody></table><p>outro</p>'
    )
    docx_utils.replace_paragraph_range_with_html(doc, target, target, html)
    assert len(doc.tables) == 1
    t = doc.tables[0]
    assert [[c.text for c in r.cells] for r in t.rows] == [["A1", "B1"], ["A2", "B2"]]
    body = [c.tag.split("}")[1] for c in doc.element.body]
    assert body[:5] == ["p", "p", "tbl", "p", "p"]  # Before, intro, table, outro, After(+sectPr)


def test_writer_keeps_a_cell_valid_when_a_table_replaces_its_only_paragraph():
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    cell = outer.cell(0, 0)
    target = cell.paragraphs[0]
    target.add_run("only")
    docx_utils.replace_paragraph_range_with_html(
        doc, target, target, '<table><tbody><tr><td data-row="r1">x</td></tr></tbody></table>'
    )
    assert cell._tc[-1].tag == qn("w:p"), "a table cell must end with a paragraph"
