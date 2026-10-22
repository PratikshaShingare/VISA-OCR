"""
Tests for invitation_letter_engine.py — the Invitation Letter (from an
inviter abroad, for exactly two invitees — the reference template's own
fixed "my Parents" wording) and the Initors Covering Letter (a real,
gendered-voice letter: the wife's-own-voice template vs. the husband's-own
-voice template), generated against the REAL
reference-templates/Invitation Letter/ and
reference-templates/Invitors letters/ files (project rule 5).
"""

import io

import pytest
from docx import Document
from PIL import Image

import invitation_letter_engine as ile
from conftest import skip_if_missing, skip_if_missing_binary


def _full_text(doc: Document) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


INVITER = {
    "fullName": "Test Inviter Abroad",
    "addressLine1": "1 Test Street",
    "addressLine2": "Apt 2",
    "cityPostcode": "London, W1 1AA",
    "country": "United Kingdom",
    "cityCountry": "London, United Kingdom",
    "passportNumber": "UKT000001",
    "fullAddress": "1 Test Street, Apt 2, London, W1 1AA, United Kingdom",
    "studyingOrWorking": "Working",
    "universityOrCompany": "Test University",
    "phone": "9876500009",
    "email": "inviter.test@example.com",
}

INVITATION_FIELDS = {
    "travelStartDate": "2026-08-01",
    "travelEndDate": "2026-08-15",
    "purpose": "Tourism",
    "accommodationDetails": "Staying with the inviter",
    "returnDate": "2026-08-16",
    "fundingArrangement": "Inviter",
}


# ---------------------------------------------------------------------------
# Invitation Letter
# ---------------------------------------------------------------------------


def test_invitation_letter_requires_exactly_two_invitees(sample_applicant, sample_companion):
    with pytest.raises(ValueError):
        ile.generate_invitation_letter_docx(INVITER, [sample_applicant], INVITATION_FIELDS)
    with pytest.raises(ValueError):
        third = dict(sample_applicant, fullName="Third")
        ile.generate_invitation_letter_docx(INVITER, [sample_applicant, sample_companion, third], INVITATION_FIELDS)


def test_invitation_letter_requires_inviter_name():
    with pytest.raises(ValueError):
        ile.generate_invitation_letter_docx({}, [{"fullName": "A"}, {"fullName": "B"}], INVITATION_FIELDS)


def test_invitation_letter_fills_real_template(sample_applicant, sample_companion):
    skip_if_missing(ile.INVITATION_LETTER_TEMPLATE_PATH)
    docx_bytes = ile.generate_invitation_letter_docx(INVITER, [sample_applicant, sample_companion], INVITATION_FIELDS)
    assert docx_bytes[:2] == b"PK"
    text = _full_text(Document(io.BytesIO(docx_bytes)))
    assert INVITER["fullName"] in text
    assert sample_applicant["fullName"] in text
    assert sample_companion["fullName"] in text
    assert INVITATION_FIELDS["purpose"] in text
    assert INVITER["email"] in text


def test_invitation_letter_embeds_a_real_uploaded_signature(sample_applicant, sample_companion, tmp_path):
    skip_if_missing(ile.INVITATION_LETTER_TEMPLATE_PATH)
    img_path = tmp_path / "sig.png"
    Image.new("RGB", (100, 40), color="white").save(img_path)
    docx_bytes = ile.generate_invitation_letter_docx(
        INVITER, [sample_applicant, sample_companion], INVITATION_FIELDS, signature_bytes=img_path.read_bytes()
    )
    doc = Document(io.BytesIO(docx_bytes))
    has_drawing = any(
        r._element.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing") is not None
        for p in doc.paragraphs
        for r in p.runs
    )
    assert has_drawing, "a real uploaded signature image must actually be embedded, not merely accepted"


def test_invitation_letter_without_a_signature_leaves_the_slot_blank(sample_applicant, sample_companion):
    """Project rule 9 — never fabricate a signature: when none is
    uploaded, the template's own blank signature paragraph must stay
    exactly as blank as it started."""
    skip_if_missing(ile.INVITATION_LETTER_TEMPLATE_PATH)
    docx_bytes = ile.generate_invitation_letter_docx(INVITER, [sample_applicant, sample_companion], INVITATION_FIELDS)
    doc = Document(io.BytesIO(docx_bytes))
    has_drawing = any(
        r._element.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing") is not None
        for p in doc.paragraphs
        for r in p.runs
    )
    assert not has_drawing


def test_invitation_letter_pdf(sample_applicant, sample_companion):
    skip_if_missing(ile.INVITATION_LETTER_TEMPLATE_PATH)
    skip_if_missing_binary("soffice")
    pdf_bytes = ile.generate_invitation_letter_pdf(INVITER, [sample_applicant, sample_companion], INVITATION_FIELDS)
    assert pdf_bytes[:5] == b"%PDF-"


# ---------------------------------------------------------------------------
# Frontend rework, Phase 4: Word-like body editor editable fields
# ---------------------------------------------------------------------------


def test_compose_invitation_editable_fields_matches_the_real_generated_text(sample_applicant, sample_companion):
    """Every composed field must be the exact text the real (non-edited)
    generator would have produced — project rule 15's zero-drift-risk
    guarantee, verified the same way the Authorization Letters' own
    compose_*_editable_fields tests are."""
    skip_if_missing(ile.INVITATION_LETTER_TEMPLATE_PATH)
    fields = ile.compose_invitation_editable_fields(INVITER, [sample_applicant, sample_companion], INVITATION_FIELDS)
    assert set(fields.keys()) == set(ile.INVITATION_EDITABLE_FIELD_KEYS)

    docx_bytes = ile.generate_invitation_letter_docx(INVITER, [sample_applicant, sample_companion], INVITATION_FIELDS)
    doc = Document(io.BytesIO(docx_bytes))
    for key, idx in ile._INVITATION_FIELD_PARAGRAPH_INDEX.items():
        composed_text = fields[key].removeprefix("<p>").removesuffix("</p>")
        assert composed_text == doc.paragraphs[idx].text, f"field {key!r} drifted from the real generated paragraph"


def test_invitation_letter_edited_field_replaces_only_that_paragraph(sample_applicant, sample_companion):
    skip_if_missing(ile.INVITATION_LETTER_TEMPLATE_PATH)
    edited = {"purposeStatement": "<p>Edited purpose sentence via the Word-like editor.</p>"}
    docx_bytes = ile.generate_invitation_letter_docx(
        INVITER, [sample_applicant, sample_companion], INVITATION_FIELDS, edited_fields=edited
    )
    text = _full_text(Document(io.BytesIO(docx_bytes)))
    assert "Edited purpose sentence via the Word-like editor." in text
    assert "It would be an incredibly special and proud moment" not in text
    # Every OTHER field must still auto-generate normally.
    assert INVITER["fullName"] in text
    assert sample_applicant["fullName"] in text
    assert INVITER["email"] in text


def test_invitation_letter_edited_fields_before_and_after_the_signature_slot(sample_applicant, sample_companion):
    """Fields both BEFORE (e.g. inviterDetails, paragraph 12) and AFTER
    (e.g. signOffName/phone/email, paragraphs 21-23) the blank signature
    slot (paragraph 20) must all still land correctly even when an
    earlier edit changes the document's total paragraph count — the
    stable-paragraph-reference regression this test guards against."""
    skip_if_missing(ile.INVITATION_LETTER_TEMPLATE_PATH)
    edited = {
        "inviterDetails": "<p>Line one of a longer edit.</p><p>Line two of the same edit.</p>",
        "phone": "<p>Phone No.: +44 7000 000000</p>",
    }
    docx_bytes = ile.generate_invitation_letter_docx(
        INVITER, [sample_applicant, sample_companion], INVITATION_FIELDS, edited_fields=edited
    )
    text = _full_text(Document(io.BytesIO(docx_bytes)))
    assert "Line one of a longer edit." in text
    assert "Line two of the same edit." in text
    assert "+44 7000 000000" in text
    # Untouched fields after the edited/shifted range must still be correct.
    assert INVITER["email"] in text
    assert INVITER["fullName"] in text


def test_invitation_letter_edited_field_survives_into_the_pdf(sample_applicant, sample_companion):
    skip_if_missing(ile.INVITATION_LETTER_TEMPLATE_PATH)
    skip_if_missing_binary("soffice")
    edited = {"fundingArrangement": "<p>All costs are covered by the inviter alone.</p>"}
    pdf_bytes = ile.generate_invitation_letter_pdf(
        INVITER, [sample_applicant, sample_companion], INVITATION_FIELDS, edited_fields=edited
    )
    assert pdf_bytes[:5] == b"%PDF-"


# ---------------------------------------------------------------------------
# Initors Covering Letter
# ---------------------------------------------------------------------------

SUBJECT_FIELDS = {
    "country": "Canada",
    "purpose": "Tourism",
    "travelStartDate": "2026-07-01",
    "travelEndDate": "2026-07-20",
    "returnDate": "2026-07-21",
    "accommodationDetails": "Hotel booking attached",
    "otherCommitments": "Salaried employment in India",
    "invitationSupportingDocuments": "Invitation letter and hotel bookings",
}

INVITING_PERSON = {
    "fullName": "Test Inviting Person",
    "passportNumber": "CAT000002",
    "countryOfResidence": "Canada",
    "visaResidenceStatus": "Permanent Resident",
}


def test_initors_letter_requires_subject_and_spouse_names():
    with pytest.raises(ValueError):
        ile.generate_initors_covering_letter_docx({}, {"fullName": "Spouse"}, INVITING_PERSON, SUBJECT_FIELDS)
    with pytest.raises(ValueError):
        ile.generate_initors_covering_letter_docx({"fullName": "Subject", "sex": "F"}, {}, INVITING_PERSON, SUBJECT_FIELDS)


def test_initors_letter_rejects_an_unknown_sex():
    subject = {"fullName": "Ambiguous Subject", "sex": ""}
    spouse = {"fullName": "Spouse Person"}
    with pytest.raises(ValueError):
        ile.generate_initors_covering_letter_docx(subject, spouse, INVITING_PERSON, SUBJECT_FIELDS)


def test_initors_letter_picks_the_wife_voice_template_for_a_female_subject(sample_companion, sample_applicant):
    skip_if_missing(ile.INITORS_WIFE_VOICE_TEMPLATE_PATH)
    subject = dict(sample_companion, sex="F", relationWithInvitingPerson="Sister")
    spouse = sample_applicant
    docx_bytes = ile.generate_initors_covering_letter_docx(subject, spouse, INVITING_PERSON, SUBJECT_FIELDS)
    text = _full_text(Document(io.BytesIO(docx_bytes)))
    assert subject["fullName"] in text
    assert spouse["fullName"] in text
    assert INVITING_PERSON["fullName"] in text
    assert "harsha_sushil@yahoo.co.in" not in text, "the template's own leftover real hyperlink email must never leak"


def test_initors_letter_picks_the_husband_voice_template_for_a_male_subject(sample_applicant, sample_companion):
    skip_if_missing(ile.INITORS_HUSBAND_VOICE_TEMPLATE_PATH)
    subject = dict(sample_applicant, sex="M", relationWithInvitingPerson="Brother")
    spouse = sample_companion
    docx_bytes = ile.generate_initors_covering_letter_docx(subject, spouse, INVITING_PERSON, SUBJECT_FIELDS)
    text = _full_text(Document(io.BytesIO(docx_bytes)))
    assert subject["fullName"] in text
    assert spouse["fullName"] in text
    assert f"Email Id: {subject['email']}" in text


def test_initors_letter_pdf(sample_applicant, sample_companion):
    skip_if_missing(ile.INITORS_HUSBAND_VOICE_TEMPLATE_PATH)
    skip_if_missing_binary("soffice")
    subject = dict(sample_applicant, sex="M", relationWithInvitingPerson="Brother")
    pdf_bytes = ile.generate_initors_covering_letter_pdf(subject, sample_companion, INVITING_PERSON, SUBJECT_FIELDS)
    assert pdf_bytes[:5] == b"%PDF-"
