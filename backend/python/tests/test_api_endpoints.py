"""
HTTP-level tests for app.py, using FastAPI's TestClient (an in-process ASGI
client — no real network socket, no separately-running uvicorn process
required). These complement the engine-level unit tests above by covering
the actual request/response contract the frontend (js/core/api.js) relies
on: status codes, JSON/multipart request shapes, streamed
Content-Disposition filenames, and honest 400/500 error surfacing.

Every document-generation route here still goes through the real engine
against the real reference template (nothing is mocked), so a genuine
template-structure regression will fail here too, exactly as it does in
the engine-level tests.
"""

import io

import openpyxl
import pytest
from docx import Document
from fastapi.testclient import TestClient

import app as app_module
import data_model
import hotel_voucher_engine
from conftest import skip_if_missing

client = TestClient(app_module.app)


def _all_paragraphs_including_nested_tables(container):
    paragraphs = list(container.paragraphs)
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(_all_paragraphs_including_nested_tables(cell))
    return paragraphs


def _full_text(doc: Document) -> str:
    return "\n".join(p.text for p in _all_paragraphs_including_nested_tables(doc))


def test_health_endpoint():
    resp = client.get("/api/health")
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "ok"
    assert body["service"] == "khanna-backend"


def test_health_endpoint_reports_real_ocr_engine_status():
    # Whatever the real answer is on this machine (installed or not), the
    # health check must report the true, live state — never a hardcoded
    # "available": true (project rule 9). This asserts the shape and, on a
    # machine with no Google Cloud credentials set (this test's own
    # environment), that "available" actually matches whether the
    # Tesseract provider could find a working binary, not just that the
    # key exists.
    import ocr_providers

    resp = client.get("/api/health")
    body = resp.json()
    assert "ocrEngine" in body
    engine = body["ocrEngine"]
    assert isinstance(engine["available"], bool)
    assert engine["provider"] in ("tesseract", "google")
    if engine["provider"] == "tesseract":
        assert engine["available"] == (ocr_providers.find_tesseract_cmd() is not None)
        if not engine["available"]:
            assert "hint" in engine and "error" in engine


def test_status_list_endpoint_matches_data_model():
    resp = client.get("/api/applications/status-list")
    assert resp.status_code == 200
    assert resp.json()["statuses"] == data_model.STATUSES


def test_hotel_voucher_endpoint_streams_a_real_docx(sample_hotel):
    skip_if_missing(hotel_voucher_engine.TEMPLATE_PATH)
    resp = client.post(
        "/api/documents/hotel-voucher",
        json={"hotels": [sample_hotel], "format": "docx", "applicantName": "Test Applicant"},
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "attachment" in resp.headers["content-disposition"]
    assert "Test_Applicant" in resp.headers["content-disposition"]
    assert resp.content[:2] == b"PK"

    doc = Document(io.BytesIO(resp.content))
    full_text = _full_text(doc)
    assert sample_hotel["hotelName"] in full_text


def test_hotel_voucher_endpoint_rejects_an_empty_hotel_list():
    resp = client.post("/api/documents/hotel-voucher", json={"hotels": [], "format": "docx"})
    assert resp.status_code == 400


def test_hotel_voucher_endpoint_rejects_a_malformed_payload():
    resp = client.post("/api/documents/hotel-voucher", json={"not_hotels_at_all": True})
    assert resp.status_code == 422, "FastAPI/pydantic must reject a payload missing the required 'hotels' field"


def test_compose_hotel_fields_endpoint_seeds_from_the_same_real_data(sample_hotel):
    skip_if_missing(hotel_voucher_engine.TEMPLATE_PATH)
    resp = client.post("/api/documents/hotel-voucher/compose-fields", json={"hotels": [sample_hotel]})
    assert resp.status_code == 200
    body = resp.json()
    assert len(body["fields"]) == 1
    assert sample_hotel["hotelName"] in body["fields"][0]["hotelName"]
    assert sample_hotel["confirmationNumber"] in body["fields"][0]["confirmationSentence"]


def test_compose_hotel_fields_endpoint_rejects_an_empty_hotel_list():
    resp = client.post("/api/documents/hotel-voucher/compose-fields", json={"hotels": []})
    assert resp.status_code == 400


def test_hotel_voucher_endpoint_honors_edited_fields(sample_hotel):
    skip_if_missing(hotel_voucher_engine.TEMPLATE_PATH)
    hotel = dict(sample_hotel, editedFields={"hotelName": "<p>Edited Via HTTP Hotel</p>"})
    resp = client.post(
        "/api/documents/hotel-voucher",
        json={"hotels": [hotel], "format": "docx", "applicantName": "Test Applicant"},
    )
    assert resp.status_code == 200
    doc = Document(io.BytesIO(resp.content))
    full_text = _full_text(doc)
    assert "Edited Via HTTP Hotel" in full_text
    assert sample_hotel["hotelName"] not in full_text
    assert sample_hotel["confirmationNumber"] in full_text  # untouched field still auto-generated


def test_hotel_voucher_endpoint_accepts_a_null_edited_field_as_clear_override(sample_hotel):
    """js/hotels.js's own "Reset to auto-generated text" sends {key: null}
    (the frontend's clear-this-one-override patch shape, via
    KhannaState's deepMerge) rather than omitting the key — the request
    model must accept a null dict value here, not 422 on it."""
    skip_if_missing(hotel_voucher_engine.TEMPLATE_PATH)
    hotel = dict(sample_hotel, editedFields={"hotelName": None})
    resp = client.post(
        "/api/documents/hotel-voucher",
        json={"hotels": [hotel], "format": "docx", "applicantName": "Test Applicant"},
    )
    assert resp.status_code == 200
    doc = Document(io.BytesIO(resp.content))
    full_text = _full_text(doc)
    assert sample_hotel["hotelName"] in full_text  # null override == auto-generated, unchanged


def test_passport_authorization_endpoint(sample_applicant):
    from authorization_letter_engine import PASSPORT_SINGLE_TEMPLATE_PATH

    skip_if_missing(PASSPORT_SINGLE_TEMPLATE_PATH)
    resp = client.post(
        "/api/documents/passport-authorization",
        json={
            "people": [
                {
                    "fullName": sample_applicant["fullName"],
                    "passportNumber": sample_applicant["passportNumber"],
                    "sex": sample_applicant["sex"],
                    "phone": sample_applicant["phone"],
                    "email": sample_applicant["email"],
                }
            ],
            "recipientCentreName": "Test Visa Centre",
            "recipientAddress": "Test Address",
            "collectorName": "Test Collector",
            "format": "docx",
        },
    )
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"


def test_passport_authorization_endpoint_returns_400_for_missing_collector():
    resp = client.post(
        "/api/documents/passport-authorization",
        json={
            "people": [{"fullName": "Someone", "passportNumber": "X1"}],
            "recipientCentreName": "Test Centre",
            "collectorName": "",
            "format": "docx",
        },
    )
    assert resp.status_code == 400


def test_company_authorization_endpoint(sample_applicant):
    from authorization_letter_engine import COMPANY_TEMPLATE_PATH

    skip_if_missing(COMPANY_TEMPLATE_PATH)
    resp = client.post(
        "/api/documents/company-authorization",
        json={
            "people": [{"fullName": sample_applicant["fullName"], "passportNumber": sample_applicant["passportNumber"]}],
            "recipientText": "Visa Officer, Test Consulate",
            "format": "docx",
        },
    )
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"


def test_compose_passport_fields_endpoint(sample_applicant):
    from authorization_letter_engine import PASSPORT_SINGLE_TEMPLATE_PATH

    skip_if_missing(PASSPORT_SINGLE_TEMPLATE_PATH)
    resp = client.post(
        "/api/documents/passport-authorization/compose-fields",
        json={
            "people": [{"fullName": sample_applicant["fullName"], "passportNumber": sample_applicant["passportNumber"]}],
            "recipientCentreName": "Test Visa Centre",
            "recipientAddress": "Test Address",
            "collectorName": "Test Collector",
        },
    )
    assert resp.status_code == 200
    fields = resp.json()["fields"]
    assert sample_applicant["fullName"] in fields["applicantSentence"]
    assert "signature" not in fields


def test_passport_authorization_endpoint_honors_edited_fields(sample_applicant):
    from authorization_letter_engine import PASSPORT_SINGLE_TEMPLATE_PATH

    skip_if_missing(PASSPORT_SINGLE_TEMPLATE_PATH)
    resp = client.post(
        "/api/documents/passport-authorization",
        json={
            "people": [
                {
                    "fullName": sample_applicant["fullName"],
                    "passportNumber": sample_applicant["passportNumber"],
                    "phone": sample_applicant["phone"],
                    "email": sample_applicant["email"],
                }
            ],
            "recipientCentreName": "Test Visa Centre",
            "recipientAddress": "Test Address",
            "collectorName": "Test Collector",
            "format": "docx",
            "editedFields": {"applicantSentence": "<p>Edited authorization sentence via HTTP.</p>"},
        },
    )
    assert resp.status_code == 200
    doc = Document(io.BytesIO(resp.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Edited authorization sentence via HTTP." in text
    assert sample_applicant["email"] in text  # signature block untouched


def test_compose_company_fields_endpoint(sample_applicant):
    from authorization_letter_engine import COMPANY_TEMPLATE_PATH

    skip_if_missing(COMPANY_TEMPLATE_PATH)
    resp = client.post(
        "/api/documents/company-authorization/compose-fields",
        json={
            "people": [{"fullName": sample_applicant["fullName"], "passportNumber": sample_applicant["passportNumber"]}],
            "recipientText": "Visa Officer, Test Consulate",
        },
    )
    assert resp.status_code == 200
    fields = resp.json()["fields"]
    assert sample_applicant["fullName"] in fields["applicant_0"]


def test_company_authorization_endpoint_honors_edited_fields(sample_applicant):
    from authorization_letter_engine import COMPANY_TEMPLATE_PATH

    skip_if_missing(COMPANY_TEMPLATE_PATH)
    resp = client.post(
        "/api/documents/company-authorization",
        json={
            "people": [{"fullName": sample_applicant["fullName"], "passportNumber": sample_applicant["passportNumber"]}],
            "recipientText": "Visa Officer, Test Consulate",
            "format": "docx",
            "editedFields": {"applicant_0": "<p>Edited via HTTP applicant line</p>"},
        },
    )
    assert resp.status_code == 200
    doc = Document(io.BytesIO(resp.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Edited via HTTP applicant line" in text
    assert sample_applicant["fullName"] not in text


def test_cover_letter_endpoint_europe(sample_applicant, sample_companion):
    from cover_letter_engine import EUROPE_TEMPLATE_PATH

    skip_if_missing(EUROPE_TEMPLATE_PATH)
    applicant = dict(sample_applicant, isApplicant=True)
    companion = dict(sample_companion, isApplicant=False)
    resp = client.post(
        "/api/documents/cover-letter",
        json={
            "region": "Europe",
            "people": [applicant, companion],
            "recipientText": "Consulate General of France, Mumbai",
            "destinationCountry": "France",
            "travelStartDate": "2026-11-10",
            "travelEndDate": "2026-11-20",
            "fundingArrangement": "Applicant",
            "cityCountryOfResidence": "Mumbai, India",
            "numberOfNights": "10 nights",
            "applicantEmploymentStatus": "Salaried",
            "applicantJobTitle": "Software Engineer",
            "applicantEmployerName": "Acme Corp",
            "companionJobTitle": "Designer",
            "companionEmployerName": "Beta LLC",
            "companionEmploymentStartYear": "2015",
            "format": "docx",
        },
    )
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"


def test_cover_letter_endpoint_europe_rejects_wrong_companion_count(sample_applicant):
    resp = client.post(
        "/api/documents/cover-letter",
        json={
            "region": "Europe",
            "people": [dict(sample_applicant, isApplicant=True)],  # zero companions
            "format": "docx",
        },
    )
    assert resp.status_code == 400


def _europe_cover_letter_payload(sample_applicant, sample_companion, **overrides):
    applicant = dict(sample_applicant, isApplicant=True)
    companion = dict(sample_companion, isApplicant=False)
    payload = {
        "region": "Europe",
        "people": [applicant, companion],
        "recipientText": "Consulate General of France, Mumbai",
        "destinationCountry": "France",
        "travelStartDate": "2026-11-10",
        "travelEndDate": "2026-11-20",
        "fundingArrangement": "Applicant",
        "cityCountryOfResidence": "Mumbai, India",
        "numberOfNights": "10 nights",
        "applicantEmploymentStatus": "Salaried",
        "applicantJobTitle": "Software Engineer",
        "applicantEmployerName": "Acme Corp",
        "companionJobTitle": "Designer",
        "companionEmployerName": "Beta LLC",
        "companionEmploymentStartYear": "2015",
        "format": "docx",
    }
    payload.update(overrides)
    return payload


def test_cover_letter_compose_body_endpoint_returns_real_seed_html(sample_applicant, sample_companion):
    """Phase 4 proof of concept — the new JSON endpoint that seeds the
    Europe Cover Letter's Quill body editor with the SAME real composed
    wording the actual generator writes."""
    from cover_letter_engine import EUROPE_TEMPLATE_PATH

    skip_if_missing(EUROPE_TEMPLATE_PATH)
    resp = client.post(
        "/api/documents/cover-letter/compose-body",
        json=_europe_cover_letter_payload(sample_applicant, sample_companion),
    )
    assert resp.status_code == 200
    body = resp.json()
    assert sample_applicant["fullName"] in body["html"]
    assert body["html"].count("ql-align-justify") == 3


def test_cover_letter_compose_body_endpoint_rejects_non_europe_region(sample_applicant, sample_companion):
    resp = client.post(
        "/api/documents/cover-letter/compose-body",
        json=_europe_cover_letter_payload(sample_applicant, sample_companion, region="Japan"),
    )
    assert resp.status_code == 400


def test_cover_letter_compose_body_endpoint_rejects_wrong_companion_count(sample_applicant):
    resp = client.post(
        "/api/documents/cover-letter/compose-body",
        json={"region": "Europe", "people": [dict(sample_applicant, isApplicant=True)]},
    )
    assert resp.status_code == 400


def test_cover_letter_endpoint_with_edited_body_html_regenerates_the_real_docx(sample_applicant, sample_companion):
    """The existing /api/documents/cover-letter route now also accepts an
    optional editedBodyHtml (Phase 4) — omitted/None keeps the previous,
    unedited behaviour (the tests above already cover that); this exercises
    the new edited path end-to-end over real HTTP."""
    from cover_letter_engine import EUROPE_TEMPLATE_PATH

    skip_if_missing(EUROPE_TEMPLATE_PATH)
    edited_html = '<p class="ql-align-justify">A completely different, staff-edited body paragraph.</p>'
    resp = client.post(
        "/api/documents/cover-letter",
        json=_europe_cover_letter_payload(sample_applicant, sample_companion, editedBodyHtml=edited_html),
    )
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"
    text = "\n".join(p.text for p in Document(io.BytesIO(resp.content)).paragraphs)
    assert "A completely different, staff-edited body paragraph." in text
    # The signature block (after the shrunken body range) is still real.
    assert sample_applicant["fullName"] in text
    assert sample_applicant["email"] in text


def test_invitation_letter_endpoint_is_multipart_with_a_json_payload_field(sample_applicant, sample_companion):
    from invitation_letter_engine import INVITATION_LETTER_TEMPLATE_PATH

    skip_if_missing(INVITATION_LETTER_TEMPLATE_PATH)
    import json

    payload = {
        "inviter": {"fullName": "Test Inviter Abroad", "email": "inviter@example.com"},
        "invitees": [
            {"fullName": sample_applicant["fullName"], "passportNumber": sample_applicant["passportNumber"]},
            {"fullName": sample_companion["fullName"], "passportNumber": sample_companion["passportNumber"]},
        ],
        "travelStartDate": "2026-08-01",
        "travelEndDate": "2026-08-15",
        "purpose": "Tourism",
        "format": "docx",
    }
    resp = client.post(
        "/api/documents/invitation-letter",
        data={"payload": json.dumps(payload)},
    )
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"


def test_invitation_letter_endpoint_rejects_malformed_json_payload_field():
    resp = client.post("/api/documents/invitation-letter", data={"payload": "not valid json"})
    assert resp.status_code == 400


def test_compose_invitation_fields_endpoint(sample_applicant, sample_companion):
    from invitation_letter_engine import INVITATION_LETTER_TEMPLATE_PATH

    skip_if_missing(INVITATION_LETTER_TEMPLATE_PATH)
    resp = client.post(
        "/api/documents/invitation-letter/compose-fields",
        json={
            "inviter": {"fullName": "Test Inviter Abroad", "email": "inviter@example.com"},
            "invitees": [
                {"fullName": sample_applicant["fullName"], "passportNumber": sample_applicant["passportNumber"]},
                {"fullName": sample_companion["fullName"], "passportNumber": sample_companion["passportNumber"]},
            ],
            "purpose": "Tourism",
        },
    )
    assert resp.status_code == 200
    fields = resp.json()["fields"]
    assert sample_applicant["fullName"] in fields["inviteesDetails"]
    assert "Tourism" in fields["purposeStatement"]


def test_invitation_letter_endpoint_honors_edited_fields(sample_applicant, sample_companion):
    from invitation_letter_engine import INVITATION_LETTER_TEMPLATE_PATH

    skip_if_missing(INVITATION_LETTER_TEMPLATE_PATH)
    import json

    payload = {
        "inviter": {"fullName": "Test Inviter Abroad", "email": "inviter@example.com"},
        "invitees": [
            {"fullName": sample_applicant["fullName"], "passportNumber": sample_applicant["passportNumber"]},
            {"fullName": sample_companion["fullName"], "passportNumber": sample_companion["passportNumber"]},
        ],
        "travelStartDate": "2026-08-01",
        "travelEndDate": "2026-08-15",
        "purpose": "Tourism",
        "format": "docx",
        "editedFields": {"fundingArrangement": "<p>Edited funding sentence via HTTP.</p>"},
    }
    resp = client.post(
        "/api/documents/invitation-letter",
        data={"payload": json.dumps(payload)},
    )
    assert resp.status_code == 200
    from docx import Document

    doc = Document(io.BytesIO(resp.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Edited funding sentence via HTTP." in text
    assert "inviter@example.com" in text  # untouched field after the edited one still auto-generates


def _required_document_letter_payload(**overrides):
    payload = {
        "country": "Germany",
        "visaType": "Tourist",
        "employmentStatus": "Employed",
        "documentIds": ["passport_copy", "bank_statement"],
        "bankStatementDuration": "6 Months",
        "itrYears": "2 Years",
        "processingTime": "15",
        "visaFees": "5,000",
        "format": "docx",
    }
    payload.update(overrides)
    return payload


def test_required_document_letter_endpoint():
    from required_document_letter_engine import TEMPLATE_PATH

    skip_if_missing(TEMPLATE_PATH)
    resp = client.post("/api/documents/required-document-letter", json=_required_document_letter_payload())
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"


def test_compose_required_document_letter_fields_endpoint():
    from required_document_letter_engine import TEMPLATE_PATH

    skip_if_missing(TEMPLATE_PATH)
    resp = client.post(
        "/api/documents/required-document-letter/compose-fields",
        json={
            "country": "Germany",
            "visaType": "Tourist",
            "employmentStatus": "Employed",
            "documentIds": ["passport_copy", "bank_statement"],
            "bankStatementDuration": "6 Months",
            "itrYears": "2 Years",
            "processingTime": "15",
            "visaFees": "5,000",
        },
    )
    assert resp.status_code == 200
    fields = resp.json()["fields"]
    assert "Germany" in fields["countryLine"]
    assert "Passport Copy" in fields["checklistItems"]


def test_required_document_letter_endpoint_honors_edited_fields():
    from required_document_letter_engine import TEMPLATE_PATH

    skip_if_missing(TEMPLATE_PATH)
    payload = _required_document_letter_payload(
        editedFields={"checklistItems": "<p>A staff-edited checklist line.</p>"}
    )
    resp = client.post("/api/documents/required-document-letter", json=payload)
    assert resp.status_code == 200
    from docx import Document

    doc = Document(io.BytesIO(resp.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "A staff-edited checklist line." in text
    assert "Passport Copy" not in text
    assert "Country: Germany" in text  # untouched field still auto-generates


def test_admin_export_excel_endpoint():
    app_record = data_model.create_empty_application(created_by="qa@khannatravels.com")
    app_record["applicant"]["fullName"] = "Endpoint Test Applicant"
    resp = client.post("/api/admin/export-excel", json={"applications": [app_record]})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert len(wb.sheetnames) == 2  # Applications + Travellers (status sheet removed)


def test_admin_export_excel_endpoint_with_empty_applications_list():
    resp = client.post("/api/admin/export-excel", json={"applications": []})
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"


def test_passport_process_endpoint_rejects_unsupported_file_type():
    resp = client.post(
        "/api/passport/process",
        files={"file": ("not_a_passport.txt", b"plain text, not an image", "text/plain")},
    )
    assert resp.status_code == 400


def test_passport_process_endpoint_rejects_empty_file():
    resp = client.post(
        "/api/passport/process",
        files={"file": ("empty.png", b"", "image/png")},
    )
    assert resp.status_code == 400
